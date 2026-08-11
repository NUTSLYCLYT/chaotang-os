from __future__ import annotations

import struct
from collections import Counter
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
from typing import Final

import xlrd
from openpyxl import load_workbook
from xlrd.compdoc import CompDoc

from .models import FinancialStatementRow, NormalizedLedgerRow, SourceRef
from .source_manifest import (
    BALANCE_LEDGER_V1,
    FINANCIAL_STATEMENTS_V1,
    AccountingSourceFile,
    AccountingSourceManifest,
    SourceRole,
)
from .sources import AccountingSourceError

_BALANCE_HEADER: Final = {
    (7, 2): "科目类别",
    (7, 3): "科目编码",
    (7, 4): "科目名称",
    (7, 5): "期初余额",
    (7, 7): "本期发生",
    (7, 9): "期末余额",
    (8, 5): "借方",
    (8, 6): "贷方",
    (8, 7): "借方",
    (8, 8): "贷方",
    (8, 9): "借方",
    (8, 10): "贷方",
}
_STATEMENT_HEADERS: Final = {
    "资产负债表": (
        "资产",
        "行次",
        "期末余额",
        "年初余额",
        "负债和所有者权益",
        "行次",
        "期末余额",
        "年初余额",
    ),
    "利润表": ("项目", "行次", "本年累计金额", "本月金额"),
    "现金流量表": ("项目", "行次", "本年累计金额", "本月金额"),
}
_OBSERVED_FIRST_HEADERS: Final = {
    "资产负债表": {"资产", "资    产"},
    "利润表": {"项目", "项  目"},
    "现金流量表": {"项目", "项       目"},
}


@dataclass(frozen=True, slots=True, repr=False)
class AccountingDataset:
    manifest: AccountingSourceManifest
    ledger_rows: tuple[NormalizedLedgerRow, ...]
    statement_rows: tuple[FinancialStatementRow, ...]
    subject_identity: str | None = None

    def __post_init__(self) -> None:
        if not self.ledger_rows:
            raise AccountingSourceError("source_schema_invalid")
        expected = set(range(self.manifest.period.start_year, self.manifest.period.end_year + 1))
        if {row.year for row in self.ledger_rows} != expected:
            raise AccountingSourceError("source_period_conflict")
        if any(row.year not in expected for row in self.statement_rows):
            raise AccountingSourceError("source_period_conflict")


def _text(value: object, path: Path) -> str:
    if value is None or isinstance(value, bool):
        raise AccountingSourceError("source_schema_invalid", path.name)
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    rendered = str(value).strip()
    if not rendered or rendered.startswith("="):
        raise AccountingSourceError("source_schema_invalid", path.name)
    return rendered


def _amount(value: object, path: Path) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, bool):
        raise AccountingSourceError("source_schema_invalid", path.name)
    rendered = str(value).strip().replace(",", "")
    if rendered.startswith("="):
        raise AccountingSourceError("source_schema_invalid", path.name)
    try:
        result = Decimal(rendered)
    except (InvalidOperation, ValueError):
        raise AccountingSourceError("source_schema_invalid", path.name) from None
    if not result.is_finite():
        raise AccountingSourceError("source_schema_invalid", path.name)
    return result


def _source_ref(item: AccountingSourceFile, sheet: str, row: int) -> SourceRef:
    return SourceRef(item.basename, sheet, row, item.sha256)


def load_balance_ledger(
    item: AccountingSourceFile,
) -> tuple[NormalizedLedgerRow, ...]:
    if item.role is not SourceRole.BALANCE_LEDGER or item.schema_id != BALANCE_LEDGER_V1:
        raise AccountingSourceError("source_schema_invalid", item.basename)
    try:
        source = BytesIO(item.content) if item.content is not None else item.path
        workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    except Exception:
        raise AccountingSourceError("source_schema_invalid", item.basename) from None
    try:
        if len(workbook.worksheets) != 1:
            raise AccountingSourceError("source_schema_invalid", item.basename)
        sheet = workbook.worksheets[0]
        for row_number in (7, 8):
            for column in range(1, max(sheet.max_column, 10) + 1):
                value = sheet.cell(row_number, column).value
                expected = _BALANCE_HEADER.get((row_number, column))
                if expected is None:
                    if value not in (None, ""):
                        raise AccountingSourceError("source_schema_invalid", item.basename)
                elif str(value).strip() != expected:
                    raise AccountingSourceError("source_schema_invalid", item.basename)
        normalized: list[NormalizedLedgerRow] = []
        for row_number in range(9, sheet.max_row + 1):
            values = [sheet.cell(row_number, column).value for column in range(1, 11)]
            if not any(value not in (None, "") for value in values):
                continue
            if values[0] not in (None, "") or any(
                sheet.cell(row_number, column).value not in (None, "")
                for column in range(11, sheet.max_column + 1)
            ):
                raise AccountingSourceError("source_schema_invalid", item.basename)
            if any(sheet.cell(row_number, column).data_type == "f" for column in range(2, 11)):
                raise AccountingSourceError("source_schema_invalid", item.basename)
            normalized.append(
                NormalizedLedgerRow(
                    year=item.year,
                    category=_text(values[1], item.path),
                    account_code=_text(values[2], item.path),
                    account_name=_text(values[3], item.path),
                    opening_debit=_amount(values[4], item.path),
                    opening_credit=_amount(values[5], item.path),
                    movement_debit=_amount(values[6], item.path),
                    movement_credit=_amount(values[7], item.path),
                    closing_debit=_amount(values[8], item.path),
                    closing_credit=_amount(values[9], item.path),
                    source=_source_ref(item, sheet.title, row_number),
                )
            )
        if not normalized:
            raise AccountingSourceError("source_schema_invalid", item.basename)
        return tuple(normalized)
    except AccountingSourceError:
        raise
    except Exception:
        raise AccountingSourceError("source_schema_invalid", item.basename) from None
    finally:
        workbook.close()


def _workbook_stream(source: bytes) -> bytes:
    try:
        compound = CompDoc(source, logfile=StringIO())
        for name in ("Workbook", "Book"):
            stream = compound.get_named_stream(name)
            if stream is not None:
                return stream
    except Exception:
        pass
    raise AccountingSourceError("source_schema_invalid")


def _reject_biff_formulas(source: bytes) -> None:
    """Reject actual BIFF FORMULA records in the OLE Workbook stream."""
    stream = _workbook_stream(source)
    offset = 0
    while offset + 4 <= len(stream):
        record_id, payload_size = struct.unpack_from("<HH", stream, offset)
        record_end = offset + 4 + payload_size
        if record_end > len(stream):
            raise AccountingSourceError("source_schema_invalid")
        if record_id == 0x0006:
            raise AccountingSourceError("source_schema_invalid")
        offset = record_end


def _read_xls_sheets(
    source: bytes | Path,
) -> tuple[tuple[str, tuple[tuple[object, ...], ...]], ...]:
    try:
        content = source if isinstance(source, bytes) else source.read_bytes()
        _reject_biff_formulas(content)
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    except Exception:
        raise AccountingSourceError("source_schema_invalid") from None
    try:
        return tuple(
            (
                sheet.name,
                tuple(
                    tuple(sheet.cell_value(row, column) for column in range(sheet.ncols))
                    for row in range(sheet.nrows)
                ),
            )
            for sheet in workbook.sheets()
        )
    except Exception:
        raise AccountingSourceError("source_schema_invalid") from None
    finally:
        workbook.release_resources()


def _normalized_cells(
    row: tuple[object, ...], sheet_name: str
) -> tuple[str, ...]:
    normalized: list[str] = []
    for column, value in enumerate(row):
        rendered = "" if value is None else str(value).strip()
        if column == 0 and rendered in _OBSERVED_FIRST_HEADERS[sheet_name]:
            rendered = _STATEMENT_HEADERS[sheet_name][0]
        normalized.append(rendered)
    return tuple(normalized)


def _find_header(
    rows: tuple[tuple[object, ...], ...],
    expected: tuple[str, ...],
    path: Path,
    sheet_name: str,
) -> int:
    matches: list[int] = []
    for index, row in enumerate(rows[:8]):
        normalized = _normalized_cells(row, sheet_name)
        prefix = normalized[: len(expected)]
        if prefix == expected and not any(normalized[len(expected) :]):
            matches.append(index)
    if len(matches) != 1:
        raise AccountingSourceError("source_schema_invalid", path.name)
    return matches[0]


def _statement_row(
    item: AccountingSourceFile,
    sheet_name: str,
    row_number: int,
    section: str,
    values: tuple[object, object, object, object],
) -> FinancialStatementRow | None:
    if not any(value not in (None, "") for value in values):
        return None
    if any(value in (None, "") for value in values):
        raise AccountingSourceError("source_schema_invalid", item.basename)
    return FinancialStatementRow(
        year=item.year,
        sheet_name=sheet_name,
        row_number=row_number,
        section=section,
        item_name=_text(values[0], item.path),
        line_number=_text(values[1], item.path),
        current_amount=_amount(values[2], item.path),
        comparison_amount=_amount(values[3], item.path),
        source=_source_ref(item, sheet_name, row_number),
    )


def load_financial_statements(
    item: AccountingSourceFile,
) -> tuple[FinancialStatementRow, ...]:
    if (
        item.role is not SourceRole.FINANCIAL_STATEMENTS
        or item.schema_id != FINANCIAL_STATEMENTS_V1
    ):
        raise AccountingSourceError("source_schema_invalid", item.basename)
    sheets = _read_xls_sheets(item.content if item.content is not None else item.path)
    names = [name for name, _ in sheets]
    if Counter(names) != Counter(_STATEMENT_HEADERS.keys()):
        raise AccountingSourceError("source_schema_invalid", item.basename)
    normalized: list[FinancialStatementRow] = []
    try:
        for sheet_name, rows in sheets:
            expected = _STATEMENT_HEADERS[sheet_name]
            header_index = _find_header(rows, expected, item.path, sheet_name)
            for zero_index, raw in enumerate(rows[header_index + 1 :], start=header_index + 1):
                row_number = zero_index + 1
                if any(value not in (None, "") for value in raw[len(expected) :]):
                    raise AccountingSourceError("source_schema_invalid", item.basename)
                padded = raw + (None,) * max(0, len(expected) - len(raw))
                if sheet_name == "资产负债表":
                    for section, offset in (("assets", 0), ("liabilities_equity", 4)):
                        result = _statement_row(
                            item,
                            sheet_name,
                            row_number,
                            section,
                            tuple(padded[offset : offset + 4]),
                        )
                        if result is not None:
                            normalized.append(result)
                else:
                    result = _statement_row(
                        item,
                        sheet_name,
                        row_number,
                        "main",
                        tuple(padded[:4]),
                    )
                    if result is not None:
                        normalized.append(result)
    except AccountingSourceError:
        raise
    except Exception:
        raise AccountingSourceError("source_schema_invalid", item.basename) from None
    if not normalized:
        raise AccountingSourceError("source_schema_invalid", item.basename)
    return tuple(normalized)


def load_accounting_dataset(manifest: AccountingSourceManifest) -> AccountingDataset:
    ledger_rows: list[NormalizedLedgerRow] = []
    statement_rows: list[FinancialStatementRow] = []
    for item in manifest.files:
        if item.role is SourceRole.BALANCE_LEDGER:
            ledger_rows.extend(load_balance_ledger(item))
        elif item.role is SourceRole.FINANCIAL_STATEMENTS:
            statement_rows.extend(load_financial_statements(item))
        else:
            raise AccountingSourceError("source_schema_invalid")
    return AccountingDataset(manifest, tuple(ledger_rows), tuple(statement_rows))


def inspect_accounting_source_dimensions(
    item: AccountingSourceFile,
) -> tuple[int, int]:
    if item.role is SourceRole.BALANCE_LEDGER:
        try:
            source = BytesIO(item.content) if item.content is not None else item.path
            workbook = load_workbook(
                source, read_only=True, data_only=False, keep_links=False
            )
        except Exception:
            raise AccountingSourceError("source_schema_invalid", item.basename) from None
        try:
            return len(workbook.worksheets), sum(
                sheet.max_row for sheet in workbook.worksheets
            )
        finally:
            workbook.close()
    if item.role is SourceRole.FINANCIAL_STATEMENTS:
        sheets = _read_xls_sheets(item.content if item.content is not None else item.path)
        return len(sheets), sum(len(rows) for _, rows in sheets)
    raise AccountingSourceError("source_schema_invalid", item.basename)


def preflight_accounting_sources(
    source_dir: Path,
    period,
) -> AccountingDataset:
    from .source_manifest import build_accounting_source_manifest

    return load_accounting_dataset(build_accounting_source_manifest(source_dir, period))
