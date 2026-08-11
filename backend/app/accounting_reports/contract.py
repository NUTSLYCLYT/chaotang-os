"""Deterministic accounting evaluator for work-product contracts."""

from __future__ import annotations

import hashlib
import re
from collections.abc import Sequence
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from app.work_products import (
    ArtifactGateReceipt,
    ArtifactGateStatus,
    ArtifactManifestItem,
    WorkProductStatus,
    evaluate_artifact_gate,
    semantic_digest,
)

from .analysis import _family, _oriented_amount, signed_closing
from .models import (
    AccountingEvaluation,
    AccountingFact,
    AccountingReportNumber,
    AccountingSourceReceipt,
    CurrencyConversionReceipt,
    NormalizedLedgerRow,
    ReportAnalysis,
    ReportCheck,
    ReportPeriod,
)

_REVIEW_REASONS = frozenset(
    {
        "DUPLICATE_ENTRY",
        "DUPLICATE_SOURCE",
        "CROSS_PERIOD_ENTRY",
        "CURRENCY_MISMATCH",
        "STALE_SOURCE",
        "SOURCE_CONFLICT",
        "MISSING_SOURCE_LABEL",
        "UNADOPTED_SOURCE",
        "ACCOUNTING_IDENTITY_MISMATCH",
        "IMPOSSIBLE_RATIO",
        "UNBACKED_REPORT_NUMBER",
    }
)
_REASON_ORDER = (
    "FORBIDDEN_EXTERNAL_ACTION",
    "MISSING_REQUIRED_FIELD",
    "INVALID_NUMERIC_VALUE",
    "MISSING_RULE_VERSION",
    "DUPLICATE_ENTRY",
    "DUPLICATE_SOURCE",
    "CROSS_PERIOD_ENTRY",
    "STALE_SOURCE",
    "CURRENCY_MISMATCH",
    "SOURCE_CONFLICT",
    "MISSING_SOURCE_LABEL",
    "UNADOPTED_SOURCE",
    "ACCOUNTING_IDENTITY_MISMATCH",
    "IMPOSSIBLE_RATIO",
    "UNBACKED_REPORT_NUMBER",
    "ZERO_DENOMINATOR_EXPLAINED",
)
_REASON_OUTPUT_KIND = {
    "FORBIDDEN_EXTERNAL_ACTION": "forbidden_external_action_notice",
    "MISSING_REQUIRED_FIELD": "missing_required_field_notice",
    "INVALID_NUMERIC_VALUE": "invalid_numeric_value_notice",
    "MISSING_RULE_VERSION": "missing_rule_version_notice",
    "DUPLICATE_ENTRY": "duplicate_entry_notice",
    "DUPLICATE_SOURCE": "duplicate_source_notice",
    "CROSS_PERIOD_ENTRY": "cross_period_entry_notice",
    "CURRENCY_MISMATCH": "currency_mismatch_notice",
    "STALE_SOURCE": "stale_source_notice",
    "SOURCE_CONFLICT": "source_conflict_notice",
    "MISSING_SOURCE_LABEL": "missing_source_label_notice",
    "UNADOPTED_SOURCE": "unadopted_source_notice",
    "ACCOUNTING_IDENTITY_MISMATCH": "accounting_identity_mismatch_notice",
    "IMPOSSIBLE_RATIO": "impossible_ratio_notice",
    "UNBACKED_REPORT_NUMBER": "unbacked_report_number_notice",
    "ZERO_DENOMINATOR_EXPLAINED": "zero_denominator_explanation",
}
ACCOUNTING_REQUIRED_ARTIFACT_KINDS = frozenset(
    {
        "management_report_xlsx",
        "work_product_envelope",
        "quality_report",
        "confirmation_request",
        "content_digest",
    }
)
_REQUIRED_QUALITY_NAMES = frozenset(
    {
        "coverage",
        "mapping",
        "balance_sheet_equation",
        "opening_continuity",
        "movement_balance",
        "account_directions",
    }
)
_SHA256_PATTERN = re.compile(r"^[0-9a-f]{64}$")
_SOURCE_REF_PATTERN = re.compile(
    r"^(?P<digest>[0-9a-f]{64}):(?P<sheet>[^:\x00]{1,128}):"
    r"(?P<row>[1-9][0-9]{0,8})$"
)


def _ledger_fact_id(
    *, source_ref: str, account_code: str, amount: Decimal
) -> str:
    return "fact:" + semantic_digest(
        {
            "source_ref": source_ref,
            "account_code": account_code,
            "amount": amount,
        }
    )[:24]


def _has_canonical_fact_semantics(
    fact: AccountingFact,
    accounting_rules_version: str,
) -> bool:
    if (
        not isinstance(fact, AccountingFact)
        or not isinstance(fact.fact_id, str)
        or not isinstance(fact.name, str)
        or not isinstance(fact.amount, Decimal)
        or not isinstance(fact.source_ref, str)
        or not isinstance(fact.source_digest, str)
        or not isinstance(fact.rule_id, str)
        or not isinstance(fact.source_human_confirmed, bool)
        or (
            fact.conversion_receipt is not None
            and not isinstance(fact.conversion_receipt, CurrencyConversionReceipt)
        )
    ):
        return False
    source_match = _SOURCE_REF_PATTERN.fullmatch(fact.source_ref)
    if (
        source_match is None
        or source_match.group("digest") != fact.source_digest
        or _SHA256_PATTERN.fullmatch(fact.source_digest) is None
        or not fact.name.strip()
        or len(fact.name) > 128
        or fact.rule_id != f"{accounting_rules_version}:ledger-closing"
        or not fact.amount.is_finite()
    ):
        return False
    expected_fact_id = _ledger_fact_id(
        source_ref=fact.source_ref,
        account_code=fact.name,
        amount=fact.amount,
    )
    return fact.fact_id == expected_fact_id


def evaluate_accounting_artifact_gate(
    *,
    artifacts: Sequence[ArtifactManifestItem],
    workbook_path: Path | None = None,
    expected_file_sha256: str | None = None,
    expected_fact_pack_digest: str | None = None,
    facts: Sequence[AccountingFact] = (),
    rows: Sequence[NormalizedLedgerRow] = (),
    source_receipts: Sequence[AccountingSourceReceipt] = (),
    accounting_rules_version: str | None = None,
    quality_checks: Sequence[ReportCheck] = (),
) -> ArtifactGateReceipt:
    """Fail closed on accounting evidence before applying the shared gate."""

    domain_reasons: list[str] = []
    management_items = tuple(
        item for item in artifacts if item.kind == "management_report_xlsx"
    )
    if workbook_path is None or expected_file_sha256 is None:
        domain_reasons.append("ACCOUNTING_WORKBOOK_EVIDENCE_MISSING")
    else:
        try:
            path = Path(workbook_path)
            actual_file_sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                from .workbook import SHEET_NAMES

                if tuple(workbook.sheetnames) != SHEET_NAMES:
                    domain_reasons.append("ACCOUNTING_SHEET_SET_INVALID")
            finally:
                workbook.close()
        except Exception:
            actual_file_sha256 = None
            domain_reasons.append("ACCOUNTING_WORKBOOK_INVALID")
        if (
            actual_file_sha256 != expected_file_sha256
            or len(management_items) != 1
            or management_items[0].content_digest != expected_file_sha256
        ):
            domain_reasons.append("ACCOUNTING_FILE_HASH_MISMATCH")

    canonical_facts = _canonical_ledger_facts(
        rows=rows,
        source_receipts=source_receipts,
        accounting_rules_version=accounting_rules_version,
    )
    if (
        not facts
        or canonical_facts is None
        or tuple(facts) != canonical_facts
        or len({fact.fact_id for fact in facts}) != len(facts)
        or any(
            not _has_canonical_fact_semantics(
                fact,
                accounting_rules_version or "",
            )
            for fact in facts
        )
    ):
        domain_reasons.append("ACCOUNTING_FACT_LINEAGE_MISSING")

    quality_names = tuple(check.name for check in quality_checks)
    if (
        len(quality_names) != len(_REQUIRED_QUALITY_NAMES)
        or frozenset(quality_names) != _REQUIRED_QUALITY_NAMES
        or len(set(quality_names)) != len(quality_names)
        or any(check.status != "PASS" for check in quality_checks)
    ):
        domain_reasons.append("ACCOUNTING_QUALITY_FAILED")

    content_digest_items = tuple(
        item for item in artifacts if item.kind == "content_digest"
    )
    try:
        actual_fact_pack_digest = semantic_digest({"facts": _fact_payload(facts)})
    except Exception:
        actual_fact_pack_digest = None
    if (
        expected_fact_pack_digest is None
        or _SHA256_PATTERN.fullmatch(expected_fact_pack_digest) is None
        or len(content_digest_items) != 1
        or actual_fact_pack_digest != expected_fact_pack_digest
        or content_digest_items[0].content_digest != expected_fact_pack_digest
    ):
        domain_reasons.append("ACCOUNTING_FACT_PACK_DIGEST_MISMATCH")

    shared = evaluate_artifact_gate(
        required_kinds=ACCOUNTING_REQUIRED_ARTIFACT_KINDS,
        artifacts=artifacts,
    )
    reason_codes = tuple(dict.fromkeys((*shared.reason_codes, *domain_reasons)))
    return ArtifactGateReceipt(
        status=(
            ArtifactGateStatus.FAILED
            if reason_codes
            else ArtifactGateStatus.PASSED
        ),
        reason_codes=reason_codes,
        missing_kinds=shared.missing_kinds,
        unexpected_kinds=shared.unexpected_kinds,
    )


def _source_ref(row: NormalizedLedgerRow) -> str:
    return (
        f"{row.source.file_sha256}:{row.source.sheet_name}:"
        f"{row.source.row_number}"
    )


def _row_amount(row: NormalizedLedgerRow) -> Decimal:
    family = _family(row)
    return signed_closing(row) if family is None else _oriented_amount(row, family)


def _canonical_ledger_facts(
    *,
    rows: Sequence[NormalizedLedgerRow],
    source_receipts: Sequence[AccountingSourceReceipt],
    accounting_rules_version: str | None,
) -> tuple[AccountingFact, ...] | None:
    if (
        not isinstance(accounting_rules_version, str)
        or not accounting_rules_version.strip()
        or not rows
        or any(not isinstance(row, NormalizedLedgerRow) for row in rows)
        or any(
            not isinstance(receipt, AccountingSourceReceipt)
            for receipt in source_receipts
        )
    ):
        return None
    row_refs = tuple(_source_ref(row) for row in rows)
    receipt_refs = tuple(receipt.source_ref for receipt in source_receipts)
    if (
        len(set(row_refs)) != len(row_refs)
        or len(set(receipt_refs)) != len(receipt_refs)
        or set(row_refs) != set(receipt_refs)
    ):
        return None
    receipts_by_ref = {
        receipt.source_ref: receipt for receipt in source_receipts
    }
    canonical: list[AccountingFact] = []
    for row, source_ref in zip(rows, row_refs, strict=True):
        receipt = receipts_by_ref[source_ref]
        if (
            not receipt.adopted
            or receipt.source_digest != row.source.file_sha256
        ):
            return None
        amount = _row_amount(row)
        canonical.append(
            AccountingFact(
                fact_id=_ledger_fact_id(
                    source_ref=source_ref,
                    account_code=row.account_code,
                    amount=amount,
                ),
                name=row.account_code,
                amount=amount,
                source_ref=source_ref,
                source_digest=receipt.source_digest,
                rule_id=f"{accounting_rules_version}:ledger-closing",
                source_human_confirmed=receipt.source_human_confirmed,
                conversion_receipt=receipt.conversion_receipt,
            )
        )
    return tuple(canonical)


def _family_total(rows: Sequence[NormalizedLedgerRow], names: frozenset[str]) -> Decimal:
    return sum(
        (_row_amount(row) for row in rows if _family(row) in names),
        Decimal("0"),
    )


def _receipt_selection_key(
    receipt: AccountingSourceReceipt,
) -> tuple[object, ...]:
    conversion = receipt.conversion_receipt
    conversion_key: tuple[object, ...] = (
        ()
        if conversion is None
        else (
            conversion.from_currency,
            conversion.to_currency,
            conversion.rate,
            conversion.period.start_year,
            conversion.period.end_year,
            conversion.source_ref,
            conversion.source_digest,
            conversion.adopted,
        )
    )
    return (
        not receipt.adopted,
        receipt.source_label is None,
        receipt.source_label or "",
        receipt.currency,
        receipt.period.start_year,
        receipt.period.end_year,
        not receipt.source_human_confirmed,
        receipt.logical_source_id,
        receipt.source_ref,
        receipt.source_digest,
        conversion_key,
    )


def _currency_code(value: str) -> str:
    return value.strip().upper()


def _currencies_are_reconciled(
    receipts: Sequence[AccountingSourceReceipt],
    *,
    period: ReportPeriod,
) -> bool:
    adopted = tuple(receipt for receipt in receipts if receipt.adopted)
    currencies = {_currency_code(receipt.currency) for receipt in adopted}
    if len(currencies) <= 1:
        return True
    for target_currency in sorted(currencies):
        if all(
            _currency_code(receipt.currency) == target_currency
            or (
                receipt.conversion_receipt is not None
                and receipt.conversion_receipt.adopted
                and receipt.conversion_receipt.period == period
                and _currency_code(
                    receipt.conversion_receipt.from_currency
                )
                == _currency_code(receipt.currency)
                and _currency_code(receipt.conversion_receipt.to_currency)
                == target_currency
            )
            for receipt in adopted
        ):
            return True
    return False


def _conversion_payload(
    conversion: CurrencyConversionReceipt | None,
) -> dict[str, object] | None:
    if conversion is None:
        return None
    return {
        "from_currency": conversion.from_currency,
        "to_currency": conversion.to_currency,
        "rate": conversion.rate,
        "period": {
            "start_year": conversion.period.start_year,
            "end_year": conversion.period.end_year,
        },
        "source_ref": conversion.source_ref,
        "source_digest": conversion.source_digest,
        "adopted": conversion.adopted,
    }


def _fact_payload(
    facts: Sequence[AccountingFact],
) -> tuple[dict[str, object], ...]:
    return tuple(
        {
            "fact_id": fact.fact_id,
            "name": fact.name,
            "amount": fact.amount,
            "source_ref": fact.source_ref,
            "source_digest": fact.source_digest,
            "rule_id": fact.rule_id,
            "source_human_confirmed": fact.source_human_confirmed,
            "conversion_receipt": _conversion_payload(fact.conversion_receipt),
        }
        for fact in facts
    )


def _is_backed_number(
    number: AccountingReportNumber,
    *,
    row_amounts: dict[str, Decimal],
    rows_by_ref: dict[str, NormalizedLedgerRow],
    valid_sources: frozenset[str],
    accounting_rules_version: str,
) -> bool:
    if number.model_generated or not isinstance(number.amount, Decimal):
        return False
    if not number.fact_id or not number.source_ref or not number.rule_id:
        return False
    if number.source_ref not in valid_sources:
        return False
    if not number.rule_id.startswith(f"{accounting_rules_version}:"):
        return False
    if number.name == "profit_margin":
        return True
    if (
        number.name == "cash_flow"
        and number.rule_id == f"{accounting_rules_version}:cash-flow-movement"
    ):
        row = rows_by_ref.get(number.source_ref)
        return (
            row is not None
            and row.movement_debit - row.movement_credit == number.amount
        )
    return row_amounts.get(number.source_ref) == number.amount


def _terminal_evaluation(
    *,
    task_id: str,
    period: ReportPeriod,
    status: WorkProductStatus,
    reasons: tuple[str, ...],
    source_receipts: Sequence[AccountingSourceReceipt],
    accounting_rules_version: str,
    renderer_version: str,
) -> AccountingEvaluation:
    return AccountingEvaluation(
        task_id=task_id,
        period=period,
        work_status=status,
        reason_codes=reasons,
        facts=(),
        fact_pack_digest=semantic_digest({"facts": ()}),
        source_digests=tuple(sorted(receipt.source_digest for receipt in source_receipts)),
        accounting_rules_version=accounting_rules_version,
        renderer_version=renderer_version,
        declared_outputs=tuple(_REASON_OUTPUT_KIND[reason] for reason in reasons),
        attempted_actions=(),
    )


def evaluate_accounting_report(
    *,
    task_id: str,
    period: ReportPeriod,
    rows: Sequence[NormalizedLedgerRow],
    analysis: ReportAnalysis,
    source_receipts: Sequence[AccountingSourceReceipt],
    accounting_rules_version: str,
    prohibited_actions: frozenset[str],
    renderer_version: str,
) -> AccountingEvaluation:
    """Evaluate normalized accounting facts without trusting model narrative."""

    if prohibited_actions:
        return _terminal_evaluation(
            task_id=task_id,
            period=period,
            status=WorkProductStatus.BLOCKED,
            reasons=("FORBIDDEN_EXTERNAL_ACTION",),
            source_receipts=source_receipts,
            accounting_rules_version=accounting_rules_version,
            renderer_version=renderer_version,
        )

    data_reasons: list[str] = []
    if (
        not task_id.strip()
        or not rows
        or not analysis.numbers
        or not source_receipts
        or not renderer_version.strip()
    ):
        data_reasons.append("MISSING_REQUIRED_FIELD")
    if any(
        number.amount is not None
        and (
            not isinstance(number.amount, Decimal)
            or not number.amount.is_finite()
        )
        for number in analysis.numbers
    ):
        data_reasons.append("INVALID_NUMERIC_VALUE")
    if not accounting_rules_version.strip():
        data_reasons.append("MISSING_RULE_VERSION")
    if data_reasons:
        reasons = tuple(reason for reason in _REASON_ORDER if reason in data_reasons)
        return _terminal_evaluation(
            task_id=task_id,
            period=period,
            status=WorkProductStatus.NEEDS_DATA,
            reasons=reasons,
            source_receipts=source_receipts,
            accounting_rules_version=accounting_rules_version,
            renderer_version=renderer_version,
        )

    receipts_by_ref: dict[str, list[AccountingSourceReceipt]] = {}
    for receipt in source_receipts:
        receipts_by_ref.setdefault(receipt.source_ref, []).append(receipt)
    selected_receipt_by_ref: dict[str, AccountingSourceReceipt] = {}
    reasons: set[str] = set()

    fact_ids = [
        number.fact_id for number in analysis.numbers if number.fact_id is not None
    ]
    if len(fact_ids) != len(set(fact_ids)):
        reasons.add("DUPLICATE_ENTRY")

    source_identities = [
        (receipt.source_ref, receipt.source_digest) for receipt in source_receipts
    ]
    if len(source_identities) != len(set(source_identities)):
        reasons.add("DUPLICATE_SOURCE")

    if any(
        row.year < period.start_year or row.year > period.end_year for row in rows
    ):
        reasons.add("CROSS_PERIOD_ENTRY")

    if any(receipt.period != period for receipt in source_receipts):
        reasons.add("STALE_SOURCE")

    digests_by_logical_source: dict[str, set[str]] = {}
    for receipt in source_receipts:
        if receipt.adopted:
            assert receipt.logical_source_id is not None
            digests_by_logical_source.setdefault(
                receipt.logical_source_id, set()
            ).add(receipt.source_digest)
    if any(
        len(digests) > 1
        for digests in digests_by_logical_source.values()
    ):
        reasons.add("SOURCE_CONFLICT")

    valid_sources: set[str] = set()
    for row in rows:
        ref = _source_ref(row)
        exact_receipts = sorted(
            (
                candidate
                for candidate in receipts_by_ref.get(ref, ())
                if candidate.source_digest == row.source.file_sha256
            ),
            key=_receipt_selection_key,
        )
        receipt = exact_receipts[0] if exact_receipts else None
        if receipt is None:
            reasons.add("UNBACKED_REPORT_NUMBER")
            continue
        selected_receipt_by_ref[ref] = receipt
        if receipt.source_label is None:
            reasons.add("MISSING_SOURCE_LABEL")
            continue
        if not receipt.adopted:
            reasons.add("UNADOPTED_SOURCE")
            continue
        valid_sources.add(ref)

    valid_receipts = tuple(
        selected_receipt_by_ref[ref] for ref in sorted(valid_sources)
    )
    if any(
        receipt.conversion_receipt is not None
        and receipt.conversion_receipt.adopted
        and receipt.conversion_receipt.period != period
        for receipt in valid_receipts
    ):
        reasons.add("STALE_SOURCE")
    if not _currencies_are_reconciled(valid_receipts, period=period):
        reasons.add("CURRENCY_MISMATCH")

    assets = _family_total(rows, frozenset({"assets"}))
    liabilities = _family_total(rows, frozenset({"liabilities"}))
    equity = _family_total(rows, frozenset({"equity"}))
    if assets != liabilities + equity:
        reasons.add("ACCOUNTING_IDENTITY_MISMATCH")

    revenue = _family_total(rows, frozenset({"revenue"}))
    costs = _family_total(rows, frozenset({"cost"}))
    expenses = _family_total(rows, frozenset({"expense"}))
    profit = revenue - costs - expenses
    margin_numbers = tuple(
        (index, number)
        for index, number in enumerate(analysis.numbers)
        if number.name == "profit_margin"
    )
    invalid_margin_indexes: set[int] = set()
    if revenue == 0:
        for index, margin_number in margin_numbers:
            if margin_number.amount is None:
                reasons.add("ZERO_DENOMINATOR_EXPLAINED")
            else:
                invalid_margin_indexes.add(index)
                reasons.add("IMPOSSIBLE_RATIO")
    else:
        expected_margin = profit / revenue
        for index, margin_number in margin_numbers:
            if (
                not isinstance(margin_number.amount, Decimal)
                or abs(margin_number.amount) > Decimal("1")
                or margin_number.amount != expected_margin
            ):
                invalid_margin_indexes.add(index)
                reasons.add("IMPOSSIBLE_RATIO")

    rows_by_ref = {_source_ref(row): row for row in rows}
    row_amounts = {ref: _row_amount(row) for ref, row in rows_by_ref.items()}
    facts: list[AccountingFact] = []
    for index, number in enumerate(analysis.numbers):
        if index in invalid_margin_indexes:
            continue
        if number.name == "profit_margin" and number.amount is None:
            continue
        if (
            number.source_ref in row_amounts
            and number.source_ref not in valid_sources
        ):
            # A known row with an inadmissible receipt is already explained by
            # the source-specific reason and must not enter the fact pack.
            continue
        if not _is_backed_number(
            number,
            row_amounts=row_amounts,
            rows_by_ref=rows_by_ref,
            valid_sources=frozenset(valid_sources),
            accounting_rules_version=accounting_rules_version,
        ):
            reasons.add("UNBACKED_REPORT_NUMBER")
            continue
        assert isinstance(number.amount, Decimal)
        assert number.fact_id is not None
        assert number.source_ref is not None
        assert number.rule_id is not None
        facts.append(
            AccountingFact(
                fact_id=number.fact_id,
                name=number.name,
                amount=number.amount,
                source_ref=number.source_ref,
                source_digest=selected_receipt_by_ref[
                    number.source_ref
                ].source_digest,
                rule_id=number.rule_id,
                source_human_confirmed=selected_receipt_by_ref[
                    number.source_ref
                ].source_human_confirmed,
                conversion_receipt=selected_receipt_by_ref[
                    number.source_ref
                ].conversion_receipt,
            )
        )

    ordered_reasons = tuple(reason for reason in _REASON_ORDER if reason in reasons)
    status = (
        WorkProductStatus.NEEDS_REVIEW
        if reasons & _REVIEW_REASONS
        else WorkProductStatus.READY_FOR_HUMAN_CONFIRMATION
    )
    fact_payload = _fact_payload(facts)
    return AccountingEvaluation(
        task_id=task_id,
        period=period,
        work_status=status,
        reason_codes=ordered_reasons,
        facts=tuple(facts),
        fact_pack_digest=semantic_digest({"facts": fact_payload}),
        source_digests=tuple(
            sorted(
                [
                    receipt.source_digest
                    for receipt in source_receipts
                    if receipt.source_ref in valid_sources
                ]
                + [
                    receipt.conversion_receipt.source_digest
                    for receipt in valid_receipts
                    if receipt.conversion_receipt is not None
                    and receipt.conversion_receipt.adopted
                ]
            )
        ),
        accounting_rules_version=accounting_rules_version,
        renderer_version=renderer_version,
        declared_outputs=tuple(
            ["traceable_fact"] if facts else []
        )
        + tuple(_REASON_OUTPUT_KIND[reason] for reason in ordered_reasons),
        attempted_actions=(),
    )
