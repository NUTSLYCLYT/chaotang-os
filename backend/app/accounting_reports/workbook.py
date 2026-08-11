from __future__ import annotations

import hashlib
import os
import re
import tempfile
import zipfile
from collections.abc import Iterable
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Final

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .analysis import ACCOUNT_FAMILY_BY_CATEGORY, source_id_for_row
from .models import AccountingReportSummary, NormalizedLedgerRow

SHEET_NAMES: Final = (
    "管理摘要",
    "核心财务报表",
    "科目趋势",
    "异常分析",
    "科目明细",
    "校验结果",
    "数据来源",
)
MONEY_FORMAT: Final = '#,##0.00;[Red](#,##0.00);-'
PERCENT_FORMAT: Final = "0.00%"
_HEADER_FILL: Final = PatternFill("solid", fgColor="1F4E78")
_SHEET_REFERENCE = re.compile(r"'([^']+)'!")
_QUALIFIED_REFERENCE = re.compile(
    r"'([^']+)'!\$?([A-Z]{1,3})\$?(\d+)"
    r"(?::\$?([A-Z]{1,3})\$?(\d+))?"
)
_UNQUOTED_QUALIFIED_REFERENCE = re.compile(
    r"(?<!')([\w.]+)!\$?([A-Z]{1,3})\$?(\d+)"
    r"(?::\$?([A-Z]{1,3})\$?(\d+))?"
)
_LOCAL_REFERENCE = re.compile(
    r"(?<![A-Z0-9_!])\$?([A-Z]{1,3})\$?(\d+)"
    r"(?::\$?([A-Z]{1,3})\$?(\d+))?"
)
_FORBIDDEN_FORMULA_MARKERS: Final = (
    "[",
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
)
_DETERMINISTIC_XLSX_DATETIME: Final = (1980, 1, 1, 0, 0, 0)
_DETERMINISTIC_ZIP_DATETIME: Final = (1980, 1, 1, 0, 0, 0)
_DETERMINISTIC_CORE_TIMESTAMP: Final = b"1980-01-01T00:00:00Z"
_CORE_CREATED_TIMESTAMP = re.compile(
    rb"(<dcterms:created\b[^>]*>)[^<]*(</dcterms:created>)"
)
_CORE_MODIFIED_TIMESTAMP = re.compile(
    rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)"
)


class AccountingWorkbookError(RuntimeError):
    pass


def _decimal_formula(value: Decimal) -> str:
    fixed = format(value, "f")
    negative = fixed.startswith("-")
    unsigned = fixed[1:] if negative else fixed
    whole, separator, fraction = unsigned.partition(".")
    digits = (whole + fraction).lstrip("0") or "0"
    numerator = f"-{digits}" if negative and digits != "0" else digits
    if not separator:
        return f"={numerator}"
    return f"={numerator}/{10 ** len(fraction)}"


def _style_header(sheet: Worksheet, row: int = 1) -> None:
    for cell in sheet[row]:
        if cell.value is not None:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center")


def _finish_table(sheet: Worksheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, sheet.max_column).column_letter}{sheet.max_row}"
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
        sheet.column_dimensions[letter].width = max(width, 10)


def _write_management_summary(
    sheet: Worksheet,
    summary: AccountingReportSummary,
    source_hashes: tuple[str, ...],
) -> None:
    sheet.append(["管理报告", "值"])
    sheet.append(
        [
            "报告期间",
            f"{summary.period.start_year}-{summary.period.end_year}",
        ]
    )
    sheet.append(["报告基准年度", summary.period.end_year])
    sheet.append(["来源哈希", ", ".join(source_hashes)])
    sheet.append(["总体校验", summary.overall_status])
    last_core_row = summary.period.end_year - summary.period.start_year + 2
    sheet.append(["末期利润", f"='核心财务报表'!H{last_core_row}"])
    sheet["B6"].number_format = MONEY_FORMAT
    _style_header(sheet)
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 72
    sheet.freeze_panes = "A2"


def _sumifs(
    amount_column: str,
    year_cell: str,
    category: str,
    detail_last_row: int,
) -> str:
    return (
        f"SUMIFS('科目明细'!${amount_column}$2:${amount_column}${detail_last_row},"
        f"'科目明细'!$A$2:$A${detail_last_row},{year_cell},"
        f"'科目明细'!$N$2:$N${detail_last_row},\"{category}\")"
    )


def _write_core(
    sheet: Worksheet,
    summary: AccountingReportSummary,
    detail_last_row: int,
) -> None:
    sheet.append(
        [
            "年度",
            "资产",
            "负债",
            "所有者权益",
            "收入",
            "成本",
            "费用",
            "利润",
            "利润同比",
            "利润同比额",
        ]
    )
    for index, year in enumerate(
        range(summary.period.start_year, summary.period.end_year + 1), start=2
    ):
        sheet.cell(index, 1, year).number_format = "0"
        year_metrics = summary.metrics_by_year[year]
        if year_metrics["assets"] is None:
            continue
        for column, category, debit_column, credit_column, credit_oriented in (
            (2, "assets", "L", "M", False),
            (3, "liabilities", "L", "M", True),
            (4, "equity", "L", "M", True),
            (5, "revenue", "G", "H", True),
            (6, "cost", "G", "H", False),
            (7, "expense", "G", "H", False),
        ):
            debit = _sumifs(debit_column, f"$A{index}", category, detail_last_row)
            credit = _sumifs(credit_column, f"$A{index}", category, detail_last_row)
            formula = f"={credit}-{debit}" if credit_oriented else f"={debit}-{credit}"
            sheet.cell(index, column, formula).number_format = MONEY_FORMAT
        sheet.cell(index, 8, f"=E{index}-F{index}-G{index}").number_format = MONEY_FORMAT
        prior_metrics = summary.metrics_by_year.get(year - 1)
        if prior_metrics is None or prior_metrics["profit"] is None:
            sheet.cell(index, 9, "")
            sheet.cell(index, 10, "")
        else:
            sheet.cell(
                index, 9, f'=IFERROR((H{index}-H{index - 1})/ABS(H{index - 1}),"")'
            ).number_format = PERCENT_FORMAT
            sheet.cell(index, 10, f"=H{index}-H{index - 1}").number_format = MONEY_FORMAT
    _style_header(sheet)
    _finish_table(sheet)


def _write_trends(sheet: Worksheet, rows: tuple[NormalizedLedgerRow, ...]) -> None:
    sheet.append(["年度", "科目编码", "科目名称", "类别", "期末净额"])
    for row_index, row in enumerate(rows, start=2):
        sheet.append(
            [
                row.year,
                row.account_code,
                row.account_name,
                row.category,
                f"='科目明细'!L{row_index}-'科目明细'!M{row_index}",
            ]
        )
        sheet.cell(row_index, 1).number_format = "0"
        sheet.cell(row_index, 2).number_format = "@"
        sheet.cell(row_index, 5).number_format = MONEY_FORMAT
    _style_header(sheet)
    _finish_table(sheet)


def _write_exceptions(sheet: Worksheet, exceptions: Iterable[str]) -> None:
    sheet.append(["序号", "异常说明"])
    values = tuple(exceptions)
    if values:
        for index, item in enumerate(values, start=1):
            sheet.append([index, item])
    else:
        sheet.append([1, "无异常"])
    _style_header(sheet)
    _finish_table(sheet)


def _write_details(sheet: Worksheet, rows: tuple[NormalizedLedgerRow, ...]) -> None:
    sheet.append(
        [
            "年度",
            "类别",
            "科目编码",
            "科目名称",
            "期初借方",
            "期初贷方",
            "本期借方",
            "本期贷方",
            "来源文件",
            "来源工作表",
            "来源行",
            "期末借方",
            "期末贷方",
            "规范类别",
        ]
    )
    for row_index, row in enumerate(rows, start=2):
        sheet.append(
            [
                row.year,
                row.category,
                row.account_code,
                row.account_name,
                _decimal_formula(row.opening_debit),
                _decimal_formula(row.opening_credit),
                _decimal_formula(row.movement_debit),
                _decimal_formula(row.movement_credit),
                Path(row.source.file_name).name,
                row.source.sheet_name,
                row.source.row_number,
                _decimal_formula(row.closing_debit),
                _decimal_formula(row.closing_credit),
                ACCOUNT_FAMILY_BY_CATEGORY.get(row.category.strip().lower(), ""),
            ]
        )
        sheet.cell(row_index, 1).number_format = "0"
        sheet.cell(row_index, 3).number_format = "@"
        for column in (*range(5, 9), 12, 13):
            sheet.cell(row_index, column).number_format = MONEY_FORMAT
    _style_header(sheet)
    _finish_table(sheet)


def _write_checks(sheet: Worksheet, summary: AccountingReportSummary) -> None:
    sheet.append(["校验项", "状态", "实际值", "期望值", "差异", "说明"])
    for check in summary.checks:
        sheet.append(
            [
                check.name,
                check.status,
                check.actual,
                check.expected,
                check.difference,
                check.detail,
            ]
        )
    _style_header(sheet)
    _finish_table(sheet)


def _write_sources(
    sheet: Worksheet,
    summary: AccountingReportSummary,
    rows: tuple[NormalizedLedgerRow, ...],
) -> None:
    sheet.append(["文件名", "工作表", "来源标识", "来源哈希", "来源行"])
    seen: set[tuple[str, str, int, str]] = set()
    for row in rows:
        identity = (
            row.source.file_sha256,
            row.source.sheet_name,
            row.source.row_number,
            source_id_for_row(row),
        )
        if identity in seen:
            continue
        seen.add(identity)
        sheet.append(
            [
                Path(row.source.file_name).name,
                row.source.sheet_name,
                source_id_for_row(row),
                row.source.file_sha256,
                row.source.row_number,
            ]
        )
    sheet.append(
        [
            "报告期间",
            f"{summary.period.start_year}-{summary.period.end_year}",
        ]
    )
    for source_hash in sorted({row.source.file_sha256 for row in rows}):
        sheet.append(["来源哈希", source_hash])
    _style_header(sheet)
    _finish_table(sheet)


def _build_workbook(
    rows: tuple[NormalizedLedgerRow, ...],
    summary: AccountingReportSummary,
) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in SHEET_NAMES:
        workbook.create_sheet(name)
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    metadata_timestamp = datetime(*_DETERMINISTIC_XLSX_DATETIME)
    workbook.properties.created = metadata_timestamp
    workbook.properties.modified = metadata_timestamp
    source_hashes = tuple(sorted({row.source.file_sha256 for row in rows}))
    _write_management_summary(workbook["管理摘要"], summary, source_hashes)
    _write_details(workbook["科目明细"], rows)
    _write_core(workbook["核心财务报表"], summary, len(rows) + 1)
    _write_trends(workbook["科目趋势"], rows)
    _write_exceptions(workbook["异常分析"], summary.exceptions)
    _write_checks(workbook["校验结果"], summary)
    _write_sources(workbook["数据来源"], summary, rows)
    return workbook


def _normalize_xlsx_archive(path: Path) -> None:
    descriptor, normalized_name = tempfile.mkstemp(
        prefix=f".{path.name}.",
        suffix=".normalized.xlsx",
        dir=path.parent,
    )
    os.close(descriptor)
    normalized_path = Path(normalized_name)
    try:
        with zipfile.ZipFile(path, "r") as source:
            members = []
            for info in sorted(source.infolist(), key=lambda item: item.filename):
                content = source.read(info.filename)
                if info.filename == "docProps/core.xml":
                    for pattern in (
                        _CORE_CREATED_TIMESTAMP,
                        _CORE_MODIFIED_TIMESTAMP,
                    ):
                        content = pattern.sub(
                            lambda match: (
                                match.group(1)
                                + _DETERMINISTIC_CORE_TIMESTAMP
                                + match.group(2)
                            ),
                            content,
                        )
                members.append((info.filename, content))
        with zipfile.ZipFile(
            normalized_path,
            "w",
            compression=zipfile.ZIP_DEFLATED,
            compresslevel=9,
        ) as destination:
            for filename, content in members:
                info = zipfile.ZipInfo(filename, _DETERMINISTIC_ZIP_DATETIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 0
                info.external_attr = 0
                destination.writestr(info, content, compress_type=zipfile.ZIP_DEFLATED)
        os.replace(normalized_path, path)
    finally:
        normalized_path.unlink(missing_ok=True)


def _audit_generated_workbook(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ValueError("invalid_workbook_archive")
    workbook = load_workbook(path, data_only=False)
    try:
        if tuple(workbook.sheetnames) != SHEET_NAMES:
            raise ValueError("invalid_workbook_sheets")
        for sheet_name in SHEET_NAMES:
            sheet = workbook[sheet_name]
            if sheet.max_row < 2 or sheet.max_column < 2:
                raise ValueError("missing_required_range")
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type != "f":
                        continue
                    formula = str(cell.value)
                    if any(marker in formula for marker in _FORBIDDEN_FORMULA_MARKERS):
                        raise ValueError("unsafe_formula")
                    references = _SHEET_REFERENCE.findall(formula)
                    if any(reference not in SHEET_NAMES for reference in references):
                        raise ValueError("unknown_formula_reference")
                    qualified = tuple(_QUALIFIED_REFERENCE.finditer(formula))
                    if len(qualified) != len(references):
                        raise ValueError("invalid_formula_range")
                    for match in qualified:
                        target = workbook[match.group(1)]
                        _validate_reference_bounds(
                            target,
                            match.group(2),
                            match.group(3),
                            match.group(4),
                            match.group(5),
                        )
                    unquoted = tuple(
                        _UNQUOTED_QUALIFIED_REFERENCE.finditer(formula)
                    )
                    for match in unquoted:
                        if match.group(1) not in SHEET_NAMES:
                            raise ValueError("unknown_formula_reference")
                        target = workbook[match.group(1)]
                        _validate_reference_bounds(
                            target,
                            match.group(2),
                            match.group(3),
                            match.group(4),
                            match.group(5),
                        )
                    local_formula = _QUALIFIED_REFERENCE.sub("", formula)
                    local_formula = _UNQUOTED_QUALIFIED_REFERENCE.sub(
                        "", local_formula
                    )
                    for match in _LOCAL_REFERENCE.finditer(local_formula):
                        _validate_reference_bounds(
                            sheet,
                            match.group(1),
                            match.group(2),
                            match.group(3),
                            match.group(4),
                        )
        checks = workbook["校验结果"]
        names = {
            checks.cell(row, 1).value for row in range(2, checks.max_row + 1)
        }
        required = {
            "coverage",
            "mapping",
            "balance_sheet_equation",
            "opening_continuity",
            "movement_balance",
            "account_directions",
        }
        if not required <= names:
            raise ValueError("missing_required_checks")
    finally:
        workbook.close()


def _validate_reference_bounds(
    sheet: Worksheet,
    start_column: str,
    start_row: str,
    end_column: str | None,
    end_row: str | None,
) -> None:
    from openpyxl.utils.cell import column_index_from_string

    coordinates = (
        (start_column, start_row),
        (end_column, end_row),
    )
    for column, row in coordinates:
        if column is None or row is None:
            continue
        if (
            column_index_from_string(column) > sheet.max_column
            or int(row) > sheet.max_row
        ):
            raise ValueError("invalid_formula_range")


def write_management_report(
    rows: tuple[NormalizedLedgerRow, ...],
    summary: AccountingReportSummary,
    destination: Path,
) -> str:
    path = Path(destination)
    temporary_path: Path | None = None
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary_name = tempfile.mkstemp(
            prefix=f".{path.name}.",
            suffix=".xlsx",
            dir=path.parent,
        )
        os.close(descriptor)
        temporary_path = Path(temporary_name)
        workbook = _build_workbook(tuple(rows), summary)
        try:
            workbook.save(temporary_path)
        finally:
            workbook.close()
        _normalize_xlsx_archive(temporary_path)
        _audit_generated_workbook(temporary_path)
        digest = hashlib.sha256(temporary_path.read_bytes()).hexdigest()
        os.replace(temporary_path, path)
        temporary_path = None
        return digest
    except Exception:
        try:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)
        except OSError:
            pass
        raise AccountingWorkbookError("workbook_generation_failed") from None
