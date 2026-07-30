from __future__ import annotations

import hashlib
import re
from collections.abc import Iterable, Sequence
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Final

import xlrd
from openpyxl import load_workbook

from .models import NormalizedLedgerRow, ReportPeriod, SourceRef

_HEADERS: Final = (
    "科目类别",
    "科目编码",
    "科目名称",
    "期初借方",
    "期初贷方",
    "本期借方",
    "本期贷方",
    "期末借方",
    "期末贷方",
)
_SUPPORTED_SUFFIXES: Final = frozenset({".xlsx", ".xls"})
_YEAR_PATTERN: Final = re.compile(r"(?<!\d)(20\d{2}|2100)(?!\d)")
_ERROR_CODES: Final = frozenset(
    {
        "source_missing",
        "source_path_invalid",
        "source_format_unsupported",
        "source_schema_invalid",
        "source_period_conflict",
    }
)


class AccountingSourceError(ValueError):
    def __init__(self, code: str, file_name: str | None = None) -> None:
        if code not in _ERROR_CODES:
            raise ValueError("unsupported accounting source error code")
        self.code = code
        message = code if file_name is None else f"{code}: {Path(file_name).name}"
        super().__init__(message)


def _safe_resolve(path: Path, root: Path) -> Path:
    try:
        resolved = path.resolve(strict=True)
        resolved.relative_to(root)
    except (OSError, ValueError):
        raise AccountingSourceError("source_path_invalid", path.name) from None
    if not resolved.is_file():
        raise AccountingSourceError("source_path_invalid", path.name)
    return resolved


def _file_year(path: Path) -> int:
    matches = _YEAR_PATTERN.findall(path.stem)
    if len(matches) != 1:
        raise AccountingSourceError("source_schema_invalid", path.name)
    return int(matches[0])


def _discover_period_files(root: Path, period: ReportPeriod) -> tuple[tuple[Path, int], ...]:
    try:
        entries = sorted(root.iterdir(), key=lambda item: item.name)
        files = [entry for entry in entries if entry.is_file() or entry.is_symlink()]
    except OSError:
        raise AccountingSourceError("source_path_invalid", root.name) from None
    if not files:
        raise AccountingSourceError("source_missing")

    supported: list[tuple[Path, int]] = []
    for entry in files:
        if entry.suffix.lower() not in _SUPPORTED_SUFFIXES:
            raise AccountingSourceError("source_format_unsupported", entry.name)
        resolved = _safe_resolve(entry, root)
        supported.append((resolved, _file_year(entry)))

    requested_years = set(range(period.start_year, period.end_year + 1))
    candidates = [(path, year) for path, year in supported if year in requested_years]
    if not candidates or {year for _, year in candidates} != requested_years:
        conflict = supported[0][0].name if supported else None
        raise AccountingSourceError("source_period_conflict", conflict)

    seen: set[tuple[int, str]] = set()
    for path, year in candidates:
        identity = (year, path.suffix.lower())
        if identity in seen:
            raise AccountingSourceError("source_period_conflict", path.name)
        seen.add(identity)
    return tuple(candidates)


def _read_xlsx_sheets(path: Path) -> Iterable[tuple[str, Iterable[Sequence[object]]]]:
    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception:
        raise AccountingSourceError("source_schema_invalid", path.name) from None
    try:
        for sheet in workbook.worksheets:
            yield sheet.title, sheet.iter_rows(values_only=True)
    finally:
        workbook.close()


def _read_xls_sheets(path: Path) -> Iterable[tuple[str, Iterable[Sequence[object]]]]:
    try:
        workbook = xlrd.open_workbook(path, on_demand=True)
    except Exception:
        raise AccountingSourceError("source_schema_invalid", path.name) from None
    try:
        for sheet in workbook.sheets():
            yield sheet.name, (
                tuple(
                    sheet.cell_value(row_index, column_index)
                    for column_index in range(sheet.ncols)
                )
                for row_index in range(sheet.nrows)
            )
    finally:
        workbook.release_resources()


def _as_text(value: object, path: Path) -> str:
    if value is None or isinstance(value, bool):
        raise AccountingSourceError("source_schema_invalid", path.name)
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if not text or text.startswith("="):
        raise AccountingSourceError("source_schema_invalid", path.name)
    return text


def _as_decimal(value: object, path: Path) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, bool):
        raise AccountingSourceError("source_schema_invalid", path.name)
    text = str(value).strip().replace(",", "")
    if not text:
        return Decimal("0")
    if text.startswith("="):
        raise AccountingSourceError("source_schema_invalid", path.name)
    try:
        amount = Decimal(text)
    except (InvalidOperation, ValueError):
        raise AccountingSourceError("source_schema_invalid", path.name) from None
    if not amount.is_finite():
        raise AccountingSourceError("source_schema_invalid", path.name)
    return amount


def _hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    try:
        with path.open("rb") as source:
            for chunk in iter(lambda: source.read(1024 * 1024), b""):
                digest.update(chunk)
    except OSError:
        raise AccountingSourceError("source_path_invalid", path.name) from None
    return digest.hexdigest()


def _load_candidate(
    candidate: tuple[Path, int],
    approved_root: Path,
) -> tuple[NormalizedLedgerRow, ...]:
    path, year = candidate
    path = _safe_resolve(path, approved_root)
    file_hash = _hash_file(path)
    reader = _read_xlsx_sheets if path.suffix.lower() == ".xlsx" else _read_xls_sheets
    normalized: list[NormalizedLedgerRow] = []

    try:
        for sheet_name, sheet_rows in reader(path):
            rows = iter(sheet_rows)
            try:
                header = next(rows)
            except StopIteration:
                raise AccountingSourceError("source_schema_invalid", path.name) from None
            if tuple(
                str(cell).strip() if cell is not None else "" for cell in header
            ) != _HEADERS:
                raise AccountingSourceError("source_schema_invalid", path.name)
            for row_number, values in enumerate(rows, start=2):
                if not any(value not in (None, "") for value in values):
                    continue
                if len(values) != len(_HEADERS):
                    raise AccountingSourceError("source_schema_invalid", path.name)
                normalized.append(
                    NormalizedLedgerRow(
                        year=year,
                        category=_as_text(values[0], path),
                        account_code=_as_text(values[1], path),
                        account_name=_as_text(values[2], path),
                        opening_debit=_as_decimal(values[3], path),
                        opening_credit=_as_decimal(values[4], path),
                        movement_debit=_as_decimal(values[5], path),
                        movement_credit=_as_decimal(values[6], path),
                        closing_debit=_as_decimal(values[7], path),
                        closing_credit=_as_decimal(values[8], path),
                        source=SourceRef(
                            file_name=path.name,
                            sheet_name=sheet_name,
                            row_number=row_number,
                            file_sha256=file_hash,
                        ),
                    )
                )
    except AccountingSourceError:
        raise
    except Exception:
        raise AccountingSourceError("source_schema_invalid", path.name) from None
    if not normalized:
        raise AccountingSourceError("source_schema_invalid", path.name)
    return tuple(normalized)


def _validate_period_coverage(
    rows: tuple[NormalizedLedgerRow, ...],
    period: ReportPeriod,
) -> None:
    expected = set(range(period.start_year, period.end_year + 1))
    if {row.year for row in rows} != expected:
        raise AccountingSourceError("source_period_conflict")


def load_ledger_rows(
    source_dir: Path,
    period: ReportPeriod,
) -> tuple[NormalizedLedgerRow, ...]:
    try:
        approved_root = source_dir.resolve(strict=True)
    except OSError:
        raise AccountingSourceError("source_missing") from None
    if not approved_root.is_dir():
        raise AccountingSourceError("source_path_invalid", source_dir.name)
    candidates = _discover_period_files(approved_root, period)
    rows = tuple(
        row
        for candidate in candidates
        for row in _load_candidate(candidate, approved_root)
    )
    _validate_period_coverage(rows, period)
    return rows
