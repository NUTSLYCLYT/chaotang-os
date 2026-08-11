from __future__ import annotations

import io
import posixpath
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Final
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from .models import CellProbe, CellRegion, ReportPeriod, SheetProbe, WorkbookProbe
from .source_adapters import MAX_MERGED_RANGES_PER_SHEET, _read_xls_probe_sheets
from .source_manifest import freeze_candidate_workbooks, resolve_accounting_source_dir_at
from .sources import AccountingSourceError

MAX_WORKBOOK_SHEETS: Final = 32
MAX_SHEET_ROWS: Final = 20_000
MAX_SHEET_COLUMNS: Final = 256
MAX_EXPANDED_WORKBOOK_BYTES: Final = 64 * 1024 * 1024
_YEAR_TOKEN: Final = re.compile(r"(?<!\d)(20\d{2})(?!\d)")
_CELL_REF = re.compile(r"([A-Z]+)([1-9][0-9]*)")
_SPREADSHEETML = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_DIMENSION_TAG = f"{{{_SPREADSHEETML}}}dimension"
_MERGE_CELL_TAG = f"{{{_SPREADSHEETML}}}mergeCell"
_XlsxSheet = tuple[
    str,
    tuple[tuple[object, ...], ...],
    tuple[tuple[str, ...], ...],
    tuple[str, ...],
]


@dataclass(frozen=True, slots=True)
class _SheetContainerInfo:
    merged_ranges: tuple[str, ...]


def _meaningful(value: object) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _cell_type(value: object, data_type: str | None = None) -> str:
    if data_type == "f" or isinstance(value, str) and value.startswith("="):
        return "formula"
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, str):
        return "string"
    return "other"


def _probe_sheet(
    name: str,
    rows: tuple[tuple[object, ...], ...],
    *,
    cell_types: tuple[tuple[str, ...], ...] = (),
    merged_ranges: tuple[str, ...] = (),
) -> SheetProbe:
    if len(rows) > MAX_SHEET_ROWS or any(len(row) > MAX_SHEET_COLUMNS for row in rows):
        raise AccountingSourceError("source_schema_invalid")
    nonempty = [index for index, row in enumerate(rows, 1) if any(map(_meaningful, row))]
    if not nonempty:
        return SheetProbe(name, len(rows), max((len(row) for row in rows), default=0), ())
    width = max(len(row) for row in rows)
    spans: list[tuple[int, int]] = []
    span_start = previous = nonempty[0]
    for row_number in nonempty[1:]:
        if row_number != previous + 1:
            spans.append((span_start, previous))
            span_start = row_number
        previous = row_number
    spans.append((span_start, previous))
    regions: list[CellRegion] = []
    for start, end in spans:
        header_start = start
        first = rows[start - 1]
        if sum(_meaningful(value) for value in first) <= 1 and start < end:
            header_start += 1
        header_depth = 0
        for row in rows[header_start - 1 : end]:
            populated = [value for value in row if _meaningful(value)]
            if len(populated) < 2 or any(isinstance(value, (int, float)) for value in populated):
                break
            header_depth += 1
        cells = tuple(
            CellProbe(
                row_index,
                column_index,
                cell_types[row_index - 1][column_index - 1]
                if row_index <= len(cell_types) and column_index <= len(cell_types[row_index - 1])
                else _cell_type(value),
                value,
            )
            for row_index in range(start, end + 1)
            for column_index, value in enumerate(rows[row_index - 1], 1)
            if _meaningful(value)
        )
        formulas = tuple(
            (cell.row, cell.column, str(cell.value))
            for cell in cells
            if cell.value_type == "formula"
        )
        regions.append(CellRegion(start, end, 1, width, header_depth, cells, formulas))
    return SheetProbe(
        name=name,
        row_count=len(rows),
        column_count=width,
        regions=tuple(regions),
        merged_ranges=merged_ranges,
    )


def _column_number(token: bytes) -> int:
    value = 0
    for byte in token:
        value = value * 26 + byte - 64
    return value


def _dimension_bounds(reference: str) -> tuple[int, int, int, int]:
    parts = reference.strip().split(":")
    if len(parts) not in {1, 2}:
        raise AccountingSourceError("source_schema_invalid")
    matches = [_CELL_REF.fullmatch(part) for part in parts]
    if any(match is None for match in matches):
        raise AccountingSourceError("source_schema_invalid")
    start = matches[0]
    end = matches[-1]
    assert start is not None and end is not None
    start_column = _column_number(start.group(1).encode("ascii"))
    end_column = _column_number(end.group(1).encode("ascii"))
    start_row = int(start.group(2))
    end_row = int(end.group(2))
    if start_column > end_column or start_row > end_row:
        raise AccountingSourceError("source_schema_invalid")
    return start_row, end_row, start_column, end_column


def _merged_references(sheet_root: ElementTree.Element) -> tuple[str, ...]:
    references: list[str] = []
    for element in sheet_root.iter():
        if element.tag != _MERGE_CELL_TAG:
            continue
        if set(element.attrib) != {"ref"}:
            raise AccountingSourceError("source_schema_invalid")
        reference = element.attrib["ref"]
        _start_row, end_row, _start_column, end_column = _dimension_bounds(reference)
        if end_row > MAX_SHEET_ROWS or end_column > MAX_SHEET_COLUMNS:
            raise AccountingSourceError("source_schema_invalid")
        references.append(reference)
        if len(references) > MAX_MERGED_RANGES_PER_SHEET:
            raise AccountingSourceError("source_schema_invalid")
    return tuple(references)


def _preflight_xlsx(source: bytes) -> tuple[_SheetContainerInfo, ...]:
    try:
        with ZipFile(io.BytesIO(source)) as archive:
            infos = archive.infolist()
            names = {item.filename.casefold() for item in infos}
            if sum(item.file_size for item in infos) > MAX_EXPANDED_WORKBOOK_BYTES:
                raise AccountingSourceError("source_schema_invalid")
            try:
                content_types_root = ElementTree.fromstring(
                    archive.read("[Content_Types].xml")
                )
            except ElementTree.ParseError:
                raise AccountingSourceError("source_schema_invalid") from None
            content_type_values = tuple(
                str(value).strip().casefold()
                for element in content_types_root.iter()
                for name, value in element.attrib.items()
                if name.rsplit("}", 1)[-1].casefold() == "contenttype"
            )
            relationship_roots = []
            for item in archive.namelist():
                if not item.casefold().endswith(".rels"):
                    continue
                try:
                    relationship_roots.append(
                        ElementTree.fromstring(archive.read(item))
                    )
                except ElementTree.ParseError:
                    raise AccountingSourceError("source_schema_invalid") from None
            if (
                any(
                    "vbaproject.bin" in name
                    or name.startswith("xl/externallinks/")
                    for name in names
                )
                or any(
                    "macroenabled" in value or "vbaproject" in value
                    for value in content_type_values
                )
                or any(
                    str(attribute_value).strip().casefold() == "external"
                    for root in relationship_roots
                    for relationship in root
                    for attribute_name, attribute_value in relationship.attrib.items()
                    if attribute_name.rsplit("}", 1)[-1].casefold() == "targetmode"
                )
            ):
                raise AccountingSourceError("source_schema_invalid")
            workbook_xml = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            relationships_xml = ElementTree.fromstring(
                archive.read("xl/_rels/workbook.xml.rels")
            )
            relationship_targets: dict[str, str] = {}
            for relationship in relationships_xml:
                if not relationship.attrib.get("Type", "").endswith("/worksheet"):
                    continue
                relationship_id = relationship.attrib.get("Id")
                target = relationship.attrib.get("Target")
                if not relationship_id or not target:
                    raise AccountingSourceError("source_schema_invalid")
                normalized = (
                    posixpath.normpath(target).lstrip("/")
                    if target.startswith("/")
                    else posixpath.normpath(posixpath.join("xl", target))
                )
                if (
                    not normalized.startswith("xl/worksheets/")
                    or normalized not in archive.namelist()
                ):
                    raise AccountingSourceError("source_schema_invalid")
                if relationship_id in relationship_targets:
                    raise AccountingSourceError("source_schema_invalid")
                relationship_targets[relationship_id] = normalized
            relation_key = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            sheet_names = []
            for sheet in workbook_xml.findall(
                ".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet"
            ):
                relationship_id = sheet.attrib.get(relation_key)
                if relationship_id not in relationship_targets:
                    raise AccountingSourceError("source_schema_invalid")
                sheet_names.append(relationship_targets[relationship_id])
            if len(set(sheet_names)) != len(sheet_names):
                raise AccountingSourceError("source_schema_invalid")
            if len(sheet_names) > MAX_WORKBOOK_SHEETS:
                raise AccountingSourceError("source_schema_invalid")
            result = []
            for sheet_name in sheet_names:
                xml = archive.read(sheet_name)
                try:
                    sheet_root = ElementTree.fromstring(xml)
                except ElementTree.ParseError:
                    raise AccountingSourceError("source_schema_invalid") from None
                dimensions = tuple(
                    element
                    for element in sheet_root.iter()
                    if element.tag == _DIMENSION_TAG
                )
                if any(
                    element.tag.rsplit("}", 1)[-1] in {"dimension", "mergeCell"}
                    and element.tag not in {_DIMENSION_TAG, _MERGE_CELL_TAG}
                    for element in sheet_root.iter()
                ):
                    raise AccountingSourceError("source_schema_invalid")
                if len(dimensions) != 1 or set(dimensions[0].attrib) != {"ref"}:
                    raise AccountingSourceError("source_schema_invalid")
                _start_row, end_row, _start_column, end_column = _dimension_bounds(
                    dimensions[0].attrib["ref"]
                )
                if end_row > MAX_SHEET_ROWS or end_column > MAX_SHEET_COLUMNS:
                    raise AccountingSourceError("source_schema_invalid")
                result.append(
                    _SheetContainerInfo(
                        _merged_references(sheet_root)
                    )
                )
            return tuple(result)
    except AccountingSourceError:
        raise
    except (BadZipFile, KeyError, OSError, ValueError):
        raise AccountingSourceError("source_schema_invalid") from None


def _xlsx_sheets(
    source: bytes,
) -> tuple[_XlsxSheet, ...]:
    container = _preflight_xlsx(source)
    workbook = None
    try:
        workbook = load_workbook(
            io.BytesIO(source), read_only=True, data_only=False, keep_links=False
        )
        if len(workbook.worksheets) > MAX_WORKBOOK_SHEETS:
            raise AccountingSourceError("source_schema_invalid")
        result = []
        for index, sheet in enumerate(workbook.worksheets):
            rows: list[tuple[object, ...]] = []
            types: list[tuple[str, ...]] = []
            for row_count, row in enumerate(sheet.iter_rows(), 1):
                if row_count > MAX_SHEET_ROWS or len(row) > MAX_SHEET_COLUMNS:
                    raise AccountingSourceError("source_schema_invalid")
                rows.append(tuple(cell.value for cell in row))
                types.append(tuple(_cell_type(cell.value, cell.data_type) for cell in row))
            result.append((sheet.title, tuple(rows), tuple(types), container[index].merged_ranges))
        return tuple(result)
    except AccountingSourceError:
        raise
    except Exception:
        raise AccountingSourceError("source_schema_invalid") from None
    finally:
        if workbook is not None:
            workbook.close()


def probe_accounting_sources(source_dir: Path, period: ReportPeriod) -> tuple[WorkbookProbe, ...]:
    root = resolve_accounting_source_dir_at(source_dir)
    frozen = freeze_candidate_workbooks(root)
    if not frozen:
        raise AccountingSourceError("source_missing")
    probes: list[WorkbookProbe] = []
    for item in frozen:
        is_xlsx = item.path.suffix.casefold() == ".xlsx"
        sheets = _xlsx_sheets(item.content) if is_xlsx else _read_xls_probe_sheets(item.content)
        if len(sheets) > MAX_WORKBOOK_SHEETS:
            raise AccountingSourceError("source_schema_invalid")
        sheet_probes = tuple(
            _probe_sheet(name, rows, cell_types=types, merged_ranges=merged)
            for name, rows, types, merged in sheets
        ) if is_xlsx else tuple(
            _probe_sheet(
                name,
                rows,
                cell_types=types,
                merged_ranges=tuple(
                    f"R{row_start}C{column_start}:R{row_end}C{column_end}"
                    for row_start, row_end, column_start, column_end in merged
                ),
            )
            for name, rows, types, merged in sheets
        )
        tokens = {
            int(match.group(1))
            for sheet in sheets
            for rows in (sheet[1],)
            for row in rows
            for value in row
            if isinstance(value, str)
            for match in _YEAR_TOKEN.finditer(value)
            if period.start_year <= int(match.group(1)) <= period.end_year
        }
        probes.append(WorkbookProbe(item.sha256, tuple(sorted(tokens)), sheet_probes))
    return tuple(probes)
