"""Persist cash-safety evaluations through the existing report artifact ledger."""

from __future__ import annotations

import hashlib
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook

from app.accounting_reports.models import ReportPeriod
from app.accounting_reports.storage import ArtifactStorage
from app.work_products import (
    ArtifactGateReceipt,
    ArtifactGateStatus,
    ArtifactManifestItem,
    ArtifactState,
    ConfirmationStatus,
    WorkProductEnvelope,
    WorkProductStatus,
    semantic_digest,
)

from .calculation import CashSafetyEvaluation
from .contracts import HubuCashSafetyCaseInput

_CAPABILITY_ID = "analyze-cash-safety"


def persist_cash_safety_work_product(
    *,
    owner_user_id: str,
    payload: HubuCashSafetyCaseInput,
    evaluation: CashSafetyEvaluation,
    recommendations: tuple[str, ...],
    storage: ArtifactStorage,
) -> tuple[str, WorkProductEnvelope]:
    """Create, bind and publish one immutable owner-scoped cash-safety workbook."""
    run_id, reply_id = uuid4().hex, uuid4().hex
    temporary_path = storage.artifact_dir / f".{uuid4().hex}.xlsx"
    try:
        _write_workbook(temporary_path, payload, evaluation, recommendations)
        file_digest = _sha256(temporary_path)
        source_refs = evaluation.evidence_links or ("attestation:missing",)
        source_hashes = tuple(hashlib.sha256(v.encode()).hexdigest() for v in source_refs)
        year = int(payload.period[:4])
        pending = storage.create_pending(
            owner_user_id=owner_user_id,
            run_id=run_id,
            report_type="cash-safety",
            display_name=f"{payload.period}现金安全评估.xlsx",
            period=ReportPeriod(year, year),
            source_hashes=source_hashes,
            file_sha256=file_digest,
            pending_path=temporary_path,
        )
        envelope = _build_envelope(
            owner_user_id,
            run_id,
            reply_id,
            file_digest,
            payload,
            evaluation,
            recommendations,
        )
        storage.create_work_product(owner_user_id, pending.artifact_id, envelope)
        storage.publish_run(owner_user_id, run_id, reply_id)
        return pending.artifact_id, envelope.model_copy(
            update={"artifact_state": ArtifactState.PUBLISHED}
        )
    except Exception:
        try:
            storage.abort_run(owner_user_id, run_id)
        except Exception:
            pass
        raise
    finally:
        temporary_path.unlink(missing_ok=True)


def _build_envelope(
    owner_user_id: str,
    run_id: str,
    reply_id: str,
    file_digest: str,
    payload: HubuCashSafetyCaseInput,
    evaluation: CashSafetyEvaluation,
    recommendations: tuple[str, ...],
) -> WorkProductEnvelope:
    missing = evaluation.missing_evidence_codes
    status = (
        WorkProductStatus.NEEDS_DATA if missing else WorkProductStatus.READY_FOR_HUMAN_CONFIRMATION
    )
    facts = (
        {
            "period": payload.period,
            "as_of": payload.as_of,
            "currency": payload.currency,
            "risk_level": evaluation.risk_level.value,
            "runway_state": evaluation.runway_state.value,
            "available_cash": f"{evaluation.available_cash:.2f}",
            "monthly_net_burn": None
            if evaluation.monthly_net_burn is None
            else f"{evaluation.monthly_net_burn:.2f}",
            "runway_months": None
            if evaluation.runway_months is None
            else f"{evaluation.runway_months:.2f}",
            "forecast_ending_cash": f"{evaluation.forecast_ending_cash:.2f}",
            "rule_version": "hubu-cash-safety-v1",
        },
    )
    manifests = (
        ArtifactManifestItem(
            kind="cash_safety_workbook_xlsx",
            ref="cash-safety-assessment.xlsx",
            content_digest=file_digest,
            traceable=True,
        ),
        ArtifactManifestItem(
            kind="cash_safety_evidence_manifest",
            ref="cash-safety-evidence.json",
            content_digest=semantic_digest(
                {"evidence_links": evaluation.evidence_links, "missing": missing}
            ),
            traceable=True,
        ),
    )
    envelope = WorkProductEnvelope(
        work_product_id=uuid4().hex,
        version=1,
        owner_user_id=owner_user_id,
        run_id=run_id,
        reply_id=reply_id,
        capability_id=_CAPABILITY_ID,
        work_status=status,
        confirmation_status=ConfirmationStatus.PENDING,
        artifact_state=ArtifactState.PENDING,
        decision="现金安全评估已就绪，等待人工确认。"
        if not missing
        else "现金安全评估已保存，但必须补齐证据后重新创建。",
        facts=facts,
        assumptions=(),
        recommendations=recommendations,
        evidence_used=evaluation.evidence_links,
        missing_evidence=missing,
        conflicts=(),
        risk_register=evaluation.risk_reason_codes,
        artifact_manifest=manifests,
        artifact_gate=ArtifactGateReceipt(
            status=ArtifactGateStatus.PASSED, reason_codes=(), missing_kinds=(), unexpected_kinds=()
        ),
        content_digest="0" * 64,
        created_at=datetime.now(UTC),
    )
    return envelope.model_copy(
        update={"content_digest": semantic_digest(envelope.model_dump(mode="python"))}
    )


def _write_workbook(
    path: Path,
    payload: HubuCashSafetyCaseInput,
    evaluation: CashSafetyEvaluation,
    recommendations: tuple[str, ...],
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "现金安全摘要"
    summary.append(("字段", "值"))
    for key, value in (
        ("期间", payload.period),
        ("基准日", payload.as_of),
        ("币种", payload.currency),
        ("风险等级", evaluation.risk_level.value),
        ("runway 状态", evaluation.runway_state.value),
        ("可用资金", f"{evaluation.available_cash:.2f}"),
        ("月净消耗", evaluation.monthly_net_burn),
        ("runway 月数", evaluation.runway_months),
        ("第 31 天余额", f"{evaluation.forecast_ending_cash:.2f}"),
        ("规则版本", "hubu-cash-safety-v1"),
    ):
        summary.append((key, "" if value is None else str(value)))
    forecast = workbook.create_sheet("31天预测")
    forecast.append(("日期", "期末现金"))
    for point in evaluation.daily_forecast:
        forecast.append((point.forecast_date, f"{point.ending_cash:.2f}"))
    evidence = workbook.create_sheet("证据与检查")
    evidence.append(("类型", "代码或引用"))
    for kind, values in (
        ("证据", evaluation.evidence_links),
        ("缺证", evaluation.missing_evidence_codes),
        ("检查", evaluation.check_codes),
    ):
        for value in values:
            evidence.append((kind, value))
    advice = workbook.create_sheet("建议")
    advice.append(("序号", "建议"))
    for index, value in enumerate(recommendations, 1):
        advice.append((index, value))
    workbook.save(path)
    workbook.close()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
