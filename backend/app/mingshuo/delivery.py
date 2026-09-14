"""Deterministic, non-authorizing Mingshuo delivery work-product producer."""

from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from typing import Any
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook

from app.work_products import (
    ArtifactManifestItem,
    ArtifactState,
    ConfirmationStatus,
    WorkProductEnvelope,
    WorkProductStatus,
    evaluate_artifact_gate,
    semantic_digest,
)

CAPABILITY_ID = "mingshuo.first-delivery.work-product.v1"
PRODUCER_POLICY_VERSION = "mingshuo.delivery.producer.v1"
REPORT_TYPE = "MINGSHUO_SOLUTION_QUOTATION_DRAFT_V1"
DISPLAY_NAME = "mingshuo-solution-quotation-draft.xlsx"
SHEET_NAMES = ("封面与限制", "事实与证据", "方案草案", "报价草案", "缺失与风险")
REQUIRED_MANIFEST_KINDS = frozenset(
    {
        "confirmation_request",
        "mingshuo_delivery_binding",
        "mingshuo_delivery_projection",
        "mingshuo_fact_pack",
        "mingshuo_solution_quote_xlsx",
    }
)
_DIGEST = re.compile(r"sha256:[0-9a-f]{64}")
_MAX_ZIP_ENTRIES = 32
_MAX_ZIP_ITEM_BYTES = 4 * 1024 * 1024
_MAX_ZIP_TOTAL_BYTES = 16 * 1024 * 1024
_ALLOWED_PARTS = frozenset(
    {
        "[Content_Types].xml",
        "_rels/.rels",
        "docProps/app.xml",
        "docProps/core.xml",
        "xl/_rels/workbook.xml.rels",
        "xl/styles.xml",
        "xl/theme/theme1.xml",
        "xl/workbook.xml",
        *(f"xl/worksheets/sheet{index}.xml" for index in range(1, 6)),
    }
)
_FORBIDDEN_XML = (
    b"TargetMode=\"External\"",
    b"TargetMode='External'",
    b"<hyperlink",
    b"externalLink",
    b"oleObject",
    b"embeddedObject",
    b"connection",
    b"queryTable",
    b"vbaProject",
)
_FORMULA_XML = re.compile(br"<f(?:[\s>])")


class DeliveryValidationError(RuntimeError):
    """A delivery projection or workbook failed the closed producer contract."""


@dataclass(frozen=True)
class DeliveryArtifact:
    artifact_id: str
    work_product_id: str
    binding_digest: str
    binding_json: str
    cell_projection_digest: str
    workbook_bytes: bytes
    workbook_sha256: str
    envelope: WorkProductEnvelope


def _canonical_bytes(value: Any) -> bytes:
    return json.dumps(
        value,
        ensure_ascii=False,
        allow_nan=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")


def _sha256(value: bytes) -> str:
    return "sha256:" + hashlib.sha256(value).hexdigest()


def _hex_digest(value: Any) -> str:
    return hashlib.sha256(_canonical_bytes(value)).hexdigest()


def _opaque_id(domain: str, binding_digest: str) -> str:
    return hashlib.sha256(f"{domain}\0{binding_digest}".encode()).hexdigest()[:32]


def _safe_cell(value: Any) -> str:
    if isinstance(value, (dict, list, tuple)):
        text = _canonical_bytes(value).decode("utf-8")
    elif value is None:
        text = ""
    else:
        text = str(value)
    normalized = unicodedata.normalize("NFKC", text)
    visible = normalized.lstrip(
        "\x00\t\n\r\v\f\u0085\u00a0\u2000\u2001\u2002\u2003\u2004\u2005\u2006"
        "\u2007\u2008\u2009\u200a\u2028\u2029\u202f\u205f\u3000"
    )
    return "'" + normalized if visible.startswith(("=", "+", "-", "@")) else normalized


def _closed_projection(
    pack: dict[str, Any], evaluation: dict[str, Any], requirements_text: str
) -> dict[str, Any]:
    adopted = {
        item["id"]: item
        for item in pack["evidence"]
        if item["adoptionStatus"] == "ADOPTED"
    }
    facts = [
        {
            "id": item["id"],
            "kind": item["kind"],
            "subject": item["subject"],
            "value": item["value"],
            "evidenceIds": sorted(item["evidenceRefs"]),
        }
        for item in sorted(pack["facts"], key=lambda item: item["id"])
    ]
    claims = [
        {
            "id": item["id"],
            "text": item["text"],
            "evidenceIds": sorted(item["evidenceRefs"]),
            "approvalStatus": item["approvalStatus"],
        }
        for item in sorted(pack["claims"], key=lambda item: item["id"])
    ]
    evidence = [
        {
            "id": item["id"],
            "sourceClass": item["sourceClass"],
            "digest": item["digest"],
            "validUntil": item["validUntil"],
            "adoptionStatus": item["adoptionStatus"],
        }
        for item in sorted(pack["evidence"], key=lambda item: item["id"])
    ]
    missing = sorted(
        {
            *evaluation["holdReasons"],
            *evaluation["blockReasons"],
            *(f"UNBOUND_CLAIM:{item['id']}" for item in claims if not item["evidenceIds"]),
        }
    )
    risks = [
        "报价金额、MOQ、交期、质保、税费和贸易条款均待商务人工填写与批准。",
        "认证、市场适用性和性能结论不得从本草案推断。",
    ]
    return {
        "schemaVersion": "mingshuo.delivery-projection.v1",
        "requirements": requirements_text,
        "facts": facts,
        "claims": claims,
        "evidence": evidence,
        "missing": missing,
        "conflicts": [],
        "risks": risks,
        "adoptedEvidenceIds": sorted(adopted),
    }


def _cell_projection(projection: dict[str, Any], pack: dict[str, Any]) -> list[dict[str, Any]]:
    fact_rows = [
        ["事实", item["id"], item["kind"], item["subject"], item["value"], item["evidenceIds"]]
        for item in projection["facts"]
    ]
    claim_rows = [
        ["主张", item["id"], item["approvalStatus"], item["text"], "", item["evidenceIds"]]
        for item in projection["claims"]
    ]
    evidence_rows = [
        [
            "证据",
            item["id"],
            item["sourceClass"],
            item["digest"],
            item["validUntil"],
            item["adoptionStatus"],
        ]
        for item in projection["evidence"]
    ]
    return [
        {
            "name": SHEET_NAMES[0],
            "rows": [
                ["字段", "内容"],
                ["成果性质", "NON_BINDING_DRAFT"],
                ["审批状态", "COMMERCIAL_APPROVAL_REQUIRED"],
                ["项目", pack["project"]["name"]],
                ["需求", projection["requirements"]],
                ["限制", "仅供内部人工审阅；不得作为报价、认证或生产放行。"],
            ],
        },
        {
            "name": SHEET_NAMES[1],
            "rows": [
                ["类别", "ID", "类型/状态", "内容/摘要", "值/有效期", "证据/采用状态"],
                *fact_rows,
                *claim_rows,
                *evidence_rows,
            ],
        },
        {
            "name": SHEET_NAMES[2],
            "rows": [
                ["类别", "内容", "证据边界"],
                ["事实", "仅采用事实与证据页中已列项目", "逐项保留证据 ID"],
                ["假设", "未提供的信息不作事实推断", "需要人工补充"],
                ["建议", "基于现有资料形成中性方案草案", "不构成工程或商业批准"],
            ],
        },
        {
            "name": SHEET_NAMES[3],
            "rows": [
                ["状态", "NON_BINDING_DRAFT"],
                ["审批", "COMMERCIAL_APPROVAL_REQUIRED"],
                ["价格", ""],
                ["MOQ", ""],
                ["交期", ""],
                ["质保", ""],
                ["有效期", ""],
                ["提示", "待商务人工填写与批准"],
            ],
        },
        {
            "name": SHEET_NAMES[4],
            "rows": [
                ["类别", "内容"],
                *[["缺失", item] for item in projection["missing"]],
                *[["冲突", item] for item in projection["conflicts"]],
                *[["风险", item] for item in projection["risks"]],
            ],
        },
    ]


def _workbook_bytes(cell_projection: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_projection in cell_projection:
        sheet = workbook.create_sheet(sheet_projection["name"])
        sheet.sheet_state = "visible"
        for row in sheet_projection["rows"]:
            sheet.append([_safe_cell(value) for value in row])
    output = BytesIO()
    workbook.save(output)
    generated = output.getvalue()
    raw_output = BytesIO()
    with ZipFile(BytesIO(generated)) as source, ZipFile(
        raw_output, "w", ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o100644 << 16
            target.writestr(info, source.read(name))
    raw = raw_output.getvalue()
    validate_workbook_bytes(raw)
    return raw


def validate_workbook_bytes(raw: bytes) -> None:
    """Fail closed on unsafe or non-canonical generated OOXML surface."""

    try:
        with ZipFile(BytesIO(raw)) as archive:
            items = archive.infolist()
            names = [item.filename for item in items]
            if (
                len(items) > _MAX_ZIP_ENTRIES
                or len(names) != len(set(names))
                or set(names) != _ALLOWED_PARTS
                or any(
                    name.startswith(("/", "\\"))
                    or "\\" in name
                    or any(part in {"", ".", ".."} for part in name.split("/"))
                    for name in names
                )
                or any(item.file_size > _MAX_ZIP_ITEM_BYTES for item in items)
                or sum(item.file_size for item in items) > _MAX_ZIP_TOTAL_BYTES
            ):
                raise DeliveryValidationError("unsafe workbook")
            for item in items:
                body = archive.read(item)
                if _FORMULA_XML.search(body) or any(
                    marker in body for marker in _FORBIDDEN_XML
                ):
                    raise DeliveryValidationError("unsafe workbook")
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False)
        if tuple(workbook.sheetnames) != SHEET_NAMES or any(
            sheet.sheet_state != "visible" for sheet in workbook.worksheets
        ):
            raise DeliveryValidationError("unsafe workbook")
        if any(
            cell.data_type == "f"
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        ):
            raise DeliveryValidationError("unsafe workbook")
    except DeliveryValidationError:
        raise
    except (BadZipFile, KeyError, OSError, TypeError, ValueError):
        raise DeliveryValidationError("unsafe workbook") from None


def build_delivery_artifact(
    *,
    pack: dict[str, Any],
    evaluation: dict[str, Any],
    requirements_text: str,
    tenant_id: str,
    owner_user_id: str,
    project_id: str,
    draft_request_id: str,
    fact_pack_version: int,
    fact_pack_digest: str,
    created_at: datetime,
) -> DeliveryArtifact:
    if evaluation.get("decision") != "PASS" or not _DIGEST.fullmatch(fact_pack_digest):
        raise DeliveryValidationError("fact pack is not deliverable")
    for key in ("evidenceDigest", "factDigest", "claimDigest"):
        if not _DIGEST.fullmatch(str(evaluation.get(key, ""))):
            raise DeliveryValidationError("fact pack is not deliverable")
    binding = {
        "schemaVersion": "mingshuo.delivery-binding.v1",
        "tenantId": tenant_id,
        "ownerUserId": owner_user_id,
        "projectId": project_id,
        "draftRequestId": draft_request_id,
        "factPackVersion": fact_pack_version,
        "factPackDigest": fact_pack_digest,
        "evidenceDigest": evaluation["evidenceDigest"],
        "factDigest": evaluation["factDigest"],
        "claimDigest": evaluation["claimDigest"],
        "producerPolicyVersion": PRODUCER_POLICY_VERSION,
    }
    binding_digest = _sha256(_canonical_bytes(binding))
    projection = _closed_projection(pack, evaluation, requirements_text)
    cells = _cell_projection(projection, pack)
    cell_projection_digest = _sha256(_canonical_bytes(cells))
    workbook_bytes = _workbook_bytes(cells)
    workbook_sha256 = _sha256(workbook_bytes)
    artifact_id = _opaque_id("mingshuo-artifact-v1", binding_digest)
    work_product_id = _opaque_id("mingshuo-work-product-v1", binding_digest)
    manifest = (
        ArtifactManifestItem(
            kind="confirmation_request",
            ref="confirmation-request.json",
            content_digest=_hex_digest(
                {"confirmationStatus": "PENDING", "meaning": "INTERNAL_DRAFT_REVIEW_ONLY"}
            ),
            traceable=True,
        ),
        ArtifactManifestItem(
            kind="mingshuo_delivery_binding",
            ref="mingshuo-delivery-binding.json",
            content_digest=binding_digest.removeprefix("sha256:"),
            traceable=True,
        ),
        ArtifactManifestItem(
            kind="mingshuo_delivery_projection",
            ref="mingshuo-delivery-projection.json",
            content_digest=_hex_digest(projection),
            traceable=True,
        ),
        ArtifactManifestItem(
            kind="mingshuo_fact_pack",
            ref="mingshuo-project-fact-pack.json",
            content_digest=fact_pack_digest.removeprefix("sha256:"),
            traceable=True,
        ),
        ArtifactManifestItem(
            kind="mingshuo_solution_quote_xlsx",
            ref=DISPLAY_NAME,
            content_digest=workbook_sha256.removeprefix("sha256:"),
            traceable=True,
        ),
    )
    gate = evaluate_artifact_gate(required_kinds=REQUIRED_MANIFEST_KINDS, artifacts=manifest)
    public_binding = {
        "bindingDigest": binding_digest,
        "cellProjectionDigest": cell_projection_digest,
        "factPackVersion": fact_pack_version,
        "factPackDigest": fact_pack_digest,
        "evidenceDigest": evaluation["evidenceDigest"],
        "factDigest": evaluation["factDigest"],
        "claimDigest": evaluation["claimDigest"],
        "producerPolicyVersion": PRODUCER_POLICY_VERSION,
    }
    draft = WorkProductEnvelope(
        work_product_id=work_product_id,
        version=1,
        owner_user_id=owner_user_id,
        run_id=draft_request_id,
        reply_id=None,
        capability_id=CAPABILITY_ID,
        work_status=WorkProductStatus.READY_FOR_HUMAN_CONFIRMATION,
        confirmation_status=ConfirmationStatus.PENDING,
        artifact_state=ArtifactState.PENDING,
        decision="NON_BINDING_DRAFT",
        facts=(public_binding,),
        assumptions=("未提供的信息不作为事实推断。",),
        recommendations=("由工程、质量、合规与商务人员完成独立复核。",),
        evidence_used=tuple(sorted(projection["adoptedEvidenceIds"])),
        missing_evidence=tuple(projection["missing"]),
        conflicts=tuple(projection["conflicts"]),
        risk_register=tuple(projection["risks"]),
        artifact_manifest=manifest,
        artifact_gate=gate,
        content_digest="0" * 64,
        created_at=created_at,
    )
    envelope = draft.model_copy(update={"content_digest": semantic_digest(draft.model_dump())})
    return DeliveryArtifact(
        artifact_id=artifact_id,
        work_product_id=work_product_id,
        binding_digest=binding_digest,
        binding_json=_canonical_bytes(binding).decode("utf-8"),
        cell_projection_digest=cell_projection_digest,
        workbook_bytes=workbook_bytes,
        workbook_sha256=workbook_sha256,
        envelope=envelope,
    )


__all__ = [
    "CAPABILITY_ID",
    "DISPLAY_NAME",
    "DeliveryArtifact",
    "DeliveryValidationError",
    "PRODUCER_POLICY_VERSION",
    "REPORT_TYPE",
    "build_delivery_artifact",
    "validate_workbook_bytes",
]
