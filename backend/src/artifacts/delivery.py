"""Deterministic, in-memory R0-W06 artifact rendering."""

from __future__ import annotations

import hashlib
import io
import json
import textwrap
import zipfile
from dataclasses import dataclass
from typing import Any, Callable

from docx import Document


@dataclass(frozen=True)
class DeliveryArtifact:
    artifact_id: str
    kind: str
    mime_type: str
    content: bytes
    status: str
    content_hash: str
    byte_size: int


@dataclass(frozen=True)
class DeliveryResult:
    task_id: str
    final_memorial_id: str
    final_memorial_version: int
    artifacts: tuple[DeliveryArtifact, ...]
    overall_status: str


ArtifactRenderer = Callable[..., DeliveryArtifact]


def _render_lines(payload: dict[str, Any]) -> list[str]:
    source_lines = json.dumps(
        payload,
        ensure_ascii=False,
        indent=2,
        sort_keys=True,
    ).splitlines()
    lines = ["ContractReviewPackV1"]
    for line in source_lines:
        lines.extend(
            textwrap.wrap(
                line,
                width=88,
                break_long_words=True,
                break_on_hyphens=False,
                drop_whitespace=False,
                replace_whitespace=False,
            )
            or [""]
        )
    return lines


def _to_unicode_cmap(lines: list[str]) -> bytes:
    code_units = sorted(
        {
            unit
            for line in lines
            for unit in (
                int.from_bytes(line.encode("utf-16-be")[index : index + 2], "big")
                for index in range(0, len(line.encode("utf-16-be")), 2)
            )
        }
    )
    entries = [
        b"/CIDInit /ProcSet findresource begin",
        b"12 dict begin",
        b"begincmap",
        b"/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def",
        b"/CMapName /Adobe-Identity-UCS def",
        b"/CMapType 2 def",
        b"1 begincodespacerange",
        b"<0000> <FFFF>",
        b"endcodespacerange",
    ]
    for offset in range(0, len(code_units), 100):
        chunk = code_units[offset : offset + 100]
        entries.append(f"{len(chunk)} beginbfchar".encode())
        entries.extend(
            f"<{unit:04X}> <{unit:04X}>".encode()
            for unit in chunk
        )
        entries.append(b"endbfchar")
    entries.extend(
        [
            b"endcmap",
            b"CMapName currentdict /CMap defineresource pop",
            b"end",
            b"end",
        ]
    )
    return b"\n".join(entries)


def _pdf_page_stream(lines: list[str]) -> bytes:
    commands = [b"BT", b"/F1 8 Tf", b"48 760 Td", b"11 TL"]
    for index, line in enumerate(lines):
        if index:
            commands.append(b"T*")
        encoded = line.encode("utf-16-be").hex().upper()
        commands.append(f"<{encoded}> Tj".encode())
    commands.append(b"ET")
    return b"\n".join(commands)


def _pdf_bytes(lines: list[str]) -> bytes:
    pages = [lines[index : index + 62] for index in range(0, len(lines), 62)]
    to_unicode = _to_unicode_cmap(lines)
    page_ids = [6 + index * 2 for index in range(len(pages))]
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        (
            b"<< /Type /Pages /Kids ["
            + b" ".join(f"{page_id} 0 R".encode() for page_id in page_ids)
            + f"] /Count {len(page_ids)} >>".encode()
        ),
        b"<< /Type /Font /Subtype /Type0 /BaseFont /STSong-Light "
        b"/Encoding /Identity-H /DescendantFonts [4 0 R] /ToUnicode 5 0 R >>",
        b"<< /Type /Font /Subtype /CIDFontType0 /BaseFont /STSong-Light "
        b"/CIDSystemInfo << /Registry (Adobe) /Ordering (Identity) /Supplement 0 >> "
        b"/DW 1000 >>",
        (
            b"<< /Length "
            + str(len(to_unicode)).encode()
            + b" >>\nstream\n"
            + to_unicode
            + b"\nendstream"
        ),
    ]
    for page_id, page_lines in zip(page_ids, pages, strict=True):
        content_id = page_id + 1
        stream = _pdf_page_stream(page_lines)
        objects.extend(
            [
                (
                    b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    b"/Resources << /Font << /F1 3 0 R >> >> "
                    + f"/Contents {content_id} 0 R >>".encode()
                ),
                (
                    b"<< /Length "
                    + str(len(stream)).encode()
                    + b" >>\nstream\n"
                    + stream
                    + b"\nendstream"
                ),
            ]
        )

    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
    output.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    output.extend(f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(output)


def _render(kind: str, payload: dict[str, Any]) -> tuple[bytes, str]:
    lines = _render_lines(payload)
    if kind == "JSON":
        return json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode(), "application/json"
    if kind == "DOCX":
        document = Document()
        document.add_heading(lines[0], level=1)
        for line in lines[1:]:
            document.add_paragraph(line)
        stream = io.BytesIO()
        document.save(stream)
        normalized = io.BytesIO()
        with (
            zipfile.ZipFile(io.BytesIO(stream.getvalue())) as source,
            zipfile.ZipFile(
                normalized,
                "w",
                compression=zipfile.ZIP_DEFLATED,
            ) as target,
        ):
            for name in sorted(source.namelist()):
                info = source.getinfo(name)
                entry = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
                entry.compress_type = zipfile.ZIP_DEFLATED
                entry.external_attr = info.external_attr
                target.writestr(entry, source.read(name))
        return normalized.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if kind == "PDF":
        return _pdf_bytes(lines), "application/pdf"
    raise ValueError(f"unsupported artifact kind: {kind}")


def render_one_artifact(
    *,
    task_id: str,
    final_memorial_id: str,
    final_memorial_version: int,
    payload: dict[str, Any],
    kind: str,
) -> DeliveryArtifact:
    content, mime = _render(kind, payload)
    digest = hashlib.sha256(content).hexdigest()
    return DeliveryArtifact(
        artifact_id=f"artifact-{task_id}-{final_memorial_version}-{kind.lower()}",
        kind=kind,
        mime_type=mime,
        content=content,
        status="READY",
        content_hash=digest,
        byte_size=len(content),
    )


def render_artifacts(*, task_id: str, final_memorial_id: str, final_memorial_version: int, payload: dict[str, Any], required_kinds: list[str], disabled_kinds: set[str] | None = None) -> DeliveryResult:
    disabled = disabled_kinds or set()
    artifacts: list[DeliveryArtifact] = []
    for kind in required_kinds:
        artifact_id = f"artifact-{task_id}-{final_memorial_version}-{kind.lower()}"
        if kind in disabled:
            empty = b""
            artifacts.append(DeliveryArtifact(artifact_id, kind, "application/octet-stream", empty, "UNAVAILABLE", hashlib.sha256(empty).hexdigest(), 0))
            continue
        artifacts.append(
            render_one_artifact(
                task_id=task_id,
                final_memorial_id=final_memorial_id,
                final_memorial_version=final_memorial_version,
                payload=payload,
                kind=kind,
            )
        )
    statuses = {artifact.status for artifact in artifacts}
    overall = "READY" if statuses == {"READY"} else "PARTIAL"
    return DeliveryResult(task_id, final_memorial_id, final_memorial_version, tuple(artifacts), overall)
