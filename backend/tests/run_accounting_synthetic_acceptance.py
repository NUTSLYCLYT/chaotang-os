"""Run the offline accounting report flow through real FastAPI and Next servers."""

from __future__ import annotations

import argparse
import ctypes
import hashlib
import http.client
import json
import os
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
from contextlib import closing
from http.cookiejar import CookieJar
from io import BytesIO
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import quote
from urllib.request import HTTPCookieProcessor, Request, build_opener, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
FRONTEND = ROOT / "frontend"
ACCOUNTING_DECREE = "我想用本地数据分析出2025年的财务数据分析一下"
SHEETS = [
    "管理摘要",
    "核心财务报表",
    "科目趋势",
    "异常分析",
    "科目明细",
    "校验结果",
    "数据来源",
]

PROTECTED_PORTS = frozenset({3000, 8000, 13000, 18000, 13381, 18381})
_PROCESS_JOBS: dict[int, int] = {}


def _tracked_popen(args, **kwargs) -> subprocess.Popen[bytes]:
    process = subprocess.Popen(args, **kwargs)
    if sys.platform == "win32":
        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        job = kernel32.CreateJobObjectW(None, None)
        if not job:
            process.terminate()
            raise OSError(ctypes.get_last_error(), "CreateJobObjectW failed")
        information = (ctypes.c_uint32 * 36)()
        information[4] = 0x00002000  # JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE
        if not kernel32.SetInformationJobObject(job, 9, information, ctypes.sizeof(information)):
            kernel32.CloseHandle(job)
            process.terminate()
            raise OSError(ctypes.get_last_error(), "SetInformationJobObject failed")
        if not kernel32.AssignProcessToJobObject(job, process._handle):
            kernel32.CloseHandle(job)
            process.terminate()
            raise OSError(ctypes.get_last_error(), "AssignProcessToJobObject failed")
        _PROCESS_JOBS[process.pid] = job
    return process


def _free_port(excluded_ports: set[int] | frozenset[int]) -> int:
    for _attempt in range(20):
        with socket.socket() as sock:
            sock.bind(("127.0.0.1", 0))
            port = int(sock.getsockname()[1])
        if port not in excluded_ports:
            return port
    raise RuntimeError("unable to allocate an unprotected unique loopback port")


def _wait(url: str, process: subprocess.Popen[bytes]) -> None:
    deadline = time.monotonic() + 60
    while time.monotonic() < deadline:
        if process.poll() is not None:
            raise RuntimeError(f"server exited early: {process.returncode}")
        try:
            with urlopen(url, timeout=1) as response:
                if response.status == 200:
                    return
        except (HTTPError, URLError, TimeoutError, http.client.BadStatusLine):
            time.sleep(0.2)
    raise RuntimeError(f"server did not become ready: {url}")


def _json_request(opener, url: str, payload: dict[str, object], cookie: str | None = None):
    headers = {"content-type": "application/json"}
    if cookie is not None:
        headers["cookie"] = cookie
    request = Request(
        url,
        data=json.dumps(payload, ensure_ascii=False).encode(),
        method="POST",
        headers=headers,
    )
    try:
        with opener.open(request, timeout=30) as response:
            return response.status, json.loads(response.read())
    except HTTPError as error:
        detail = error.read().decode("utf-8", errors="replace")
        raise AssertionError(f"JSON request failed: status={error.code} body={detail}") from error


def _poll_job(
    opener,
    base: str,
    job_id: str,
    cookie: str,
    *,
    diagnostic_base: str | None = None,
    diagnostic_owner_user_id: str | None = None,
) -> dict[str, object]:
    deadline = time.monotonic() + 60
    last: dict[str, object] | None = None
    observations: list[dict[str, object]] = []
    while time.monotonic() < deadline:
        request = Request(
            f"{base}/api/decree-jobs/{job_id}",
            headers={"cookie": cookie},
        )
        try:
            with opener.open(request, timeout=30) as response:
                assert response.status == 200
                last = json.loads(response.read())
                bff_headers = {
                    key: response.headers.get(key)
                    for key in ("date", "age", "x-nextjs-cache", "cache-control")
                }
        except HTTPError as error:
            detail = error.read().decode("utf-8", errors="replace")
            if error.code == 503:
                if diagnostic_base is not None and len(observations) < 20:
                    token = cookie.split("=", 1)[1]
                    direct_request = Request(
                        f"{diagnostic_base}/api/v1/decree-jobs/{job_id}",
                        headers={"authorization": f"Bearer {token}"},
                    )
                    try:
                        with urlopen(direct_request, timeout=5) as direct_response:
                            direct_body = json.loads(direct_response.read())
                    except (HTTPError, URLError, TimeoutError, ValueError) as direct_error:
                        direct_body = {"error": type(direct_error).__name__}
                    observations.append(
                        {"bff_status": 503, "bff_body": detail, "direct": direct_body}
                    )
                time.sleep(0.1)
                continue
            raise AssertionError(f"job poll failed: status={error.code} body={detail}") from error
        if last.get("state") in {"SUCCEEDED", "FAILED", "CANCELLED"}:
            return last
        if diagnostic_base is not None and len(observations) < 20:
            token = cookie.split("=", 1)[1]
            direct_request = Request(
                f"{diagnostic_base}/api/v1/decree-jobs/{job_id}",
                headers={"authorization": f"Bearer {token}"},
            )
            try:
                with urlopen(direct_request, timeout=5) as direct_response:
                    direct_body = json.loads(direct_response.read())
                    direct_headers = {
                        key: direct_response.headers.get(key)
                        for key in ("date", "age", "x-nextjs-cache", "cache-control")
                    }
            except (HTTPError, URLError, TimeoutError, ValueError) as error:
                direct_body = {"error": type(error).__name__}
                direct_headers = {}
            observations.append(
                {
                    "bff": last,
                    "bff_headers": bff_headers,
                    "direct": direct_body,
                    "direct_headers": direct_headers,
                }
            )
        time.sleep(0.1)
    diagnostic: object = None
    if diagnostic_base is not None:
        try:
            with urlopen(
                f"{diagnostic_base}/__synthetic__/worker-diagnostics/{job_id}"
                f"?owner_user_id={quote(diagnostic_owner_user_id or '')}",
                timeout=5,
            ) as response:
                diagnostic = json.loads(response.read())
        except (HTTPError, URLError, TimeoutError, ValueError) as error:
            diagnostic = {"diagnostic_error": type(error).__name__}
    raise AssertionError(
        f"decree job did not reach a terminal state: {last}; diagnostic={diagnostic}; "
        f"observations={observations}"
    )


def _stop(process: subprocess.Popen[bytes]) -> None:
    job = _PROCESS_JOBS.pop(process.pid, None)
    if job is not None:
        ctypes.WinDLL("kernel32", use_last_error=True).CloseHandle(job)
        process.wait(timeout=10)
        return
    if process.poll() is not None:
        return
    process.terminate()
    try:
        process.wait(timeout=10)
    except subprocess.TimeoutExpired:
        process.kill()
        process.wait(timeout=10)


def _expect_http_error(opener, request: Request, expected: int) -> None:
    try:
        opener.open(request, timeout=30)
    except HTTPError as error:
        try:
            assert error.code == expected
        finally:
            error.close()
    else:
        raise AssertionError(f"request unexpectedly succeeded; expected HTTP {expected}")


def _synthetic_diagnostics(backend_port: int) -> dict[str, object]:
    url = f"http://127.0.0.1:{backend_port}/__synthetic__/diagnostics"
    with urlopen(url, timeout=5) as response:
        return json.loads(response.read())


def _start_servers(
    backend_port: int,
    frontend_port: int,
    environment: dict[str, str],
) -> tuple[subprocess.Popen[bytes], subprocess.Popen[bytes]]:
    backend: subprocess.Popen[bytes] | None = None
    frontend: subprocess.Popen[bytes] | None = None
    log_root = Path(
        environment.get(
            "CHAOTANG_SYNTHETIC_SERVER_LOG_DIR",
            environment.get(
                "CHAOTANG_SYNTHETIC_ACCEPTANCE_TMP",
                tempfile.gettempdir(),
            ),
        )
    )
    log_root.mkdir(parents=True, exist_ok=True)
    backend_log = (log_root / "backend-server.log").open("ab")
    frontend_log = (log_root / "frontend-server.log").open("ab")
    try:
        backend = _tracked_popen(
            [
                sys.executable,
                "-m",
                "uvicorn",
                "tests.synthetic_accounting_acceptance_app:app",
                "--host",
                "127.0.0.1",
                "--port",
                str(backend_port),
            ],
            cwd=BACKEND,
            env=environment,
            stdout=backend_log,
            stderr=subprocess.STDOUT,
        )
        next_environment = environment.copy()
        next_environment["BACKEND_BASE_URL"] = f"http://127.0.0.1:{backend_port}"
        frontend = _tracked_popen(
            [
                "node",
                "node_modules/next/dist/bin/next",
                "start",
                "--hostname",
                "127.0.0.1",
                "--port",
                str(frontend_port),
            ],
            cwd=FRONTEND,
            env=next_environment,
            stdout=frontend_log,
            stderr=subprocess.STDOUT,
        )
        return backend, frontend
    except Exception:
        if frontend is not None:
            _stop(frontend)
        if backend is not None:
            _stop(backend)
        raise
    finally:
        backend_log.close()
        frontend_log.close()


def main(
    *,
    decree: str = ACCOUNTING_DECREE,
    rounds: int = 1,
    expected_decree_sha256: str | None = None,
) -> int:
    requested_decree = decree
    actual_decree_sha256 = hashlib.sha256(decree.encode()).hexdigest()
    if expected_decree_sha256 is not None and actual_decree_sha256 != expected_decree_sha256:
        raise ValueError(
            "acceptance_decree_sha256_mismatch:"
            f"actual={actual_decree_sha256}:expected={expected_decree_sha256}"
        )
    if not decree.strip() or "2025" not in decree or rounds != 1:
        raise ValueError("acceptance_requires_nonblank_2025_decree_and_one_round")
    formal_round = int(os.environ.get("CHAOTANG_ACCEPTANCE_ROUND", "1"))
    if not 1 <= formal_round <= 10:
        raise ValueError("acceptance_round_out_of_range")
    command = (
        "python tests/run_accounting_synthetic_acceptance.py "
        f"--decree {json.dumps(decree, ensure_ascii=False)} --rounds 1"
    )
    snapshot_path = os.environ.get("CHAOTANG_ACCEPTANCE_INPUT_SNAPSHOT")
    if snapshot_path:
        selected_environment = {
            key: hashlib.sha256(value.encode()).hexdigest()
            for key, value in sorted(os.environ.items())
            if key.startswith(("CHAOTANG_", "PYTHON", "NODE"))
        }
        Path(snapshot_path).write_text(
            json.dumps(
                {
                    "argv": sys.argv,
                    "decree_sha256": hashlib.sha256(decree.encode("utf-8")).hexdigest(),
                    "decree_utf8_hex": decree.encode("utf-8").hex(),
                    "cwd": str(Path.cwd()),
                    "formal_round": formal_round,
                    "environment_hashes": selected_environment,
                },
                ensure_ascii=False,
                sort_keys=True,
            ),
            encoding="utf-8",
        )
    used_ports = set(PROTECTED_PORTS)
    backend_port = _free_port(used_ports)
    used_ports.add(backend_port)
    frontend_port = _free_port(used_ports)
    with tempfile.TemporaryDirectory(prefix="chaotang-accounting-acceptance-") as tmp:
        environment = os.environ.copy()
        environment["CHAOTANG_SYNTHETIC_ACCEPTANCE_TMP"] = tmp
        environment["CHAOTANG_SYNTHETIC_ACCEPTANCE_DECREE"] = decree
        environment["CHAOTANG_DECREE_JOB_WORKER_ENABLED"] = "1"
        environment["DEEPSEEK_API_KEY"] = ""
        environment["JINYIWEI_EXTERNAL_NETWORK_ENABLED"] = ""
        environment["JINYIWEI_MCP_CREDENTIAL_SOURCE"] = "env"
        backend, frontend = _start_servers(backend_port, frontend_port, environment)
        try:
            _wait(f"http://127.0.0.1:{backend_port}/health", backend)
            _wait(f"http://127.0.0.1:{frontend_port}/study", frontend)
            startup_diagnostic = _synthetic_diagnostics(backend_port)
            assert startup_diagnostic["pid"] == backend.pid or startup_diagnostic["start_nonce"]
            startup_nonce = startup_diagnostic["start_nonce"]
            base = f"http://127.0.0.1:{frontend_port}"
            owner_jar, other_jar, failure_jar = CookieJar(), CookieJar(), CookieJar()
            owner = build_opener(HTTPCookieProcessor(owner_jar))
            other = build_opener(HTTPCookieProcessor(other_jar))
            failure_owner = build_opener(HTTPCookieProcessor(failure_jar))
            session_cookies: dict[str, str] = {}
            for opener, jar, username in (
                (owner, owner_jar, "synthetic-owner"),
                (other, other_jar, "synthetic-other"),
                (failure_owner, failure_jar, "synthetic-failure"),
            ):
                status, _body = _json_request(
                    opener,
                    f"{base}/api/auth/register",
                    {
                        "username": username,
                        "email": f"{username}@example.test",
                        "password": "six-or-more",
                    },
                )
                assert status == 201
                status, _body = _json_request(
                    opener,
                    f"{base}/api/auth/login",
                    {"identifier": username, "password": "six-or-more"},
                )
                assert status == 200
                cookie = next(iter(jar))
                session_cookies[username] = f"{cookie.name}={cookie.value}"

            status, draft = _json_request(
                owner,
                f"{base}/api/drafts/chancellor",
                {
                    "messages": [{"role": "user", "content": decree}],
                    "version": 1,
                },
                session_cookies["synthetic-owner"],
            )
            assert status == 200
            assert draft["status"] == "DRAFT_READY"
            assert "2025" in draft["expert_example"]
            assert "2025" in draft["decree_text"]
            assert "2025" in decree
            assert draft["draft"]["departments"] == [
                {
                    "department": "户部",
                    "bureaus": ["会计司"],
                    "role": "主审",
                    "reason": "负责财务报表。",
                    "responsibility": "生成并校验报表。",
                    "expected_output": "可下载的 Excel 文件。",
                }
            ]

            replacement = Request(
                f"{base}/api/decrees/chancellor",
                data=json.dumps(
                    {
                        "decreeText": "客户端替换的 2024 年财务报表",
                        "draftVersion": draft["version"],
                        "draftFingerprint": draft["fingerprint"],
                        "idempotencyKey": "synthetic-owner-replacement",
                    },
                    ensure_ascii=False,
                ).encode(),
                method="POST",
                headers={
                    "content-type": "application/json",
                    "cookie": session_cookies["synthetic-owner"],
                },
            )
            _expect_http_error(owner, replacement, 409)

            status, decree = _json_request(
                owner,
                f"{base}/api/decrees/chancellor",
                {
                    "decreeText": draft["decree_text"],
                    "draftVersion": draft["version"],
                    "draftFingerprint": draft["fingerprint"],
                    "idempotencyKey": "synthetic-owner-success",
                },
                session_cookies["synthetic-owner"],
            )
            assert status == 202
            assert decree["state"] == "QUEUED"
            job = _poll_job(
                owner,
                base,
                decree["jobId"],
                session_cookies["synthetic-owner"],
                diagnostic_base=f"http://127.0.0.1:{backend_port}",
            )
            with closing(sqlite3.connect(Path(tmp) / "decree_jobs.sqlite3")) as connection:
                job_diagnostic = connection.execute(
                    "SELECT state, error_code, error_stage, error_category "
                    "FROM decree_jobs WHERE job_id = ?",
                    (decree["jobId"],),
                ).fetchone()
            assert job["state"] == "SUCCEEDED", (
                job,
                job_diagnostic,
                {
                    "expected": "SUCCEEDED",
                    "decree_sha256": hashlib.sha256(requested_decree.encode()).hexdigest(),
                    "startup_nonce": startup_nonce,
                    "failure_diagnostic": _synthetic_diagnostics(backend_port),
                },
            )
            assert job["error"] is None
            decree = job["result"]
            assert isinstance(decree, dict)
            assert decree["routeType"] == "single"
            assert decree["departments"] == ["户部"]
            assert decree["processingPath"] == [
                "上书房",
                "丞相（首次分流）",
                "户部",
                "户部·会计司",
                "户部（部级补充）",
                "丞相（最终汇总）",
            ]
            assert decree["councilVerdict"] is None
            assert decree["deliveryKind"] == "accounting_report"
            assert len(decree["artifacts"]) == 1
            artifact = decree["artifacts"][0]
            assert artifact["periodStart"] == artifact["periodEnd"] == 2025
            assert artifact["kind"] == "ACCOUNTING_MANAGEMENT_REPORT_XLSX"
            assert artifact["displayName"].endswith(".xlsx")
            artifact_id = artifact["artifactId"]
            download_url = f"{base}/api/report-artifacts/{artifact_id}"
            owner_request = Request(
                download_url,
                headers={"cookie": session_cookies["synthetic-owner"]},
            )
            with owner.open(owner_request, timeout=30) as response:
                assert response.status == 200
                assert response.headers.get_content_type() == (
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
                assert response.headers["content-disposition"].startswith(
                    "attachment; filename*=UTF-8''"
                )
                assert response.headers["content-disposition"].endswith(".xlsx")
                workbook_bytes = response.read()
            try:
                other.open(
                    Request(
                        download_url,
                        headers={"cookie": session_cookies["synthetic-other"]},
                    ),
                    timeout=30,
                )
            except HTTPError as error:
                try:
                    assert error.code == 404
                finally:
                    error.close()
            else:
                raise AssertionError("second owner unexpectedly downloaded the artifact")

            unsafe_url = f"{base}/api/report-artifacts/{quote('../secret', safe='')}"
            _expect_http_error(
                owner,
                Request(unsafe_url, headers={"cookie": session_cookies["synthetic-owner"]}),
                400,
            )

            workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)
            try:
                assert workbook.sheetnames == SHEETS
                details = workbook["科目明细"]
                assert {details.cell(row, 1).value for row in range(2, details.max_row + 1)} == {
                    2025
                }
                assets = sum(
                    int(str(details.cell(row, 12).value).lstrip("="))
                    for row in range(2, details.max_row + 1)
                    if details.cell(row, 14).value == "assets"
                )
                assert assets == 1000
                checks = workbook["校验结果"]
                assert checks.cell(1, 1).value == "校验项"
                sources = workbook["数据来源"]
                assert {sources.cell(row, 1).value for row in range(1, sources.max_row + 1)} >= {
                    "报告期间",
                    "来源哈希",
                }
                assert (
                    next(
                        sources.cell(row, 2).value
                        for row in range(1, sources.max_row + 1)
                        if sources.cell(row, 1).value == "报告期间"
                    )
                    == "2025-2025"
                )
            finally:
                workbook.close()

            artifact_hash = hashlib.sha256(workbook_bytes).hexdigest()

            with closing(sqlite3.connect(Path(tmp) / "artifacts.sqlite3")) as connection:
                report_row = connection.execute(
                    "SELECT artifact_id, reply_id, state, period_start, period_end "
                    "FROM report_artifacts"
                ).fetchone()
            assert report_row is not None
            assert report_row[0] == artifact_id
            assert report_row[1]
            assert report_row[2:] == ("PUBLISHED", 2025, 2025)
            with closing(sqlite3.connect(Path(tmp) / "shiguan.sqlite3")) as connection:
                reply = connection.execute(
                    "SELECT id, type, source_kind, source_text FROM archives WHERE id = ?",
                    (report_row[1],),
                ).fetchone()
            assert reply == (report_row[1], "REPLY", "DECREE", draft["decree_text"])
            with closing(sqlite3.connect(Path(tmp) / "shiguan.sqlite3")) as connection:
                final_reply_count = connection.execute(
                    "SELECT COUNT(*) FROM archives WHERE type = 'REPLY'"
                ).fetchone()[0]
            assert final_reply_count == 1
            assert {path.name for path in (Path(tmp) / "report_artifacts").iterdir()} == {
                f"{artifact_id}.xlsx"
            }

            toggle = Request(
                f"http://127.0.0.1:{backend_port}/__synthetic__/fail-next",
                data=b"",
                method="POST",
            )
            try:
                with urlopen(toggle, timeout=30) as response:
                    assert response.status == 204
            except HTTPError as error:
                try:
                    raise AssertionError(
                        f"failure injection endpoint returned HTTP {error.code}"
                    ) from error
                finally:
                    error.close()

            status, failed_draft = _json_request(
                failure_owner,
                f"{base}/api/drafts/chancellor",
                {
                    "messages": [{"role": "user", "content": ACCOUNTING_DECREE}],
                    "version": 1,
                },
                session_cookies["synthetic-failure"],
            )
            assert status == 200
            assert failed_draft["status"] == "DRAFT_READY"
            failed_request = Request(
                f"{base}/api/decrees/chancellor",
                data=json.dumps(
                    {
                        "decreeText": failed_draft["decree_text"],
                        "draftVersion": failed_draft["version"],
                        "draftFingerprint": failed_draft["fingerprint"],
                        "idempotencyKey": "synthetic-failure",
                    },
                    ensure_ascii=False,
                ).encode(),
                method="POST",
                headers={
                    "content-type": "application/json",
                    "cookie": session_cookies["synthetic-failure"],
                },
            )
            with failure_owner.open(failed_request, timeout=30) as response:
                assert response.status == 202
                failed_acceptance = json.loads(response.read())
            failed_job = _poll_job(
                failure_owner,
                base,
                failed_acceptance["jobId"],
                session_cookies["synthetic-failure"],
                diagnostic_base=f"http://127.0.0.1:{backend_port}",
            )
            assert failed_job["state"] == "FAILED"
            assert failed_job["result"] is None
            assert failed_job["error"] == {
                "code": "job_failed",
                "stage": "execution",
                "category": "internal",
            }
            with closing(sqlite3.connect(Path(tmp) / "artifacts.sqlite3")) as connection:
                states = [
                    row[0]
                    for row in connection.execute(
                        "SELECT state FROM report_artifacts ORDER BY created_at"
                    )
                ]
            assert states == ["PUBLISHED", "ABORTED"]
            assert {path.name for path in (Path(tmp) / "report_artifacts").iterdir()} == {
                f"{artifact_id}.xlsx"
            }
            os.environ["CHAOTANG_SYNTHETIC_ACCEPTANCE_TMP"] = tmp
            from tests.synthetic_accounting_acceptance_app import _run_dynamic_layout_matrix

            dynamic_matrix = _run_dynamic_layout_matrix(Path(tmp) / "dynamic-layouts")
            evidence = {
                "round": formal_round,
                "exit_code": 0,
                "command": command,
                "status": "PASS",
                "artifact_sha256": artifact_hash,
                "workbook_open": True,
                "tool_audit_refs": dynamic_matrix["alternate_tool_audit_refs"],
                "final_reply_count": final_reply_count,
                "error_classification": failed_job["error"],
                "dynamic_layout_matrix": dynamic_matrix,
            }
            print(json.dumps(evidence, ensure_ascii=False, sort_keys=True))
            print(
                "synthetic-accounting-acceptance: PASS "
                "(explicit-analysis-2025, canonical authority, published identity, "
                "owner isolation, "
                "unsafe ID rejection, seven sheets, source/check disclosures)"
            )
            return 0
        finally:
            _stop(frontend)
            _stop(backend)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Run one frozen synthetic accounting acceptance round."
    )
    parser.add_argument("--decree", default=ACCOUNTING_DECREE)
    parser.add_argument("--rounds", type=int, default=1)
    parser.add_argument("--expected-decree-sha256")
    arguments = parser.parse_args()
    sys.exit(
        main(
            decree=arguments.decree,
            rounds=arguments.rounds,
            expected_decree_sha256=arguments.expected_decree_sha256,
        )
    )
