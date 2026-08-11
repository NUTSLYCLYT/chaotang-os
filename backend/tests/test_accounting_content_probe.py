from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from app.accounting_reports.content_probe import probe_accounting_sources
from app.accounting_reports.models import ReportPeriod
from app.accounting_reports.sources import AccountingSourceError


def _write_workbook(path: Path, sheets: dict[str, list[list[object]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    workbook.close()


def _rewrite_xlsx(path: Path, replacements: dict[str, bytes]) -> None:
    source = BytesIO(path.read_bytes())
    target = BytesIO()
    with ZipFile(source) as original, ZipFile(target, "w", ZIP_DEFLATED) as rewritten:
        for item in original.infolist():
            rewritten.writestr(item, replacements.get(item.filename, original.read(item)))
        for name, value in replacements.items():
            if name not in original.namelist():
                rewritten.writestr(name, value)
    path.write_bytes(target.getvalue())


def _rename_xlsx_parts(path: Path, renames: dict[str, str], replacements: dict[str, bytes]) -> None:
    target = BytesIO()
    with ZipFile(path) as original, ZipFile(target, "w", ZIP_DEFLATED) as rewritten:
        for item in original.infolist():
            rewritten.writestr(
                renames.get(item.filename, item.filename),
                replacements.get(item.filename, original.read(item)),
            )
    path.write_bytes(target.getvalue())


def test_probe_finds_year_and_multiline_header_without_fixed_filename(tmp_path: Path) -> None:
    _write_workbook(
        tmp_path / "internal-A.xlsx",
        {"余额明细": [["2025年度余额表"], ["科目", "期初余额"], ["编码", "借方"], ["1001", 3]]},
    )

    probes = probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))

    assert probes[0].years == (2025,)
    assert probes[0].sheets[0].regions[0].header_depth == 2


def test_probe_preserves_title_section_and_auxiliary_detail_rows(tmp_path: Path) -> None:
    _write_workbook(
        tmp_path / "financial-material.xlsx",
        {
            "Sheet1": [
                ["2025财务资料"],
                ["项目", "金额"],
                ["资产类"],
                ["库存现金", 2],
                ["其中：受限资金", 1],
            ]
        },
    )

    sheet = probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))[0].sheets[0]

    assert sheet.rejected_rows == ()
    assert sheet.regions[0].start_row == 1
    assert sheet.regions[0].end_row == 5


def test_probe_inventories_multiple_sheets(tmp_path: Path) -> None:
    _write_workbook(
        tmp_path / "book.xlsx",
        {
            "资产负债表": [["2025年"], ["项目", "金额"], ["资产", 1]],
            "辅助明细": [["2025年"], ["名称", "金额"], ["A", 2]],
        },
    )

    probe = probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))[0]

    assert tuple(sheet.name for sheet in probe.sheets) == ("资产负债表", "辅助明细")
    assert all(sheet.regions for sheet in probe.sheets)


def test_probe_records_formula_as_untrusted_data_without_evaluating(tmp_path: Path) -> None:
    _write_workbook(
        tmp_path / "formula.xlsx",
        {"Sheet1": [["2025年"], ["项目", "金额"], ["合计", "=1+1"]]},
    )

    region = probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))[0].sheets[0].regions[0]

    assert region.formula_cells == ((3, 2, "=1+1"),)


def test_probe_preserves_values_types_and_merged_ranges(tmp_path: Path) -> None:
    path = tmp_path / "typed.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "2025 title"
    sheet.merge_cells("A1:B1")
    sheet.append(["item", "amount"])
    sheet.append(["cash", 3])
    workbook.save(path)
    workbook.close()

    probe = probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))[0]

    assert probe.sheets[0].merged_ranges == ("A1:B1",)
    cells = probe.sheets[0].regions[0].cells
    assert (cells[0].value, cells[0].value_type) == ("2025 title", "string")
    assert (cells[-1].value, cells[-1].value_type) == (3, "number")
    assert "2025 title" not in repr(cells[0])
    assert "cash" not in repr(probe)


def test_probe_returns_each_bounded_nonempty_region(tmp_path: Path) -> None:
    _write_workbook(
        tmp_path / "regions.xlsx",
        {
            "Sheet1": [
                ["2025 title"],
                ["item", "amount"],
                ["cash", 3],
                [],
                ["aux", "amount"],
                ["detail", 1],
            ]
        },
    )

    regions = probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))[0].sheets[0].regions

    assert tuple((region.start_row, region.end_row) for region in regions) == ((1, 3), (5, 6))


@pytest.mark.parametrize(
    "hazard", ["vba", "macro_content_type", "external_link", "external_relationship"]
)
def test_probe_rejects_hazardous_ooxml_packages(
    tmp_path: Path, hazard: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "hazard.xlsx"
    _write_workbook(path, {"Sheet1": [["2025"]]})
    with ZipFile(path) as archive:
        content_types = archive.read("[Content_Types].xml")
        workbook_rels = archive.read("xl/_rels/workbook.xml.rels")
    replacements = {
        "vba": {"xl/vbaProject.bin": b"macro"},
        "macro_content_type": {
            "[Content_Types].xml": content_types.replace(
                b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
            )
        },
        "external_link": {"xl/externalLinks/externalLink1.xml": b"<externalLink/>"},
        "external_relationship": {
            "xl/_rels/workbook.xml.rels": workbook_rels.replace(
                b"</Relationships>",
                b'<Relationship Id="external" '
                b'Target="https://example.invalid/x" '
                b'TargetMode="External"/></Relationships>',
            )
        },
    }[hazard]
    _rewrite_xlsx(path, replacements)
    from app.accounting_reports import content_probe

    monkeypatch.setattr(
        content_probe,
        "load_workbook",
        lambda *_args, **_kwargs: pytest.fail("hazard must be rejected before parsing"),
    )

    with pytest.raises(AccountingSourceError, match="^source_schema_invalid$"):
        probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))


def test_probe_rejects_declared_oversized_ooxml_dimension_before_iteration(tmp_path: Path) -> None:
    path = tmp_path / "oversized.xlsx"
    _write_workbook(path, {"Sheet1": [["2025"]]})
    with ZipFile(path) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml")
    _rewrite_xlsx(
        path,
        {
            "xl/worksheets/sheet1.xml": sheet_xml.replace(
                b'ref="A1:A1"', b'ref="A1:ZZ30000"'
            )
        },
    )

    with pytest.raises(AccountingSourceError, match="^source_schema_invalid$"):
        probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))


def test_ooxml_sheet_parts_follow_workbook_relationships_not_part_names(tmp_path: Path) -> None:
    path = tmp_path / "custom-parts.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "First"
    first["A1"] = "2025 first"
    first.merge_cells("A1:B1")
    second = workbook.create_sheet("Second")
    second["A1"] = "2025 second"
    second.merge_cells("A1:C1")
    workbook.save(path)
    workbook.close()
    with ZipFile(path) as archive:
        rels = archive.read("xl/_rels/workbook.xml.rels")
    rels = rels.replace(b"/xl/worksheets/sheet1.xml", b"/xl/worksheets/custom-b.xml")
    rels = rels.replace(b"/xl/worksheets/sheet2.xml", b"/xl/worksheets/custom-a.xml")
    relationships = ElementTree.fromstring(rels)
    relationships[:] = reversed(relationships[:])
    rels = ElementTree.tostring(relationships)
    _rename_xlsx_parts(
        path,
        {
            "xl/worksheets/sheet1.xml": "xl/worksheets/custom-b.xml",
            "xl/worksheets/sheet2.xml": "xl/worksheets/custom-a.xml",
        },
        {"xl/_rels/workbook.xml.rels": rels},
    )

    sheets = probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))[0].sheets

    assert tuple(sheet.name for sheet in sheets) == ("First", "Second")
    assert tuple(sheet.merged_ranges for sheet in sheets) == (("A1:B1",), ("A1:C1",))


def test_openpyxl_workbook_closes_when_iteration_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "close.xlsx"
    _write_workbook(path, {"Sheet1": [["2025"]]})
    from app.accounting_reports import content_probe

    class BrokenSheet:
        title = "Sheet1"

        def iter_rows(self):
            raise RuntimeError("iteration failed")

    class BrokenWorkbook:
        worksheets = (BrokenSheet(),)
        closed = False

        def close(self) -> None:
            self.closed = True

    workbook = BrokenWorkbook()
    monkeypatch.setattr(content_probe, "load_workbook", lambda *_args, **_kwargs: workbook)

    with pytest.raises(AccountingSourceError, match="^source_schema_invalid$"):
        probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))
    assert workbook.closed is True


def test_ooxml_external_relationship_is_rejected_by_xml_semantics(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "semantic-external.xlsx"
    _write_workbook(path, {"Sheet1": [["2025"]]})
    with ZipFile(path) as archive:
        rels = archive.read("xl/_rels/workbook.xml.rels")
    rels = rels.replace(
        b"</Relationships>",
        b'<Relationship Id="external" Target="https://example.invalid/x" '
        b'TargetMode="  ExTeRnAl  "/></Relationships>',
    )
    _rewrite_xlsx(path, {"xl/_rels/workbook.xml.rels": rels})
    from app.accounting_reports import content_probe

    monkeypatch.setattr(
        content_probe,
        "load_workbook",
        lambda *_args, **_kwargs: pytest.fail("external rel must fail before parsing"),
    )

    with pytest.raises(AccountingSourceError, match="^source_schema_invalid$"):
        probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))


def test_ooxml_merged_ranges_use_xml_attribute_semantics(tmp_path: Path) -> None:
    path = tmp_path / "merged-xml.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "2025"
    sheet.merge_cells("A1:B1")
    workbook.save(path)
    workbook.close()
    with ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml")
    xml = xml.replace(b'<mergeCell ref="A1:B1"/>', b"<mergeCell   ref = 'A1:B1' />")
    _rewrite_xlsx(path, {"xl/worksheets/sheet1.xml": xml})

    merged = probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))[0].sheets[0].merged_ranges

    assert merged == ("A1:B1",)


def test_ooxml_merge_reference_cannot_exceed_system_sheet_bounds(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "oversized-merge.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "2025"
    sheet.merge_cells("A1:B1")
    workbook.save(path)
    workbook.close()
    with ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    merge = next(
        element
        for element in root.iter()
        if element.tag.rsplit("}", 1)[-1] == "mergeCell"
    )
    merge.attrib["ref"] = "A1:A20001"
    _rewrite_xlsx(
        path,
        {"xl/worksheets/sheet1.xml": ElementTree.tostring(root)},
    )
    from app.accounting_reports import content_probe

    monkeypatch.setattr(
        content_probe,
        "load_workbook",
        lambda *_args, **_kwargs: pytest.fail("invalid merge must fail before OpenPyXL"),
    )

    with pytest.raises(AccountingSourceError, match="^source_schema_invalid$"):
        probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))


def test_ooxml_merge_count_is_bounded_before_openpyxl(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.accounting_reports import content_probe

    path = tmp_path / "many-merges.xlsx"
    _write_workbook(path, {"Sheet1": [["2025"]]})
    with ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml")
    marker = b"</worksheet>"
    merge_cells = b"<mergeCells>" + b"".join(
        b'<mergeCell ref="A1:A1"/>'
        for _ in range(content_probe.MAX_MERGED_RANGES_PER_SHEET + 1)
    ) + b"</mergeCells>"
    _rewrite_xlsx(path, {"xl/worksheets/sheet1.xml": xml.replace(marker, merge_cells + marker)})

    monkeypatch.setattr(
        content_probe,
        "load_workbook",
        lambda *_args, **_kwargs: pytest.fail("excess merges must fail before OpenPyXL"),
    )

    with pytest.raises(AccountingSourceError, match="^source_schema_invalid$"):
        probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))


def test_ooxml_dimension_uses_parsed_attribute_semantics_before_openpyxl(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "semantic-dimension.xlsx"
    _write_workbook(path, {"Sheet1": [["2025"]]})
    with ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml")
    root = ElementTree.fromstring(xml)
    dimension = next(
        element
        for element in root.iter()
        if element.tag.rsplit("}", 1)[-1] == "dimension"
    )
    dimension.attrib["ref"] = "A1:ZZ30000"
    xml = ElementTree.tostring(root).replace(
        b'ref="A1:ZZ30000"', b"ref = 'A1:ZZ30000'"
    )
    _rewrite_xlsx(path, {"xl/worksheets/sheet1.xml": xml})
    from app.accounting_reports import content_probe

    monkeypatch.setattr(
        content_probe,
        "load_workbook",
        lambda *_args, **_kwargs: pytest.fail("dimension must fail before OpenPyXL"),
    )

    with pytest.raises(AccountingSourceError, match="^source_schema_invalid$"):
        probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))


def test_content_types_utf16_macro_value_is_rejected_semantically(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "semantic-content-type.xlsx"
    _write_workbook(path, {"Sheet1": [["2025"]]})
    with ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("[Content_Types].xml"))
    for element in root:
        if element.attrib.get("PartName") == "/xl/workbook.xml":
            element.attrib["ContentType"] = (
                "  Application/Vnd.Ms-Excel.Sheet.MacroEnabled.Main+Xml  "
            )
    encoded = ElementTree.tostring(root, encoding="utf-16", xml_declaration=True)
    _rewrite_xlsx(path, {"[Content_Types].xml": encoded})
    from app.accounting_reports import content_probe

    monkeypatch.setattr(
        content_probe,
        "load_workbook",
        lambda *_args, **_kwargs: pytest.fail("macro type must fail before OpenPyXL"),
    )

    with pytest.raises(AccountingSourceError, match="^source_schema_invalid$"):
        probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))


@pytest.mark.parametrize("element_kind", ["dimension", "mergeCell", "namespaced_ref"])
def test_ooxml_structural_elements_require_official_spreadsheetml_namespace(
    tmp_path: Path, element_kind: str
) -> None:
    path = tmp_path / "namespace.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "2025"
    sheet.merge_cells("A1:B1")
    workbook.save(path)
    workbook.close()
    with ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    target = next(
        element
        for element in root.iter()
        if element.tag.rsplit("}", 1)[-1]
        == ("dimension" if element_kind in {"dimension", "namespaced_ref"} else "mergeCell")
    )
    if element_kind == "namespaced_ref":
        value = target.attrib.pop("ref")
        target.attrib["{urn:evil}ref"] = value
    else:
        target.tag = f"{{urn:evil}}{element_kind}"
    _rewrite_xlsx(
        path,
        {"xl/worksheets/sheet1.xml": ElementTree.tostring(root)},
    )

    with pytest.raises(AccountingSourceError, match="^source_schema_invalid$"):
        probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))


def test_probe_rejects_invalid_container_without_leaking_path(tmp_path: Path) -> None:
    sensitive = tmp_path / "private-name.xlsx"
    sensitive.write_bytes(b"not-a-workbook")

    with pytest.raises(AccountingSourceError) as caught:
        probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))

    assert str(caught.value) == "source_schema_invalid"
    assert "private-name" not in repr(caught.value)


def test_probe_rejects_symlink_candidate(tmp_path: Path) -> None:
    outside = tmp_path.parent / "outside-probe.xlsx"
    _write_workbook(outside, {"Sheet1": [["2025年"]]})
    link = tmp_path / "linked.xlsx"
    try:
        link.symlink_to(outside)
    except OSError:
        pytest.skip("symlink creation is unavailable")

    with pytest.raises(AccountingSourceError, match="^source_path_invalid$"):
        probe_accounting_sources(tmp_path, ReportPeriod(2025, 2025))
