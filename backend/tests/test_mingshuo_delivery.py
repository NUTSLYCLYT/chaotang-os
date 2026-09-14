"""Mingshuo Fact Pack to non-binding work-product delivery tests."""

from __future__ import annotations

import hashlib
import json
import sqlite3
import threading
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.accounting_reports.storage import ArtifactStorage, ArtifactStorageError
from app.api import mingshuo as mingshuo_api
from app.api import report_artifacts as report_api
from app.api.auth import register_auth_exception_handlers, require_current_user
from app.auth.models import AuthenticatedPrincipal
from app.mingshuo import delivery, service, storage


def _owner(*, user_id: str = "owner-a", tenant_id: str = "tenant-a") -> AuthenticatedPrincipal:
    return AuthenticatedPrincipal(
        id=user_id,
        username=user_id,
        email=f"{user_id}@example.test",
        tenant_id=tenant_id,
        membership_id=f"membership-{user_id}",
        tenant_role="OWNER",
    )


def _project_payload(*, request_key: str = "mingshuo-delivery-create-0001") -> dict[str, object]:
    evidence_digest = "sha256:" + hashlib.sha256(b"synthetic-evidence").hexdigest()
    return {
        "requestKey": request_key,
        "projectName": "Synthetic Mingshuo delivery",
        "requirementsText": "=HYPERLINK(\"https://invalid.test\") must remain text",
        "productLines": ["PACK_POWER"],
        "markets": ["DE"],
        "languages": ["en"],
        "skuCandidates": [
            {
                "id": f"sku-{suffix}",
                "label": f"Candidate {suffix.upper()}",
                "status": "EVIDENCE_BOUND" if suffix == "a" else "RESERVED",
                "parameterStatus": "EVIDENCE_BOUND" if suffix == "a" else "PARTIAL",
            }
            for suffix in ("a", "b", "c")
        ],
        "evidence": [
            {
                "id": "evidence-a",
                "sourceClass": "THIRD_PARTY_VERIFIED",
                "digest": evidence_digest,
                "validUntil": "2099-12-31",
                "adoptionStatus": "ADOPTED",
            }
        ],
        "facts": [
            {
                "id": "fact-a",
                "kind": "PARAMETER",
                "subject": "nominal voltage",
                "value": "51.2V synthetic",
                "evidenceRefs": ["evidence-a"],
            }
        ],
        "claims": [
            {
                "id": "claim-a",
                "text": "Synthetic evidence-bound claim.",
                "evidenceRefs": ["evidence-a"],
            }
        ],
        "channels": [{"id": "WEBSITE"}],
        "quoteRequested": False,
    }


def _pack() -> dict[str, object]:
    payload = _project_payload()
    return service._build_pack(  # noqa: SLF001 - exact pure producer fixture
        mingshuo_api.CreateProjectRequest.model_validate(payload),
        _owner(),
        project_id="1" * 32,
        project_name=str(payload["projectName"]),
    )


def _evaluation(pack: dict[str, object]) -> dict[str, object]:
    result = service.fact_pack.evaluate_pack(pack, now="2026-09-13")
    assert result["decision"] == "PASS"
    return result


def test_delivery_builder_creates_safe_non_binding_five_sheet_workbook() -> None:
    pack = _pack()
    product = delivery.build_delivery_artifact(
        pack=pack,
        evaluation=_evaluation(pack),
        requirements_text="=HYPERLINK(\"https://invalid.test\") must remain text",
        tenant_id="tenant-a",
        owner_user_id="owner-a",
        project_id="1" * 32,
        draft_request_id="2" * 32,
        fact_pack_version=1,
        fact_pack_digest=service.fact_pack.fact_pack_digest(pack),
        created_at=datetime(2026, 9, 13, tzinfo=UTC),
    )

    assert len(product.artifact_id) == len(product.work_product_id) == 32
    assert product.artifact_id != product.work_product_id
    assert product.envelope.confirmation_status.value == "PENDING"
    assert product.envelope.artifact_state.value == "PENDING"
    assert product.envelope.work_status.value == "READY_FOR_HUMAN_CONFIRMATION"
    assert product.envelope.artifact_gate.status.value == "PASSED"
    assert {item.kind for item in product.envelope.artifact_manifest} == {
        "confirmation_request",
        "mingshuo_delivery_binding",
        "mingshuo_delivery_projection",
        "mingshuo_fact_pack",
        "mingshuo_solution_quote_xlsx",
    }
    public_facts = json.dumps(product.envelope.model_dump(mode="json")["facts"])
    assert "tenant-a" not in public_facts
    assert "owner-a" not in public_facts

    workbook = load_workbook(BytesIO(product.workbook_bytes), data_only=False, read_only=True)
    assert workbook.sheetnames == ["封面与限制", "事实与证据", "方案草案", "报价草案", "缺失与风险"]
    all_values = [
        cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    ]
    assert any(value.startswith("'=HYPERLINK") for value in all_values)
    assert "NON_BINDING_DRAFT" in all_values
    assert "COMMERCIAL_APPROVAL_REQUIRED" in all_values
    assert all(
        cell.data_type != "f"
        for sheet in workbook.worksheets
        for row in sheet
        for cell in row
    )
    delivery.validate_workbook_bytes(product.workbook_bytes)


def test_delivery_builder_identity_is_deterministic_but_zip_bytes_are_not_identity() -> None:
    pack = _pack()
    kwargs = {
        "pack": pack,
        "evaluation": _evaluation(pack),
        "requirements_text": "Bound requirements",
        "tenant_id": "tenant-a",
        "owner_user_id": "owner-a",
        "project_id": "1" * 32,
        "draft_request_id": "2" * 32,
        "fact_pack_version": 1,
        "fact_pack_digest": service.fact_pack.fact_pack_digest(pack),
        "created_at": datetime(2026, 9, 13, tzinfo=UTC),
    }
    first = delivery.build_delivery_artifact(**kwargs)
    second = delivery.build_delivery_artifact(**kwargs)
    assert first.artifact_id == second.artifact_id
    assert first.work_product_id == second.work_product_id
    assert first.binding_digest == second.binding_digest
    assert first.cell_projection_digest == second.cell_projection_digest


def test_workbook_validator_rejects_formula_and_zip_traversal() -> None:
    pack = _pack()
    product = delivery.build_delivery_artifact(
        pack=pack,
        evaluation=_evaluation(pack),
        requirements_text="Bound requirements",
        tenant_id="tenant-a",
        owner_user_id="owner-a",
        project_id="1" * 32,
        draft_request_id="2" * 32,
        fact_pack_version=1,
        fact_pack_digest=service.fact_pack.fact_pack_digest(pack),
        created_at=datetime(2026, 9, 13, tzinfo=UTC),
    )
    source = ZipFile(BytesIO(product.workbook_bytes))
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            body = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                body = body.replace(b"</c>", b"<f>1+1</f></c>", 1)
            target.writestr(item, body)
    with pytest.raises(delivery.DeliveryValidationError, match="unsafe workbook"):
        delivery.validate_workbook_bytes(output.getvalue())

    traversal = BytesIO()
    with ZipFile(traversal, "w", ZIP_DEFLATED) as archive:
        archive.writestr("../escape.xml", "<xml/>")
    with pytest.raises(delivery.DeliveryValidationError, match="unsafe workbook"):
        delivery.validate_workbook_bytes(traversal.getvalue())


@pytest.fixture
def client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    mingshuo_db = tmp_path / "mingshuo.sqlite3"
    artifact_db = tmp_path / "report_artifacts.sqlite3"
    artifact_dir = tmp_path / "report_artifacts"
    monkeypatch.setattr(storage, "_DEFAULT_DB_PATH", mingshuo_db)
    monkeypatch.setattr(
        service,
        "_artifact_storage",
        lambda: ArtifactStorage(artifact_dir=artifact_dir, db_path=artifact_db),
    )
    report_api.configure_report_artifact_db(artifact_db)
    application = FastAPI()
    application.include_router(mingshuo_api.router)
    application.include_router(report_api.router)
    register_auth_exception_handlers(application)
    application.dependency_overrides[require_current_user] = _owner
    return TestClient(application)


def _ready_draft(client: TestClient) -> tuple[str, str]:
    project = client.post("/api/v1/mingshuo/projects", json=_project_payload()).json()["project"]
    pack = project["currentFactPack"]
    draft = client.post(
        f"/api/v1/mingshuo/projects/{project['projectId']}/draft-requests",
        json={
            "requestKey": "mingshuo-delivery-draft-0001",
            "factPackVersion": pack["version"],
            "factPackDigest": pack["factPackDigest"],
        },
    ).json()["draftRequest"]
    return project["projectId"], draft["draftRequestId"]


def _vary_workbook_container(
    monkeypatch: pytest.MonkeyPatch,
) -> list[str]:
    """Give each build a valid, semantically irrelevant ZIP comment."""

    original = delivery._workbook_bytes  # noqa: SLF001 - controlled container fixture
    digests: list[str] = []
    lock = threading.Lock()

    def varied(cells: list[dict[str, object]]) -> bytes:
        raw = original(cells)
        with lock:
            comment = f"container-{len(digests)}".encode()
            value = raw[:-2] + len(comment).to_bytes(2, "little") + comment
            digests.append(hashlib.sha256(value).hexdigest())
        delivery.validate_workbook_bytes(value)
        return value

    monkeypatch.setattr(delivery, "_workbook_bytes", varied)
    return digests


def test_api_creates_one_pending_reviewable_product_and_replays_exact_identity(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    container_digests = _vary_workbook_container(monkeypatch)
    project_id, draft_id = _ready_draft(client)
    path = f"/api/v1/mingshuo/projects/{project_id}/draft-requests/{draft_id}/work-product"
    created = client.post(path, json={})
    replay = client.post(path, json={})
    assert created.status_code == 201
    assert replay.status_code == 200
    assert len(container_digests) == 2
    assert container_digests[0] != container_digests[1]
    assert created.json() | {"created": False} == replay.json()
    body = created.json()
    assert set(body) == {
        "created",
        "artifactId",
        "workProductId",
        "factPackVersion",
        "factPackDigest",
        "workStatus",
        "confirmationStatus",
        "artifactState",
        "nonAuthorizing",
    }
    assert body["nonAuthorizing"] is True
    assert body["confirmationStatus"] == body["artifactState"] == "PENDING"

    public = client.get(f"/api/v1/report-artifacts/{body['artifactId']}/work-product")
    assert public.status_code == 200
    serialized = json.dumps(public.json(), ensure_ascii=False).lower()
    assert "tenant-a" not in serialized
    assert "owner-a" not in serialized
    assert "requestkey" not in serialized
    assert "requirements_text" not in serialized
    assert client.get(f"/api/v1/report-artifacts/{body['artifactId']}/download").status_code == 404

    confirmed = client.post(
        f"/api/v1/report-artifacts/{body['artifactId']}/confirmation",
        json={"decision": "CONFIRMED", "structured_reason": "Synthetic human review."},
    )
    assert confirmed.status_code == 200
    assert confirmed.json()["confirmation_status"] == "CONFIRMED"
    assert client.get(f"/api/v1/report-artifacts/{body['artifactId']}/download").status_code == 404
    with pytest.raises(ArtifactStorageError):
        service._artifact_storage().publish_run("owner-a", draft_id, "reply-a")


@pytest.mark.parametrize("lagging_state", ["PREPARED", "ARTIFACT_PENDING"])
def test_replay_recovers_lagging_intent_without_using_new_container_sha(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
    lagging_state: str,
) -> None:
    container_digests = _vary_workbook_container(monkeypatch)
    project_id, draft_id = _ready_draft(client)
    path = f"/api/v1/mingshuo/projects/{project_id}/draft-requests/{draft_id}/work-product"
    original_transition = service._transition_delivery_intent  # noqa: SLF001

    def leave_state(
        principal: AuthenticatedPrincipal,
        request_id: str,
        old_state: str,
        new_state: str,
    ) -> None:
        if lagging_state == "PREPARED" or new_state == "WORK_PRODUCT_BOUND":
            return
        original_transition(principal, request_id, old_state, new_state)

    monkeypatch.setattr(service, "_transition_delivery_intent", leave_state)
    created = client.post(path, json={})
    assert created.status_code == 201
    monkeypatch.setattr(service, "_transition_delivery_intent", original_transition)

    replay = client.post(path, json={})

    assert replay.status_code == 200
    assert created.json() | {"created": False} == replay.json()
    assert len(container_digests) == 2
    assert container_digests[0] != container_digests[1]
    with storage.read_connection() as connection:
        state = connection.execute(
            "SELECT state FROM mingshuo_delivery_intents WHERE draft_request_id=?",
            (draft_id,),
        ).fetchone()["state"]
    assert state == "WORK_PRODUCT_BOUND"


def test_concurrent_identical_delivery_requests_bind_one_durable_identity(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    container_digests = _vary_workbook_container(monkeypatch)
    project_id, draft_id = _ready_draft(client)
    path = f"/api/v1/mingshuo/projects/{project_id}/draft-requests/{draft_id}/work-product"
    barrier = threading.Barrier(2)

    def post() -> object:
        barrier.wait(timeout=5)
        return client.post(path, json={})

    with ThreadPoolExecutor(max_workers=2) as executor:
        responses = list(executor.map(lambda _value: post(), range(2)))

    assert sorted(response.status_code for response in responses) == [200, 201]
    assert responses[0].json() | {"created": False} == responses[1].json() | {
        "created": False
    }
    assert len(set(container_digests)) == 2
    report_store = service._artifact_storage()  # noqa: SLF001
    with sqlite3.connect(report_store.db_path) as connection:
        assert connection.execute("SELECT COUNT(*) FROM report_artifacts").fetchone()[0] == 1
        assert connection.execute("SELECT COUNT(*) FROM work_products").fetchone()[0] == 1
        assert connection.execute("SELECT COUNT(*) FROM work_product_artifacts").fetchone()[0] == 1


def test_incomplete_prepared_intent_fails_closed_without_silent_repair(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    _vary_workbook_container(monkeypatch)
    project_id, draft_id = _ready_draft(client)
    path = f"/api/v1/mingshuo/projects/{project_id}/draft-requests/{draft_id}/work-product"
    original_secure = service._secure_pending_file  # noqa: SLF001
    monkeypatch.setattr(
        service,
        "_secure_pending_file",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("synthetic interruption")),
    )
    interrupted = client.post(path, json={})
    assert interrupted.status_code == 503
    monkeypatch.setattr(service, "_secure_pending_file", original_secure)

    replay = client.post(path, json={})

    assert replay.status_code == 503
    assert replay.json() == interrupted.json()
    report_store = service._artifact_storage()  # noqa: SLF001
    with sqlite3.connect(report_store.db_path) as connection:
        assert connection.execute("SELECT COUNT(*) FROM report_artifacts").fetchone()[0] == 0
        assert connection.execute("SELECT COUNT(*) FROM work_products").fetchone()[0] == 0


@pytest.mark.parametrize(
    "tamper",
    [
        "intent_binding",
        "intent_artifact_sha",
        "artifact_file",
        "work_product_payload",
        "work_product_binding",
    ],
)
def test_replay_fails_closed_on_frozen_identity_or_durable_entity_tampering(
    client: TestClient, tamper: str
) -> None:
    project_id, draft_id = _ready_draft(client)
    path = f"/api/v1/mingshuo/projects/{project_id}/draft-requests/{draft_id}/work-product"
    created = client.post(path, json={})
    assert created.status_code == 201
    artifact_id = created.json()["artifactId"]
    work_product_id = created.json()["workProductId"]
    report_store = service._artifact_storage()  # noqa: SLF001
    if tamper.startswith("intent_"):
        column = "binding_digest" if tamper == "intent_binding" else "artifact_sha256"
        with sqlite3.connect(storage._DEFAULT_DB_PATH) as connection:
            connection.execute("DROP TRIGGER mingshuo_delivery_intents_guard_update")
            connection.execute(
                f"UPDATE mingshuo_delivery_intents SET {column}=? WHERE draft_request_id=?",
                ("sha256:" + "e" * 64, draft_id),
            )
    elif tamper == "artifact_file":
        (report_store.artifact_dir / f"{artifact_id}.pending.xlsx").write_bytes(b"tampered")
    elif tamper == "work_product_payload":
        with sqlite3.connect(report_store.db_path) as connection:
            connection.execute("DROP TRIGGER work_products_guard_update")
            connection.execute(
                "UPDATE work_products SET payload_json=? WHERE work_product_id=?",
                ("{}", work_product_id),
            )
    else:
        with sqlite3.connect(report_store.db_path) as connection:
            connection.execute("DROP TRIGGER work_product_artifacts_no_update")
            connection.execute(
                "UPDATE work_product_artifacts SET artifact_id=? WHERE work_product_id=?",
                ("f" * 32, work_product_id),
            )

    replay = client.post(path, json={})

    assert replay.status_code in {409, 503}
    assert replay.json() in (
        {"status": "error", "reason": "conflict"},
        {"status": "error", "reason": "unavailable"},
    )


@pytest.mark.parametrize(
    "raw",
    [b"", b"[]", b'{"extra":true}', b'{"x":1,"x":2}', b"null"],
)
def test_work_product_api_requires_strict_empty_json_object(
    client: TestClient, raw: bytes
) -> None:
    project_id, draft_id = _ready_draft(client)
    response = client.post(
        f"/api/v1/mingshuo/projects/{project_id}/draft-requests/{draft_id}/work-product",
        content=raw,
        headers={"content-type": "application/json"},
    )
    assert response.status_code == 422


def test_cross_tenant_work_product_guess_is_same_404_as_unknown(client: TestClient) -> None:
    project_id, draft_id = _ready_draft(client)
    client.app.dependency_overrides[require_current_user] = lambda: _owner(
        user_id="owner-b", tenant_id="tenant-b"
    )
    cross = client.post(
        f"/api/v1/mingshuo/projects/{project_id}/draft-requests/{draft_id}/work-product",
        json={},
    )
    unknown = client.post(
        f"/api/v1/mingshuo/projects/{'f' * 32}/draft-requests/{'e' * 32}/work-product",
        json={},
    )
    assert cross.status_code == unknown.status_code == 404
    assert cross.json() == unknown.json()


def test_schema_v1_migrates_forward_and_delivery_intent_is_append_only(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(storage, "_DEFAULT_DB_PATH", tmp_path / "mingshuo.sqlite3")
    storage.initialize_database()
    with sqlite3.connect(storage._DEFAULT_DB_PATH) as connection:
        assert connection.execute("PRAGMA user_version").fetchone()[0] == 2
        names = {
            row[0]
            for row in connection.execute("SELECT name FROM sqlite_master WHERE type='trigger'")
        }
    assert {
        "mingshuo_delivery_intents_no_delete",
        "mingshuo_delivery_intents_guard_update",
    } <= names


def test_expired_current_fact_pack_fails_before_artifact_or_intent(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    project_id, draft_id = _ready_draft(client)
    monkeypatch.setattr(service, "_utc_day", lambda: date(2100, 1, 1))
    response = client.post(
        f"/api/v1/mingshuo/projects/{project_id}/draft-requests/{draft_id}/work-product",
        json={},
    )
    assert response.status_code == 409
    assert storage.count_rows("mingshuo_delivery_intents") == 0
