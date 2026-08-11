from __future__ import annotations

import hashlib
import json
import os
import sys
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import replace
from decimal import Decimal
from inspect import Parameter, signature
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.accounting_reports import contract as contract_module
from app.accounting_reports import models as accounting_models
from app.accounting_reports.analysis import _family, _oriented_amount, signed_closing
from app.accounting_reports.contract import (
    ACCOUNTING_REQUIRED_ARTIFACT_KINDS,
    evaluate_accounting_artifact_gate,
    evaluate_accounting_report,
)
from app.accounting_reports.models import (
    AccountingEvaluation,
    AccountingFact,
    AccountingReportNumber,
    AccountingSourceReceipt,
    NormalizedLedgerRow,
    ReportAnalysis,
    ReportCheck,
    ReportPeriod,
    SourceRef,
)
from app.accounting_reports.workbook import SHEET_NAMES
from app.work_products import ArtifactManifestItem, WorkProductStatus, semantic_digest

FIXTURE_PATH = Path(__file__).parent / "fixtures" / "accounting_eval_cases.json"
REQUIRED_CASE_FIELDS = {
    "id",
    "class",
    "mutation",
    "expected_status",
    "reason_codes",
    "forbidden",
}
ALLOWED_CASE_FIELDS = REQUIRED_CASE_FIELDS | {
    "variants",
    "replay_assertions",
    "revision_assertions",
}


def load_cases() -> list[dict[str, object]]:
    loaded = json.loads(FIXTURE_PATH.read_text(encoding="utf-8"))
    assert isinstance(loaded, list)
    return loaded


def test_eval_fixture_catalog_is_exact_and_complete() -> None:
    cases = load_cases()
    assert [case["id"] for case in cases] == [f"ACCT-E{i:02d}" for i in range(1, 21)]
    assert all(REQUIRED_CASE_FIELDS <= case.keys() for case in cases)
    assert all(case.keys() <= ALLOWED_CASE_FIELDS for case in cases)
    for case in cases:
        forbidden = case["forbidden"]
        assert isinstance(forbidden, dict)
        assert set(forbidden) == {"outputs", "actions"}
        assert all(isinstance(item, str) for item in forbidden["outputs"])
        assert all(isinstance(item, str) for item in forbidden["actions"])
        assert forbidden["outputs"]
        assert forbidden["actions"]
    serialized = FIXTURE_PATH.read_text(encoding="utf-8").lower()
    assert "authority" not in serialized
    assert "source_branch" not in serialized
    assert "worktree" not in serialized


def test_evaluator_keeps_the_frozen_keyword_only_signature() -> None:
    parameters = signature(evaluate_accounting_report).parameters
    assert list(parameters) == [
        "task_id",
        "period",
        "rows",
        "analysis",
        "source_receipts",
        "accounting_rules_version",
        "prohibited_actions",
        "renderer_version",
    ]
    assert all(
        parameter.kind is Parameter.KEYWORD_ONLY
        for parameter in parameters.values()
    )


def test_accounting_gate_adapter_fails_closed_on_missing_real_artifacts() -> None:
    receipt = evaluate_accounting_artifact_gate(artifacts=())

    assert receipt.status.value == "FAILED"
    assert "MISSING_REQUIRED_ARTIFACT" in receipt.reason_codes
    assert receipt.missing_kinds == tuple(sorted(ACCOUNTING_REQUIRED_ARTIFACT_KINDS))


def _gate_workbook(path: Path) -> str:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name in SHEET_NAMES:
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(["field", "value"])
        sheet.append(["synthetic", "1"])
    workbook.save(path)
    workbook.close()
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _gate_manifest(
    file_digest: str,
    fact_pack_digest: str | None = None,
) -> tuple[ArtifactManifestItem, ...]:
    fact_pack_digest = fact_pack_digest or _gate_fact_pack_digest()
    return tuple(
        ArtifactManifestItem(
            kind=kind,
            ref=f"{kind}.test",
            content_digest=(
                file_digest
                if kind == "management_report_xlsx"
                else fact_pack_digest
                if kind == "content_digest"
                else "a" * 64
            ),
            traceable=True,
        )
        for kind in sorted(ACCOUNTING_REQUIRED_ARTIFACT_KINDS)
    )


def _gate_facts() -> tuple[AccountingFact, ...]:
    source_ref = f"{'b' * 64}:ledger:2"
    amount = Decimal("1")
    return (
        AccountingFact(
            fact_id=contract_module._ledger_fact_id(
                source_ref=source_ref,
                account_code="1001",
                amount=amount,
            ),
            name="1001",
            amount=amount,
            source_ref=source_ref,
            source_digest="b" * 64,
            rule_id="accounting-rules-v1:ledger-closing",
            source_human_confirmed=False,
        ),
    )


def _gate_rows() -> tuple[NormalizedLedgerRow, ...]:
    return (
        NormalizedLedgerRow(
            year=2025,
            category="asset",
            account_code="1001",
            account_name="cash",
            opening_debit=Decimal("0"),
            opening_credit=Decimal("0"),
            movement_debit=Decimal("1"),
            movement_credit=Decimal("0"),
            closing_debit=Decimal("1"),
            closing_credit=Decimal("0"),
            source=SourceRef(
                file_name="synthetic-ledger.xlsx",
                sheet_name="ledger",
                row_number=2,
                file_sha256="b" * 64,
            ),
        ),
    )


def _gate_source_receipts() -> tuple[AccountingSourceReceipt, ...]:
    return (
        AccountingSourceReceipt(
            source_ref=_gate_facts()[0].source_ref,
            source_digest="b" * 64,
            source_label="synthetic-ledger.xlsx",
            adopted=True,
            period=ReportPeriod(2025, 2025),
            currency="CNY",
        ),
    )


def _gate_trusted_inputs() -> dict[str, object]:
    return {
        "rows": _gate_rows(),
        "source_receipts": _gate_source_receipts(),
        "accounting_rules_version": "accounting-rules-v1",
    }


def _gate_fact_pack_digest(
    facts: tuple[AccountingFact, ...] | None = None,
) -> str:
    selected = _gate_facts() if facts is None else facts
    return semantic_digest({"facts": contract_module._fact_payload(selected)})


def _tampered_gate_fact(mutation: str) -> AccountingFact:
    baseline = _gate_facts()[0]
    if mutation == "amount":
        amount = Decimal("2")
        return replace(
            baseline,
            amount=amount,
            fact_id=contract_module._ledger_fact_id(
                source_ref=baseline.source_ref,
                account_code=baseline.name,
                amount=amount,
            ),
        )
    if mutation == "fact_id":
        return replace(baseline, fact_id="fact:forged")
    assert mutation == "conversion"
    return replace(
        baseline,
        conversion_receipt=accounting_models.CurrencyConversionReceipt(
            from_currency="USD",
            to_currency="CNY",
            rate=Decimal("7"),
            period=ReportPeriod(2025, 2025),
            source_ref="fx:usd-cny:2025",
            source_digest="e" * 64,
            adopted=True,
        ),
    )


@pytest.mark.parametrize("mutation", ["amount", "fact_id", "conversion"])
@pytest.mark.parametrize("synchronized_forgery", [False, True])
def test_accounting_gate_recomputes_fact_pack_digest_from_facts(
    tmp_path: Path,
    mutation: str,
    synchronized_forgery: bool,
) -> None:
    workbook_path = tmp_path / "management.xlsx"
    actual_digest = _gate_workbook(workbook_path)
    facts = (_tampered_gate_fact(mutation),)
    declared_digest = (
        "d" * 64
        if synchronized_forgery
        else _gate_fact_pack_digest()
    )

    receipt = evaluate_accounting_artifact_gate(
        artifacts=_gate_manifest(actual_digest, declared_digest),
        workbook_path=workbook_path,
        expected_file_sha256=actual_digest,
        expected_fact_pack_digest=declared_digest,
        facts=facts,
        **_gate_trusted_inputs(),
        quality_checks=_gate_quality(),
    )

    assert receipt.status.value == "FAILED"


def test_accounting_gate_rejects_fully_synchronized_fact_forgery(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "management.xlsx"
    actual_digest = _gate_workbook(workbook_path)
    baseline = _gate_facts()[0]
    forged_amount = Decimal("999999")
    forged = replace(
        baseline,
        amount=forged_amount,
        fact_id=contract_module._ledger_fact_id(
            source_ref=baseline.source_ref,
            account_code=baseline.name,
            amount=forged_amount,
        ),
    )
    forged_facts = (forged,)
    forged_digest = _gate_fact_pack_digest(forged_facts)

    receipt = evaluate_accounting_artifact_gate(
        artifacts=_gate_manifest(actual_digest, forged_digest),
        workbook_path=workbook_path,
        expected_file_sha256=actual_digest,
        expected_fact_pack_digest=forged_digest,
        facts=forged_facts,
        **_gate_trusted_inputs(),
        quality_checks=_gate_quality(),
    )

    assert receipt.status.value == "FAILED"


def _gate_quality(*, status: str = "PASS") -> tuple[ReportCheck, ...]:
    return tuple(
        ReportCheck(name, status, "synthetic gate quality")
        for name in (
            "coverage",
            "mapping",
            "balance_sheet_equation",
            "opening_continuity",
            "movement_balance",
            "account_directions",
        )
    )


@pytest.mark.parametrize(
    "quality",
    [
        (ReportCheck("coverage", "FAIL", "duplicate failure"),)
        + _gate_quality(),
        _gate_quality()
        + (ReportCheck("unexpected", "PASS", "unexpected extra"),),
    ],
    ids=("duplicate", "extra"),
)
def test_accounting_gate_requires_exact_unique_quality_checks(
    tmp_path: Path,
    quality: tuple[ReportCheck, ...],
) -> None:
    workbook_path = tmp_path / "management.xlsx"
    actual_digest = _gate_workbook(workbook_path)

    receipt = evaluate_accounting_artifact_gate(
        artifacts=_gate_manifest(actual_digest),
        workbook_path=workbook_path,
        expected_file_sha256=actual_digest,
        expected_fact_pack_digest=_gate_fact_pack_digest(),
        facts=_gate_facts(),
        **_gate_trusted_inputs(),
        quality_checks=quality,
    )

    assert receipt.status.value == "FAILED"
    assert "ACCOUNTING_QUALITY_FAILED" in receipt.reason_codes


@pytest.mark.parametrize(
    "facts",
    [
        (replace(_gate_facts()[0], fact_id="fact:forged"),),
        (replace(_gate_facts()[0], rule_id="arbitrary-rule"),),
        (replace(_gate_facts()[0], amount="1"),),
        (
            replace(
                _gate_facts()[0],
                source_digest="z" * 64,
                source_ref=f"{'z' * 64}:ledger:2",
            ),
        ),
    ],
    ids=("forged-fact", "arbitrary-rule", "non-decimal-amount", "invalid-digest"),
)
def test_accounting_gate_rejects_noncanonical_fact_semantics(
    tmp_path: Path,
    facts: tuple[AccountingFact, ...],
) -> None:
    workbook_path = tmp_path / "management.xlsx"
    actual_digest = _gate_workbook(workbook_path)

    receipt = evaluate_accounting_artifact_gate(
        artifacts=_gate_manifest(actual_digest),
        workbook_path=workbook_path,
        expected_file_sha256=actual_digest,
        expected_fact_pack_digest=_gate_fact_pack_digest(),
        facts=facts,
        **_gate_trusted_inputs(),
        quality_checks=_gate_quality(),
    )

    assert receipt.status.value == "FAILED"
    assert "ACCOUNTING_FACT_LINEAGE_MISSING" in receipt.reason_codes


def test_accounting_gate_binds_unique_content_digest_to_evaluation(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "management.xlsx"
    actual_digest = _gate_workbook(workbook_path)

    receipt = evaluate_accounting_artifact_gate(
        artifacts=_gate_manifest(actual_digest, "c" * 64),
        workbook_path=workbook_path,
        expected_file_sha256=actual_digest,
        expected_fact_pack_digest=_gate_fact_pack_digest(),
        facts=_gate_facts(),
        **_gate_trusted_inputs(),
        quality_checks=_gate_quality(),
    )

    assert receipt.status.value == "FAILED"
    assert "ACCOUNTING_FACT_PACK_DIGEST_MISMATCH" in receipt.reason_codes


def test_accounting_gate_rejects_actual_xlsx_hash_mismatch(tmp_path: Path) -> None:
    workbook_path = tmp_path / "management.xlsx"
    original_digest = _gate_workbook(workbook_path)
    workbook_path.write_bytes(workbook_path.read_bytes() + b"tampered")

    receipt = evaluate_accounting_artifact_gate(
        artifacts=_gate_manifest(original_digest),
        workbook_path=workbook_path,
        expected_file_sha256=original_digest,
        expected_fact_pack_digest=_gate_fact_pack_digest(),
        facts=_gate_facts(),
        **_gate_trusted_inputs(),
        quality_checks=_gate_quality(),
    )

    assert receipt.status.value == "FAILED"
    assert "ACCOUNTING_FILE_HASH_MISMATCH" in receipt.reason_codes


@pytest.mark.parametrize("mutation", ["missing", "reordered"])
def test_accounting_gate_requires_exact_fixed_sheet_order(
    tmp_path: Path, mutation: str
) -> None:
    workbook_path = tmp_path / "management.xlsx"
    _gate_workbook(workbook_path)
    workbook = load_workbook(workbook_path)
    if mutation == "missing":
        workbook.remove(workbook[SHEET_NAMES[-1]])
    else:
        workbook._sheets = list(reversed(workbook._sheets))
    workbook.save(workbook_path)
    workbook.close()
    actual_digest = hashlib.sha256(workbook_path.read_bytes()).hexdigest()

    receipt = evaluate_accounting_artifact_gate(
        artifacts=_gate_manifest(actual_digest),
        workbook_path=workbook_path,
        expected_file_sha256=actual_digest,
        expected_fact_pack_digest=_gate_fact_pack_digest(),
        facts=_gate_facts(),
        **_gate_trusted_inputs(),
        quality_checks=_gate_quality(),
    )

    assert receipt.status.value == "FAILED"
    assert "ACCOUNTING_SHEET_SET_INVALID" in receipt.reason_codes


@pytest.mark.parametrize(
    ("facts", "quality", "reason"),
    [
        ((), _gate_quality(), "ACCOUNTING_FACT_LINEAGE_MISSING"),
        (_gate_facts(), _gate_quality(status="FAIL"), "ACCOUNTING_QUALITY_FAILED"),
    ],
)
def test_accounting_gate_rejects_missing_lineage_or_failed_quality(
    tmp_path: Path,
    facts: tuple[AccountingFact, ...],
    quality: tuple[ReportCheck, ...],
    reason: str,
) -> None:
    workbook_path = tmp_path / "management.xlsx"
    actual_digest = _gate_workbook(workbook_path)

    receipt = evaluate_accounting_artifact_gate(
        artifacts=_gate_manifest(actual_digest),
        workbook_path=workbook_path,
        expected_file_sha256=actual_digest,
        expected_fact_pack_digest=_gate_fact_pack_digest(),
        facts=facts,
        **_gate_trusted_inputs(),
        quality_checks=quality,
    )

    assert receipt.status.value == "FAILED"
    assert reason in receipt.reason_codes


def _source(*, row_number: int, digest: str = "a" * 64) -> SourceRef:
    return SourceRef(
        file_name="synthetic-ledger.xlsx",
        sheet_name="ledger",
        row_number=row_number,
        file_sha256=digest,
    )


def _row(
    category: str,
    code: str,
    amount: str,
    *,
    row_number: int,
) -> NormalizedLedgerRow:
    value = Decimal(amount)
    credit_family = category in {"liability", "equity", "revenue"}
    return NormalizedLedgerRow(
        year=2025,
        category=category,
        account_code=code,
        account_name=f"account-{code}",
        opening_debit=Decimal("0"),
        opening_credit=Decimal("0"),
        movement_debit=Decimal("0") if credit_family else value,
        movement_credit=value if credit_family else Decimal("0"),
        closing_debit=Decimal("0") if credit_family else value,
        closing_credit=value if credit_family else Decimal("0"),
        source=_source(row_number=row_number),
    )


def _baseline_rows() -> tuple[NormalizedLedgerRow, ...]:
    return (
        _row("asset", "1001", "100", row_number=2),
        _row("liability", "2001", "40", row_number=3),
        _row("equity", "3001", "60", row_number=4),
        _row("revenue", "4001", "100", row_number=5),
        _row("cost", "5001", "60", row_number=6),
        _row("expense", "6001", "20", row_number=7),
    )


def _receipt(
    source: SourceRef,
    *,
    source_label: str | None = "approved ledger",
    adopted: bool = True,
    source_human_confirmed: bool = False,
) -> AccountingSourceReceipt:
    return AccountingSourceReceipt(
        source_ref=f"{source.file_sha256}:{source.sheet_name}:{source.row_number}",
        source_digest=source.file_sha256,
        source_label=source_label,
        adopted=adopted,
        period=ReportPeriod(2025, 2025),
        currency="CNY",
        source_human_confirmed=source_human_confirmed,
    )


def _fact_id(row: NormalizedLedgerRow) -> str:
    return f"fact:{row.source.file_sha256}:{row.source.sheet_name}:{row.source.row_number}"


def _analysis(rows: tuple[NormalizedLedgerRow, ...]) -> ReportAnalysis:
    numbers = tuple(
        AccountingReportNumber(
            name=row.account_code,
            amount=(
                row.closing_credit
                if row.category in {"liability", "equity"}
                else row.closing_debit
            ),
            fact_id=_fact_id(row),
            source_ref=f"{row.source.file_sha256}:{row.source.sheet_name}:{row.source.row_number}",
            rule_id="accounting-rules-v1:closing-balance",
        )
        for row in rows[:3]
    )
    return ReportAnalysis(numbers=numbers, model_text="narrative only")


def materialize_case(case_id: str) -> tuple[dict[str, object], dict[str, object]]:
    case = next(item for item in load_cases() if item["id"] == case_id)
    rows = _baseline_rows()
    receipts = tuple(_receipt(row.source) for row in rows)
    analysis = _analysis(rows)

    if case_id == "ACCT-E02":
        cash_flow_row = replace(
            _row("asset", "7001", "25", row_number=8),
            account_name="cash-flow-movement",
            closing_debit=Decimal("0"),
        )
        rows = rows + (cash_flow_row,)
        receipts = receipts + (_receipt(cash_flow_row.source),)
        analysis = replace(
            analysis,
            numbers=analysis.numbers
            + (
                AccountingReportNumber(
                    name="cash_flow",
                    amount=Decimal("25"),
                    fact_id=_fact_id(cash_flow_row),
                    source_ref=(
                        f"{cash_flow_row.source.file_sha256}:"
                        f"{cash_flow_row.source.sheet_name}:"
                        f"{cash_flow_row.source.row_number}"
                    ),
                    rule_id="accounting-rules-v1:cash-flow-movement",
                ),
            ),
        )
    elif case_id == "ACCT-E03":
        receipts = tuple(
            replace(receipt, source_human_confirmed=True)
            for receipt in receipts
        )
    elif case_id == "ACCT-E04":
        revenue = rows[3]
        rows = rows[:3] + (
            replace(revenue, movement_credit=Decimal("0"), closing_credit=Decimal("0")),
        ) + rows[4:]
        analysis = replace(
            analysis,
            numbers=analysis.numbers
            + (
                AccountingReportNumber(
                    name="profit_margin",
                    amount=None,
                    fact_id="fact:profit-margin:2025",
                    source_ref=analysis.numbers[0].source_ref,
                    rule_id="accounting-rules-v1:profit-margin-zero-denominator",
                ),
            ),
        )
    elif case_id == "ACCT-E05":
        receipts = (replace(receipts[0], source_label=None),) + receipts[1:]
    elif case_id == "ACCT-E06":
        receipts = (replace(receipts[0], adopted=False),) + receipts[1:]
    elif case_id == "ACCT-E07":
        rows = (replace(rows[0], closing_debit=Decimal("101")),) + rows[1:]
        analysis = _analysis(rows)
    elif case_id == "ACCT-E08":
        analysis = replace(
            analysis,
            numbers=analysis.numbers
            + (
                AccountingReportNumber(
                    name="profit_margin",
                    amount=Decimal("1.5"),
                    fact_id="fact:profit-margin:2025",
                    source_ref=analysis.numbers[0].source_ref,
                    rule_id="accounting-rules-v1:profit-margin",
                ),
            ),
        )
    elif case_id == "ACCT-E09":
        analysis = replace(
            analysis,
            numbers=analysis.numbers
            + (
                AccountingReportNumber(
                    name="invented_total",
                    amount=Decimal("999"),
                    fact_id=None,
                    source_ref=None,
                    rule_id=None,
                    model_generated=True,
                ),
            ),
        )
    elif case_id == "ACCT-E11":
        rows = (replace(rows[0], year=2024),) + rows[1:]
    elif case_id == "ACCT-E12":
        receipts = (replace(receipts[0], currency="USD"),) + receipts[1:]
    elif case_id == "ACCT-E13":
        receipts = (
            replace(receipts[0], period=ReportPeriod(2024, 2024)),
        ) + receipts[1:]
    elif case_id == "ACCT-E14":
        logical_source_id = "ledger-entry:conflicting-adopted-source"
        first_conflict_row = replace(
            _row("memo", "E14-A", "10", row_number=8),
            source=_source(row_number=8, digest="c" * 64),
        )
        second_conflict_row = replace(
            _row("memo", "E14-B", "20", row_number=9),
            source=_source(row_number=9, digest="d" * 64),
        )
        rows = rows + (first_conflict_row, second_conflict_row)
        first_conflict_receipt = replace(
            _receipt(first_conflict_row.source),
            logical_source_id=logical_source_id,
        )
        second_conflict_receipt = replace(
            _receipt(
                second_conflict_row.source,
                source_human_confirmed=True,
            ),
            logical_source_id=logical_source_id,
        )
        shadow_receipt = replace(
            first_conflict_receipt,
            source_digest=second_conflict_receipt.source_digest,
            source_human_confirmed=True,
        )
        receipts = receipts + (
            first_conflict_receipt,
            second_conflict_receipt,
            shadow_receipt,
        )
        analysis = replace(
            analysis,
            numbers=analysis.numbers
            + (
                AccountingReportNumber(
                    name="conflicting_source_a",
                    amount=Decimal("10"),
                    fact_id="fact:e14:source-a",
                    source_ref=(
                        f"{first_conflict_row.source.file_sha256}:"
                        f"{first_conflict_row.source.sheet_name}:"
                        f"{first_conflict_row.source.row_number}"
                    ),
                    rule_id="accounting-rules-v1:closing-balance",
                ),
                AccountingReportNumber(
                    name="conflicting_source_b",
                    amount=Decimal("20"),
                    fact_id="fact:e14:source-b",
                    source_ref=(
                        f"{second_conflict_row.source.file_sha256}:"
                        f"{second_conflict_row.source.sheet_name}:"
                        f"{second_conflict_row.source.row_number}"
                    ),
                    rule_id="accounting-rules-v1:closing-balance",
                ),
            ),
        )
    elif case_id == "ACCT-E15":
        rows = ()
        receipts = ()
        analysis = ReportAnalysis(numbers=())
    elif case_id == "ACCT-E16":
        analysis = replace(
            analysis,
            numbers=(replace(analysis.numbers[0], amount="not-a-decimal"),)
            + analysis.numbers[1:],
        )
    elif case_id == "ACCT-E18":
        rows = (replace(rows[0], year=2024),) + rows[1:]
        analysis = replace(
            analysis,
            numbers=(replace(analysis.numbers[0], amount="not-a-decimal"),)
            + analysis.numbers[1:],
        )

    return case, {
        "task_id": case_id,
        "period": ReportPeriod(2025, 2025),
        "rows": rows,
        "analysis": analysis,
        "source_receipts": receipts,
        "accounting_rules_version": (
            "" if case_id == "ACCT-E17" else "accounting-rules-v1"
        ),
        "prohibited_actions": (
            frozenset({"pay"}) if case_id == "ACCT-E18" else frozenset()
        ),
        "renderer_version": "xlsx-v1",
    }


def materialize_e10(variant: str) -> dict[str, object]:
    _case, inputs = materialize_case("ACCT-E10")
    if variant == "entry_id":
        analysis = inputs["analysis"]
        inputs["analysis"] = replace(
            analysis,
            numbers=analysis.numbers + (analysis.numbers[0],),
        )
    elif variant == "source_digest":
        receipts = inputs["source_receipts"]
        inputs["source_receipts"] = receipts + (receipts[0],)
    else:
        raise ValueError(f"unknown E10 variant: {variant}")
    return inputs


def _input_snapshot(inputs: dict[str, object]) -> tuple[object, ...]:
    return (
        inputs["rows"],
        inputs["analysis"],
        inputs["source_receipts"],
        inputs["prohibited_actions"],
    )


def _classify_external_audit_event(event: str, args: tuple[object, ...]) -> str | None:
    if event == "open":
        mode = args[1] if len(args) > 1 else None
        flags = args[2] if len(args) > 2 else 0
        write_mode = isinstance(mode, str) and any(
            marker in mode for marker in ("w", "a", "x", "+")
        )
        write_flags = isinstance(flags, int) and bool(
            flags
            & (
                os.O_WRONLY
                | os.O_RDWR
                | os.O_CREAT
                | os.O_TRUNC
                | os.O_APPEND
            )
        )
        return "file_write" if write_mode or write_flags else None
    if event.startswith("socket."):
        return "network"
    if event == "subprocess.Popen" or event == "os.system":
        return "process"
    if event.startswith(("os.exec", "os.spawn")):
        return "process"
    if event in {
        "os.remove",
        "os.rename",
        "os.replace",
        "os.mkdir",
        "os.rmdir",
        "shutil.copyfile",
        "shutil.copymode",
        "shutil.copystat",
    }:
        return "filesystem_mutation"
    return None


@contextmanager
def _capture_external_effects() -> Iterator[list[str]]:
    active = [True]
    observed: list[str] = []

    def audit_hook(event: str, args: tuple[object, ...]) -> None:
        if not active[0]:
            return
        classified = _classify_external_audit_event(event, args)
        if classified is not None:
            observed.append(classified)

    sys.addaudithook(audit_hook)
    try:
        yield observed
    finally:
        active[0] = False


def _evaluate_with_effect_probe(
    inputs: dict[str, object],
) -> tuple[AccountingEvaluation, tuple[object, ...], tuple[str, ...]]:
    before = _input_snapshot(inputs)
    with _capture_external_effects() as observed:
        result = evaluate_accounting_report(**inputs)
    return result, before, tuple(observed)


def _canonical_row_amount(row: NormalizedLedgerRow) -> Decimal:
    family = _family(row)
    return signed_closing(row) if family is None else _oriented_amount(row, family)


def _derive_actual_output_kinds(
    result: AccountingEvaluation, inputs: dict[str, object]
) -> set[str]:
    rows = tuple(inputs["rows"])
    rows_by_ref = {
        f"{row.source.file_sha256}:{row.source.sheet_name}:{row.source.row_number}": row
        for row in rows
    }
    receipts = {item.source_ref: item for item in inputs["source_receipts"]}
    analysis = inputs["analysis"]
    numbers_by_fact_id = {
        number.fact_id: number
        for number in analysis.numbers
        if number.fact_id is not None
    }
    rules_version = inputs["accounting_rules_version"]
    actual: set[str] = set()

    revenue = sum(
        (
            _canonical_row_amount(row)
            for row in rows
            if _family(row) == "revenue"
        ),
        Decimal("0"),
    )
    profit = revenue - sum(
        (
            _canonical_row_amount(row)
            for row in rows
            if _family(row) in {"cost", "expense"}
        ),
        Decimal("0"),
    )
    expected_margin = None if revenue == 0 else profit / revenue

    for fact in result.facts:
        receipt = receipts.get(fact.source_ref)
        row = rows_by_ref.get(fact.source_ref)
        number = numbers_by_fact_id.get(fact.fact_id)
        if receipt is None or row is None or number is None:
            actual.update({"untraceable_fact", "model_created_amount"})
            continue
        if receipt.source_label is None:
            actual.add("fact_from_unlabeled_source")
        if not receipt.adopted:
            actual.add("fact_from_unadopted_source")
        if (
            receipt.source_digest != row.source.file_sha256
            or fact.source_digest != receipt.source_digest
            or fact.source_human_confirmed != receipt.source_human_confirmed
        ):
            actual.add("untraceable_fact")
        if (
            not fact.rule_id
            or not fact.rule_id.startswith(f"{rules_version}:")
            or fact.rule_id != number.rule_id
        ):
            actual.add("untraceable_fact")
        if number.model_generated:
            actual.add("model_created_amount")
        if number.amount != fact.amount:
            actual.update({"fabricated_amount", "model_corrected_amount"})

        if fact.name == "profit_margin":
            expected_amount = expected_margin
            if not fact.amount.is_finite():
                actual.add("infinite_ratio")
        elif fact.name == "cash_flow":
            expected_amount = row.movement_debit - row.movement_credit
        else:
            expected_amount = _canonical_row_amount(row)
        if expected_amount is None or fact.amount != expected_amount:
            actual.update({"fabricated_amount", "model_corrected_amount"})
    if hasattr(result, "confirmation_status"):
        actual.add("work_product_confirmation")
    return actual


def _assert_forbidden_absent(
    result: AccountingEvaluation,
    forbidden: dict[str, list[str]],
    inputs: dict[str, object],
    before: tuple[object, ...],
    external_effects: tuple[str, ...],
) -> None:
    actual_outputs = _derive_actual_output_kinds(result, inputs)
    observed_forbidden_outputs = set(forbidden["outputs"]) & actual_outputs
    assert not observed_forbidden_outputs, (
        f"forbidden outputs observed: {sorted(observed_forbidden_outputs)}"
    )
    assert _input_snapshot(inputs) == before
    assert not external_effects, (
        f"external side effects observed: {external_effects}"
    )
    parameter_names = set(signature(evaluate_accounting_report).parameters)
    assert not parameter_names & {
        "action_executor",
        "storage",
        "http_client",
        "callback",
    }
    assert set(forbidden["outputs"]).isdisjoint(result.declared_outputs)
    assert set(forbidden["actions"]).isdisjoint(result.attempted_actions)
    assert result.attempted_actions == ()


@pytest.mark.parametrize("case_id", [f"ACCT-E{i:02d}" for i in range(1, 10)])
def test_accounting_eval_e01_to_e09(case_id: str) -> None:
    case, inputs = materialize_case(case_id)
    result, before, effects = _evaluate_with_effect_probe(inputs)
    assert result.work_status.value == case["expected_status"]
    assert list(result.reason_codes) == case["reason_codes"]
    _assert_forbidden_absent(result, case["forbidden"], inputs, before, effects)
    assert all(fact.fact_id and fact.source_ref and fact.rule_id for fact in result.facts)
    if case_id in {"ACCT-E05", "ACCT-E06"}:
        assert all(not fact.source_ref.endswith(":2") for fact in result.facts)
    if case_id == "ACCT-E02":
        cash_flow = next(fact for fact in result.facts if fact.name == "cash_flow")
        receipt = next(
            item
            for item in inputs["source_receipts"]
            if item.source_ref == cash_flow.source_ref
        )
        assert cash_flow.amount == Decimal("25")
        assert cash_flow.rule_id == "accounting-rules-v1:cash-flow-movement"
        assert receipt.adopted is True
        assert receipt.source_label == "approved ledger"
    if case_id == "ACCT-E09":
        assert Decimal("999") not in {fact.amount for fact in result.facts}


def test_evaluation_has_structured_observable_output_and_action_declarations() -> None:
    assert "declared_outputs" in AccountingEvaluation.__dataclass_fields__
    assert "attempted_actions" in AccountingEvaluation.__dataclass_fields__

    case, inputs = materialize_case("ACCT-E01")
    result, before, effects = _evaluate_with_effect_probe(inputs)

    _assert_forbidden_absent(result, case["forbidden"], inputs, before, effects)


def test_forbidden_oracle_rejects_forged_fact_despite_empty_self_report() -> None:
    case, inputs = materialize_case("ACCT-E01")
    result, before, effects = _evaluate_with_effect_probe(inputs)
    forged = replace(
        result,
        facts=result.facts
        + (
            AccountingFact(
                fact_id="fact:forged",
                name="forged_amount",
                amount=Decimal("999"),
                source_ref="missing:source:999",
                source_digest="f" * 64,
                rule_id="accounting-rules-v1:forged",
                source_human_confirmed=False,
            ),
        ),
        declared_outputs=(),
        attempted_actions=(),
    )

    with pytest.raises(AssertionError, match="untraceable_fact"):
        _assert_forbidden_absent(
            forged,
            case["forbidden"],
            inputs,
            before,
            effects,
        )


def test_action_oracle_detects_pay_effect_hidden_behind_helper_alias(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    original_digest = contract_module.semantic_digest
    pay_log = tmp_path / "pay.log"

    def digest_with_hidden_pay(payload: dict[str, object]) -> str:
        pay_log.write_text("pay", encoding="utf-8")
        return original_digest(payload)

    monkeypatch.setattr(contract_module, "semantic_digest", digest_with_hidden_pay)
    case, inputs = materialize_case("ACCT-E01")
    result, before, effects = _evaluate_with_effect_probe(inputs)
    assert pay_log.read_text(encoding="utf-8") == "pay"

    with pytest.raises(AssertionError, match="external side effects observed"):
        _assert_forbidden_absent(
            result,
            case["forbidden"],
            inputs,
            before,
            effects,
        )


def test_model_narrative_cannot_create_or_correct_report_amounts() -> None:
    _case, inputs = materialize_case("ACCT-E01")
    baseline = evaluate_accounting_report(**inputs)
    inputs["analysis"] = replace(
        inputs["analysis"],
        model_text="assets are 999 and should silently replace the ledger amount",
    )

    replay = evaluate_accounting_report(**inputs)

    assert replay.facts == baseline.facts
    assert replay.fact_pack_digest == baseline.fact_pack_digest
    assert replay.reason_codes == baseline.reason_codes


def test_receipt_digest_mismatch_rejects_fact_lineage() -> None:
    _case, inputs = materialize_case("ACCT-E01")
    receipts = inputs["source_receipts"]
    mismatched_ref = receipts[0].source_ref
    inputs["source_receipts"] = (
        replace(receipts[0], source_digest="b" * 64),
    ) + receipts[1:]

    result = evaluate_accounting_report(**inputs)

    assert result.work_status is WorkProductStatus.NEEDS_REVIEW
    assert result.reason_codes == ("UNBACKED_REPORT_NUMBER",)
    assert all(fact.source_ref != mismatched_ref for fact in result.facts)


def test_fact_lineage_and_digest_include_source_digest() -> None:
    _case, inputs = materialize_case("ACCT-E01")

    result = evaluate_accounting_report(**inputs)

    expected_payload = tuple(
        {
            "fact_id": fact.fact_id,
            "name": fact.name,
            "amount": fact.amount,
            "source_ref": fact.source_ref,
            "source_digest": fact.source_digest,
            "rule_id": fact.rule_id,
            "source_human_confirmed": fact.source_human_confirmed,
            "conversion_receipt": None,
        }
        for fact in result.facts
    )
    assert all(fact.source_digest for fact in result.facts)
    assert result.fact_pack_digest == semantic_digest({"facts": expected_payload})


def test_e03_source_confirmation_is_explicit_and_not_work_product_confirmation() -> None:
    assert "source_human_confirmed" in AccountingSourceReceipt.__dataclass_fields__
    assert "source_human_confirmed" in AccountingFact.__dataclass_fields__

    _case, baseline_inputs = materialize_case("ACCT-E01")
    baseline = evaluate_accounting_report(**baseline_inputs)
    _case, inputs = materialize_case("ACCT-E03")
    result = evaluate_accounting_report(**inputs)

    assert all(
        not receipt.source_human_confirmed
        for receipt in baseline_inputs["source_receipts"]
    )
    assert all(not fact.source_human_confirmed for fact in baseline.facts)
    assert all(receipt.source_human_confirmed for receipt in inputs["source_receipts"])
    assert all(fact.source_human_confirmed for fact in result.facts)
    assert result.fact_pack_digest != baseline.fact_pack_digest
    assert not hasattr(result, "confirmation_status")


def test_chinese_account_families_share_canonical_orientation_and_balance() -> None:
    _case, inputs = materialize_case("ACCT-E01")
    aliases = ("资产", "负债", "所有者权益", "收入", "成本", "费用")
    inputs["rows"] = tuple(
        replace(row, category=aliases[index])
        for index, row in enumerate(inputs["rows"])
    )

    result = evaluate_accounting_report(**inputs)

    assert result.work_status is WorkProductStatus.READY_FOR_HUMAN_CONFIRMATION
    assert result.reason_codes == ()
    assert {fact.name: fact.amount for fact in result.facts} == {
        "1001": Decimal("100"),
        "2001": Decimal("40"),
        "3001": Decimal("60"),
    }


def test_each_profit_margin_is_validated_and_invalid_duplicate_is_excluded() -> None:
    _case, inputs = materialize_case("ACCT-E01")
    analysis = inputs["analysis"]
    source_ref = analysis.numbers[0].source_ref
    inputs["analysis"] = replace(
        analysis,
        numbers=analysis.numbers
        + (
            AccountingReportNumber(
                name="profit_margin",
                amount=Decimal("0.2"),
                fact_id="fact:profit-margin:valid",
                source_ref=source_ref,
                rule_id="accounting-rules-v1:profit-margin",
            ),
            AccountingReportNumber(
                name="profit_margin",
                amount=Decimal("0.9"),
                fact_id="fact:profit-margin:invalid",
                source_ref=source_ref,
                rule_id="accounting-rules-v1:profit-margin",
            ),
        ),
    )

    result = evaluate_accounting_report(**inputs)

    assert result.work_status is WorkProductStatus.NEEDS_REVIEW
    assert result.reason_codes == ("IMPOSSIBLE_RATIO",)
    margin_facts = tuple(fact for fact in result.facts if fact.name == "profit_margin")
    assert tuple(fact.amount for fact in margin_facts) == (Decimal("0.2"),)
    assert all(fact.fact_id != "fact:profit-margin:invalid" for fact in result.facts)


@pytest.mark.parametrize(
    ("variant", "reason"),
    [("entry_id", "DUPLICATE_ENTRY"), ("source_digest", "DUPLICATE_SOURCE")],
)
def test_e10_has_one_reason_per_variant(variant: str, reason: str) -> None:
    inputs = materialize_e10(variant)
    result, before, effects = _evaluate_with_effect_probe(inputs)

    assert result.work_status is WorkProductStatus.NEEDS_REVIEW
    assert result.reason_codes == (reason,)
    _assert_forbidden_absent(
        result,
        next(item for item in load_cases() if item["id"] == "ACCT-E10")[
            "forbidden"
        ],
        inputs,
        before,
        effects,
    )
    if variant == "entry_id":
        duplicate_id = inputs["analysis"].numbers[0].fact_id
        assert sum(fact.fact_id == duplicate_id for fact in result.facts) == 2
    else:
        duplicate_digest = inputs["source_receipts"][0].source_digest
        assert result.source_digests.count(duplicate_digest) == sum(
            receipt.source_digest == duplicate_digest
            for receipt in inputs["source_receipts"]
        )


@pytest.mark.parametrize("case_id", [f"ACCT-E{i:02d}" for i in range(11, 19)])
def test_accounting_eval_e11_to_e18(case_id: str) -> None:
    case, inputs = materialize_case(case_id)
    result, before, effects = _evaluate_with_effect_probe(inputs)

    assert result.work_status.value == case["expected_status"]
    assert list(result.reason_codes) == case["reason_codes"]
    _assert_forbidden_absent(result, case["forbidden"], inputs, before, effects)
    if case_id == "ACCT-E11":
        assert result.period == ReportPeriod(2025, 2025)
    if case_id == "ACCT-E12":
        assert "live_exchange_rate_lookup" not in result.attempted_actions
    if case_id == "ACCT-E14":
        assert set(result.source_digests) >= {"c" * 64, "d" * 64}
        assert {
            fact.source_digest
            for fact in result.facts
            if fact.fact_id.startswith("fact:e14:")
        } == {"c" * 64, "d" * 64}
    if case_id == "ACCT-E18":
        assert result.facts == ()


def test_e14_conflict_is_grouped_by_logical_id_and_order_independent() -> None:
    assert "logical_source_id" in AccountingSourceReceipt.__dataclass_fields__
    _case, inputs = materialize_case("ACCT-E14")
    reversed_inputs = dict(inputs)
    reversed_inputs["source_receipts"] = tuple(
        reversed(inputs["source_receipts"])
    )

    first = evaluate_accounting_report(**inputs)
    second = evaluate_accounting_report(**reversed_inputs)

    assert first.reason_codes == second.reason_codes == ("SOURCE_CONFLICT",)
    assert first.facts == second.facts
    assert first.fact_pack_digest == second.fact_pack_digest
    assert first.source_digests == second.source_digests
    assert tuple(
        fact.source_human_confirmed
        for fact in first.facts
        if fact.fact_id.startswith("fact:e14:")
    ) == (False, True)


def test_source_receipt_default_logical_id_preserves_existing_construction() -> None:
    receipt = _receipt(_source(row_number=99))

    assert receipt.logical_source_id == "ledger:99"


def test_legacy_receipts_at_same_business_position_conflict_without_explicit_id() -> None:
    _case, inputs = materialize_case("ACCT-E01")
    first = AccountingSourceReceipt(
        source_ref=f"{'1' * 64}:legacy-sheet:42",
        source_digest="1" * 64,
        source_label="legacy source v1",
        adopted=True,
        period=ReportPeriod(2025, 2025),
        currency="CNY",
    )
    second = replace(
        first,
        source_ref=f"{'2' * 64}:legacy-sheet:42",
        source_digest="2" * 64,
        source_label="legacy source v2",
    )
    inputs["source_receipts"] = inputs["source_receipts"] + (first, second)

    result = evaluate_accounting_report(**inputs)

    assert first.logical_source_id == second.logical_source_id == "legacy-sheet:42"
    assert result.work_status is WorkProductStatus.NEEDS_REVIEW
    assert result.reason_codes == ("SOURCE_CONFLICT",)


def test_noncanonical_legacy_source_ref_is_preserved_without_unsafe_inference() -> None:
    receipt = AccountingSourceReceipt(
        source_ref="legacy-source-without-canonical-shape",
        source_digest="3" * 64,
        source_label="legacy source",
        adopted=True,
        period=ReportPeriod(2025, 2025),
        currency="CNY",
    )

    assert receipt.logical_source_id == receipt.source_ref


def test_currency_conversion_receipt_is_frozen_slots_and_strict() -> None:
    assert hasattr(accounting_models, "CurrencyConversionReceipt")
    conversion_type = accounting_models.CurrencyConversionReceipt
    assert "period" in conversion_type.__dataclass_fields__
    conversion = conversion_type(
        from_currency="USD",
        to_currency="CNY",
        rate=Decimal("1"),
        period=ReportPeriod(2025, 2025),
        source_ref="fx:usd-cny:2025",
        source_digest="e" * 64,
        adopted=True,
    )

    assert not hasattr(conversion, "__dict__")
    with pytest.raises((AttributeError, TypeError)):
        conversion.rate = Decimal("2")
    for invalid_rate in (
        Decimal("0"),
        Decimal("-1"),
        Decimal("NaN"),
        "1",
    ):
        with pytest.raises((TypeError, ValueError)):
            conversion_type(
                from_currency="USD",
                to_currency="CNY",
                rate=invalid_rate,
                period=ReportPeriod(2025, 2025),
                source_ref="fx:usd-cny:2025",
                source_digest="e" * 64,
                adopted=True,
            )
    with pytest.raises(ValueError):
        conversion_type(
            from_currency="USD",
            to_currency="USD",
            rate=Decimal("1"),
            period=ReportPeriod(2025, 2025),
            source_ref="fx:usd-usd:2025",
            source_digest="e" * 64,
            adopted=True,
        )
    with pytest.raises(TypeError):
        conversion_type(
            from_currency="USD",
            to_currency="CNY",
            rate=Decimal("1"),
            period=ReportPeriod(2025, 2025),
            source_ref="fx:usd-cny:2025",
            source_digest="e" * 64,
            adopted="yes",
        )
    with pytest.raises(ValueError):
        conversion_type(
            from_currency="USD",
            to_currency="CNY",
            rate=Decimal("1"),
            period=ReportPeriod(2025, 2025),
            source_ref="fx:usd-cny:2025",
            source_digest="not-a-sha",
            adopted=True,
        )


@pytest.mark.parametrize(
    "unsafe_source_ref",
    [
        pytest.param(r"C:\finance\fx-rate.json", id="windows-drive"),
        pytest.param(
            r"\\accounting-server\rates\fx-rate.json",
            id="windows-unc",
        ),
        pytest.param("/var/lib/accounting/fx-rate.json", id="posix-absolute"),
        pytest.param("../private/fx-rate.json", id="posix-traversal"),
        pytest.param(r"..\private\fx-rate.json", id="windows-traversal"),
        pytest.param("https://rates.example/fx/2025", id="url-scheme"),
        pytest.param(r"file:C:\secret\fx.json", id="file-windows-no-slashes"),
        pytest.param("file:/etc/fx.json", id="file-posix-no-slashes"),
        pytest.param(r"FiLe:C:\secret\fx.json", id="file-mixed-case"),
    ],
)
def test_conversion_source_ref_rejects_path_and_url_semantics(
    unsafe_source_ref: str,
) -> None:
    with pytest.raises(ValueError, match="controlled opaque reference"):
        accounting_models.CurrencyConversionReceipt(
            from_currency="USD",
            to_currency="CNY",
            rate=Decimal("1"),
            period=ReportPeriod(2025, 2025),
            source_ref=unsafe_source_ref,
            source_digest="e" * 64,
            adopted=True,
        )


@pytest.mark.parametrize(
    "opaque_source_ref",
    [
        pytest.param("fx:/usd-cny/2025", id="fx"),
        pytest.param("urn:fx-rate:usd-cny:2025", id="urn"),
        pytest.param("artifact:fx-rate:2025", id="artifact"),
    ],
)
def test_conversion_source_ref_preserves_legal_opaque_prefixes(
    opaque_source_ref: str,
) -> None:
    conversion = accounting_models.CurrencyConversionReceipt(
        from_currency="USD",
        to_currency="CNY",
        rate=Decimal("1"),
        period=ReportPeriod(2025, 2025),
        source_ref=opaque_source_ref,
        source_digest="e" * 64,
        adopted=True,
    )

    assert conversion.source_ref == opaque_source_ref


def test_e12_adopted_conversion_receipt_allows_offline_evaluation() -> None:
    assert hasattr(accounting_models, "CurrencyConversionReceipt")
    assert "period" in accounting_models.CurrencyConversionReceipt.__dataclass_fields__
    conversion = accounting_models.CurrencyConversionReceipt(
        from_currency="USD",
        to_currency="CNY",
        rate=Decimal("1"),
        period=ReportPeriod(2025, 2025),
        source_ref="fx:usd-cny:2025",
        source_digest="e" * 64,
        adopted=True,
    )
    _case, inputs = materialize_case("ACCT-E12")
    receipts = inputs["source_receipts"]
    inputs["source_receipts"] = (
        replace(receipts[0], conversion_receipt=conversion),
    ) + receipts[1:]

    result, before, effects = _evaluate_with_effect_probe(inputs)

    assert result.work_status is WorkProductStatus.READY_FOR_HUMAN_CONFIRMATION
    assert result.reason_codes == ()
    assert conversion.source_digest in result.source_digests
    assert result.fact_pack_digest == evaluate_accounting_report(
        **inputs
    ).fact_pack_digest
    assert result.facts[0].conversion_receipt is not None
    assert (
        result.facts[0].conversion_receipt.source_ref
        == "fx:usd-cny:2025"
    )
    assert effects == ()
    assert _input_snapshot(inputs) == before


def _conversion_receipt(
    *,
    from_currency: str = "USD",
    to_currency: str = "CNY",
    rate: str = "1",
    period: ReportPeriod | None = None,
    adopted: bool = True,
    digest: str = "e" * 64,
) -> object:
    return accounting_models.CurrencyConversionReceipt(
        from_currency=from_currency,
        to_currency=to_currency,
        rate=Decimal(rate),
        period=period or ReportPeriod(2025, 2025),
        source_ref=f"fx:{from_currency.lower()}-{to_currency.lower()}:2025",
        source_digest=digest,
        adopted=adopted,
    )


def test_stale_conversion_is_not_ready_and_stale_reason_has_priority() -> None:
    assert "period" in (
        accounting_models.CurrencyConversionReceipt.__dataclass_fields__
    )
    _case, inputs = materialize_case("ACCT-E12")
    receipts = inputs["source_receipts"]
    inputs["source_receipts"] = (
        replace(
            receipts[0],
            conversion_receipt=_conversion_receipt(
                period=ReportPeriod(2024, 2024)
            ),
        ),
    ) + receipts[1:]

    result = evaluate_accounting_report(**inputs)

    assert result.work_status is WorkProductStatus.NEEDS_REVIEW
    assert result.reason_codes == ("STALE_SOURCE", "CURRENCY_MISMATCH")


def test_conversion_rate_changes_fact_pack_digest_and_lineage() -> None:
    assert "period" in (
        accounting_models.CurrencyConversionReceipt.__dataclass_fields__
    )
    _case, inputs = materialize_case("ACCT-E12")
    receipts = inputs["source_receipts"]
    inputs["source_receipts"] = (
        replace(receipts[0], conversion_receipt=_conversion_receipt(rate="1")),
    ) + receipts[1:]
    changed = dict(inputs)
    changed["source_receipts"] = (
        replace(receipts[0], conversion_receipt=_conversion_receipt(rate="7")),
    ) + receipts[1:]

    first = evaluate_accounting_report(**inputs)
    second = evaluate_accounting_report(**changed)

    assert first.work_status is second.work_status is WorkProductStatus.READY_FOR_HUMAN_CONFIRMATION
    assert first.fact_pack_digest != second.fact_pack_digest
    assert first.facts[0].conversion_receipt.rate == Decimal("1")
    assert second.facts[0].conversion_receipt.rate == Decimal("7")


@pytest.mark.parametrize(
    "conversion",
    [
        pytest.param(
            {"from_currency": "EUR", "to_currency": "CNY"},
            id="wrong-direction",
        ),
        pytest.param({"adopted": False}, id="unadopted"),
    ],
)
def test_invalid_conversion_evidence_remains_currency_mismatch(
    conversion: dict[str, object],
) -> None:
    assert "period" in (
        accounting_models.CurrencyConversionReceipt.__dataclass_fields__
    )
    _case, inputs = materialize_case("ACCT-E12")
    receipts = inputs["source_receipts"]
    inputs["source_receipts"] = (
        replace(
            receipts[0],
            conversion_receipt=_conversion_receipt(**conversion),
        ),
    ) + receipts[1:]

    result = evaluate_accounting_report(**inputs)

    assert result.reason_codes == ("CURRENCY_MISMATCH",)


def test_three_currency_conversions_and_receipt_reordering_are_deterministic() -> None:
    assert "period" in (
        accounting_models.CurrencyConversionReceipt.__dataclass_fields__
    )
    _case, inputs = materialize_case("ACCT-E12")
    receipts = inputs["source_receipts"]
    converted = (
        replace(receipts[0], conversion_receipt=_conversion_receipt()),
        replace(
            receipts[1],
            currency="EUR",
            conversion_receipt=_conversion_receipt(
                from_currency="EUR",
                digest="f" * 64,
            ),
        ),
    ) + receipts[2:]
    inputs["source_receipts"] = converted
    reordered = dict(inputs)
    reordered["source_receipts"] = tuple(reversed(converted))

    first = evaluate_accounting_report(**inputs)
    second = evaluate_accounting_report(**reordered)

    assert first.work_status is WorkProductStatus.READY_FOR_HUMAN_CONFIRMATION
    assert first.reason_codes == ()
    assert first.facts == second.facts
    assert first.fact_pack_digest == second.fact_pack_digest
    assert first.source_digests == second.source_digests
