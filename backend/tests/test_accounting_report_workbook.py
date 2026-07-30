from __future__ import annotations

import hashlib
import re
from dataclasses import replace
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.accounting_reports.analysis import analyze_ledger, source_id_for_row
from app.accounting_reports.models import NormalizedLedgerRow, ReportPeriod, SourceRef
from app.accounting_reports.workbook import (
    AccountingWorkbookError,
    _audit_generated_workbook,
    write_management_report,
)

SHEETS = [
    "管理摘要",
    "核心财务报表",
    "科目趋势",
    "异常分析",
    "科目明细",
    "校验结果",
    "数据来源",
]


def _row(year: int, category: str, code: str, amount: str) -> NormalizedLedgerRow:
    debit_family = category in {"asset", "cost", "expense"}
    value = Decimal(amount)
    return NormalizedLedgerRow(
        year=year,
        category=category,
        account_code=code,
        account_name=f"科目{code}",
        opening_debit=Decimal("0"),
        opening_credit=Decimal("0"),
        movement_debit=value if debit_family else Decimal("0"),
        movement_credit=Decimal("0") if debit_family else value,
        closing_debit=value if debit_family else Decimal("0"),
        closing_credit=Decimal("0") if debit_family else value,
        source=SourceRef(
            file_name=f"synthetic-{year}.xlsx",
            sheet_name="合成余额表",
            row_number=int(code[-1]) + 1,
            file_sha256=(str(year) * 16)[:64],
        ),
    )


def _rows() -> tuple[NormalizedLedgerRow, ...]:
    return tuple(
        _row(year, category, code, amount)
        for year, values in (
            (2024, ("1000", "400", "600", "500", "200", "100")),
            (2025, ("1200", "500", "700", "600", "240", "120")),
        )
        for (category, code), amount in zip(
            (
                ("asset", "1001"),
                ("liability", "2001"),
                ("equity", "3001"),
                ("revenue", "4001"),
                ("cost", "5001"),
                ("expense", "6001"),
            ),
            values,
            strict=True,
        )
    )


def test_writer_creates_exact_auditable_seven_sheet_workbook(tmp_path: Path) -> None:
    rows = _rows()
    summary = analyze_ledger(rows, ReportPeriod(2024, 2025))
    destination = tmp_path / "management-report.xlsx"

    digest = write_management_report(rows, summary, destination)

    assert digest == hashlib.sha256(destination.read_bytes()).hexdigest()
    workbook = load_workbook(destination, data_only=False)
    assert workbook.sheetnames == SHEETS
    assert workbook.calculation.fullCalcOnLoad is True

    management = workbook["管理摘要"]
    assert management["B2"].value == "2024-2025"
    assert management["B3"].value
    assert management["B4"].value == ", ".join(
        sorted({row.source.file_sha256 for row in rows})
    )
    assert management["B6"].data_type == "f"
    assert management["B6"].value == "='核心财务报表'!H3"
    assert management["B6"].number_format == '#,##0.00;[Red](#,##0.00);-'

    core = workbook["核心财务报表"]
    assert core["H3"].data_type == "f"
    assert core["I3"].data_type == "f"
    assert core["I3"].number_format == "0.00%"

    detail = workbook["科目明细"]
    assert detail.freeze_panes == "A2"
    assert detail.auto_filter.ref == f"A1:N{detail.max_row}"
    assert detail["A2"].number_format == "0"
    assert detail["C2"].number_format == "@"
    assert detail["E2"].data_type == "f"
    assert detail["E2"].value == "=0"
    assert detail["E2"].number_format == '#,##0.00;[Red](#,##0.00);-'

    checks = workbook["校验结果"]
    assert [checks.cell(1, column).value for column in range(1, 7)] == [
        "校验项",
        "状态",
        "实际值",
        "期望值",
        "差异",
        "说明",
    ]
    assert {checks.cell(row, 2).value for row in range(2, checks.max_row + 1)} <= {
        "PASS",
        "FAIL",
    }

    sources = workbook["数据来源"]
    assert {"报告期间", "生成时间", "来源哈希"} <= {
        sources.cell(row, 1).value for row in range(1, sources.max_row + 1)
    }

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formula = str(cell.value)
                    assert "[" not in formula
                    assert not any(
                        marker in formula
                        for marker in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")
                    )
                    assert str(tmp_path) not in formula
    workbook.close()


def test_writer_removes_partial_file_and_sanitizes_unstable_errors(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.accounting_reports import workbook as workbook_module

    destination = tmp_path / "sensitive-name.xlsx"
    monkeypatch.setattr(
        workbook_module,
        "_audit_generated_workbook",
        lambda _path: (_ for _ in ()).throw(RuntimeError(str(destination))),
    )

    with pytest.raises(AccountingWorkbookError) as caught:
        write_management_report(
            _rows(),
            analyze_ledger(_rows(), ReportPeriod(2024, 2025)),
            destination,
        )

    assert str(caught.value) == "workbook_generation_failed"
    assert not destination.exists()


def test_category_aliases_share_analysis_mapping_and_feed_core_formulas(
    tmp_path: Path,
) -> None:
    aliases = ("assets", "liabilities", "所有者权益", "income", "成本", "费用")
    rows = tuple(
        replace(row, category=aliases[index % len(aliases)])
        for index, row in enumerate(_rows())
    )
    destination = tmp_path / "aliases.xlsx"

    write_management_report(
        rows, analyze_ledger(rows, ReportPeriod(2024, 2025)), destination
    )

    workbook = load_workbook(destination, data_only=False)
    detail = workbook["科目明细"]
    assert [detail.cell(row, 14).value for row in range(2, 8)] == [
        "assets",
        "liabilities",
        "equity",
        "revenue",
        "cost",
        "expense",
    ]
    assert "'科目明细'!$N$2:$N$13" in workbook["核心财务报表"]["B2"].value
    expected_criteria = (
        "assets",
        "liabilities",
        "equity",
        "revenue",
        "cost",
        "expense",
    )
    for column, expected in enumerate(expected_criteria, start=2):
        formula = workbook["核心财务报表"].cell(2, column).value
        assert re.findall(r'\$N\$2:\$N\$13,"([^"]+)"', formula) == [
            expected,
            expected,
        ]
    workbook.close()


def test_missing_year_stays_blank_and_does_not_create_false_yoy(
    tmp_path: Path,
) -> None:
    rows = tuple(
        replace(row, year=2026)
        if row.year == 2025
        else row
        for row in _rows()
    )
    summary = analyze_ledger(rows, ReportPeriod(2024, 2026))
    destination = tmp_path / "missing-year.xlsx"

    write_management_report(rows, summary, destination)

    workbook = load_workbook(destination, data_only=False)
    core = workbook["核心财务报表"]
    assert all(core.cell(3, column).value is None for column in range(2, 11))
    assert core["I4"].value is None
    assert core["J4"].value is None
    workbook.close()


def test_high_precision_decimal_is_encoded_without_binary_float(
    tmp_path: Path,
) -> None:
    exact = Decimal("1234567890.123456789")
    rows = (replace(_rows()[0], opening_debit=exact),) + _rows()[1:]
    destination = tmp_path / "decimal.xlsx"

    write_management_report(
        rows, analyze_ledger(rows, ReportPeriod(2024, 2025)), destination
    )

    workbook = load_workbook(destination, data_only=False)
    cell = workbook["科目明细"]["E2"]
    assert cell.data_type == "f"
    assert "1234567890123456789" in cell.value
    assert "1000000000" in cell.value
    workbook.close()


def test_audit_rejects_reference_outside_actual_sheet_bounds(tmp_path: Path) -> None:
    destination = tmp_path / "bad-reference.xlsx"
    rows = _rows()
    write_management_report(
        rows, analyze_ledger(rows, ReportPeriod(2024, 2025)), destination
    )
    workbook = load_workbook(destination)
    workbook["管理摘要"]["B6"] = "='核心财务报表'!Z999"
    workbook.save(destination)
    workbook.close()

    with pytest.raises(ValueError, match="invalid_formula_range"):
        _audit_generated_workbook(destination)


def test_audit_rejects_unquoted_unknown_sheet_reference(tmp_path: Path) -> None:
    destination = tmp_path / "unknown-sheet.xlsx"
    rows = _rows()
    write_management_report(
        rows, analyze_ledger(rows, ReportPeriod(2024, 2025)), destination
    )
    workbook = load_workbook(destination)
    workbook["管理摘要"]["B6"] = "=Unknown!A1"
    workbook.save(destination)
    workbook.close()

    with pytest.raises(ValueError, match="unknown_formula_reference"):
        _audit_generated_workbook(destination)


def test_failure_preserves_preexisting_destination_atomically(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.accounting_reports import workbook as workbook_module

    destination = tmp_path / "existing.xlsx"
    original = b"existing-report"
    destination.write_bytes(original)
    monkeypatch.setattr(
        workbook_module,
        "_audit_generated_workbook",
        lambda _path: (_ for _ in ()).throw(RuntimeError("audit failed")),
    )

    with pytest.raises(AccountingWorkbookError):
        write_management_report(
            _rows(),
            analyze_ledger(_rows(), ReportPeriod(2024, 2025)),
            destination,
        )

    assert destination.read_bytes() == original
    assert list(tmp_path.iterdir()) == [destination]


def test_checks_expose_measured_actual_expected_and_difference(
    tmp_path: Path,
) -> None:
    rows = _rows()[:-1]
    summary = analyze_ledger(rows, ReportPeriod(2024, 2025))
    destination = tmp_path / "failed-check.xlsx"

    write_management_report(rows, summary, destination)

    workbook = load_workbook(destination, data_only=False)
    checks = workbook["校验结果"]
    failed_row = next(
        row
        for row in range(2, checks.max_row + 1)
        if checks.cell(row, 2).value == "FAIL"
    )
    assert checks.cell(failed_row, 3).value != checks.cell(failed_row, 2).value
    assert int(checks.cell(failed_row, 3).value) > 0
    assert checks.cell(failed_row, 4).value == "0"
    assert checks.cell(failed_row, 5).value == checks.cell(failed_row, 3).value
    workbook.close()


def test_sources_map_safe_basename_sheet_source_id_and_hash(
    tmp_path: Path,
) -> None:
    secret_path = tmp_path / "private" / "synthetic-ledger.xlsx"
    first = _rows()[0]
    rows = (
        replace(first, source=replace(first.source, file_name=str(secret_path))),
    ) + _rows()[1:]
    destination = tmp_path / "sources.xlsx"

    write_management_report(
        rows, analyze_ledger(rows, ReportPeriod(2024, 2025)), destination
    )

    workbook = load_workbook(destination)
    sources = workbook["数据来源"]
    headers = [sources.cell(1, column).value for column in range(1, 6)]
    assert headers == ["文件名", "工作表", "来源标识", "来源哈希", "来源行"]
    mapped = next(
        row
        for row in sources.iter_rows(min_row=2, values_only=True)
        if row[0] == "synthetic-ledger.xlsx"
    )
    assert mapped == (
        "synthetic-ledger.xlsx",
        first.source.sheet_name,
        source_id_for_row(rows[0]),
        first.source.file_sha256,
        first.source.row_number,
    )
    assert str(tmp_path) not in repr(list(sources.values))
    assert all(
        str(tmp_path) not in str(cell.value)
        for sheet in workbook
        for row in sheet
        for cell in row
    )
    workbook.close()
