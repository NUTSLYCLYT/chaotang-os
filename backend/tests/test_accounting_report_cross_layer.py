from __future__ import annotations

import hashlib
import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import app.api.decrees as decrees_api
from app.accounting_reports.models import (
    AccountingRequestKind,
    NormalizedLedgerRow,
    ReportPeriod,
    SourceRef,
)
from app.accounting_reports.session import AccountingReportSession
from app.agents.chancellor.graph import build_chancellor_graph
from app.agents.chancellor_draft.authority import (
    AccountingAuthorityContext,
    ConsumedDraftAuthority,
    DraftAuthorityRegistry,
)
from app.agents.chancellor_draft.routing import (
    ApprovedDepartmentRoute,
    ApprovedRouteSnapshot,
)
from app.agents.ministries.agent import invoke_ministry_agent
from app.api.report_artifacts import configure_report_artifact_db
from app.auth import configure_auth_db, create_session, create_user
from app.main import app
from app.shiguan.storage import list_archives
from app.work_products import ConfirmationStatus, WorkProductStatus


def _row(year: int, category: str, code: str, closing: int) -> NormalizedLedgerRow:
    from decimal import Decimal

    debit_oriented = category in {"asset", "cost", "expense"}
    return NormalizedLedgerRow(
        year=year,
        category=category,
        account_code=code,
        account_name=f"synthetic-{code}",
        opening_debit=Decimal("0"),
        opening_credit=Decimal("0"),
        movement_debit=Decimal(closing if debit_oriented else 0),
        movement_credit=Decimal(0 if debit_oriented else closing),
        closing_debit=Decimal(closing if debit_oriented else 0),
        closing_credit=Decimal(0 if debit_oriented else closing),
        source=SourceRef(
            file_name="synthetic-ledger.xlsx",
            sheet_name=f"Synthetic-{year}",
            row_number=int(code[0]) + 1,
            file_sha256="a" * 64,
        ),
    )


SYNTHETIC_ROWS = (
    _row(2025, "asset", "1001", 1000),
    _row(2025, "liability", "2001", 400),
    _row(2025, "equity", "3001", 600),
    _row(2025, "revenue", "4001", 300),
    _row(2025, "cost", "5001", 200),
    _row(2025, "expense", "6001", 100),
)


def test_e19_session_replay_is_semantically_stable_across_runtime_identity(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "app.accounting_reports.session.load_ledger_rows",
        lambda _source, _period: SYNTHETIC_ROWS,
    )

    envelopes = []
    file_digests = []
    generated_at_values = []
    for index in (1, 2):
        root = tmp_path / f"replay-{index}"
        session = AccountingReportSession(
            owner_user_id="owner-e19",
            run_id=f"random-run-{index}",
            source_dir=root / "source",
            artifact_dir=root / "artifacts",
            db_path=root / "artifacts.sqlite3",
        )
        generation = session.maybe_generate(
            "户部",
            "会计司",
            "请生成2025年财务报表",
        )
        assert generation is not None
        envelope = session.storage.get_work_product_for_artifact(
            "owner-e19", generation.artifact_id
        )
        envelopes.append(envelope)
        pending_path = root / "artifacts" / f"{generation.artifact_id}.pending.xlsx"
        file_digests.append(hashlib.sha256(pending_path.read_bytes()).hexdigest())
        workbook = load_workbook(pending_path, data_only=False, read_only=True)
        try:
            assert workbook.properties.created.year == 1980
            assert workbook.properties.modified.year == 1980
            visible_values = tuple(
                cell.value
                for sheet_name in ("管理摘要", "数据来源")
                for row in workbook[sheet_name].iter_rows()
                for cell in row
                if cell.value is not None
            )
        finally:
            workbook.close()
        assert "生成时间" not in visible_values
        assert not any(
            getattr(value, "year", None) == 1980 for value in visible_values
        )
        published = session.publish(f"reply-e19-{index}")
        assert len(published) == 1
        generated_at_values.append(published[0].generated_at)

    first, second = envelopes
    assert first.work_status is second.work_status is (
        WorkProductStatus.READY_FOR_HUMAN_CONFIRMATION
    )
    assert first.confirmation_status is second.confirmation_status is (
        ConfirmationStatus.PENDING
    )
    assert first.artifact_gate.reason_codes == second.artifact_gate.reason_codes == ()
    assert tuple(fact["fact_id"] for fact in first.facts) == tuple(
        fact["fact_id"] for fact in second.facts
    )
    assert first.artifact_manifest == second.artifact_manifest
    assert first.content_digest == second.content_digest
    assert file_digests[0] == file_digests[1]
    assert generated_at_values[0] != generated_at_values[1]
    assert all(value.year != 1980 for value in generated_at_values)
    management_item = next(
        item for item in first.artifact_manifest if item.kind == "management_report_xlsx"
    )
    assert management_item.content_digest == file_digests[0]


def test_session_aborts_pending_artifact_when_envelope_storage_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "app.accounting_reports.session.load_ledger_rows",
        lambda _source, _period: SYNTHETIC_ROWS,
    )
    session = AccountingReportSession(
        owner_user_id="owner-fail-closed",
        run_id="run-fail-closed",
        source_dir=tmp_path / "source",
        artifact_dir=tmp_path / "artifacts",
        db_path=tmp_path / "artifacts.sqlite3",
    )
    artifact_ids: list[str] = []
    original_create_pending = session.storage.create_pending

    def capture_pending(**kwargs):
        pending = original_create_pending(**kwargs)
        artifact_ids.append(pending.artifact_id)
        return pending

    monkeypatch.setattr(session.storage, "create_pending", capture_pending)
    monkeypatch.setattr(
        session.storage,
        "create_work_product",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            RuntimeError("synthetic envelope failure")
        ),
    )

    with pytest.raises(RuntimeError, match="synthetic envelope failure"):
        session.maybe_generate(
            "户部",
            "会计司",
            "请生成2025年财务报表",
        )

    assert len(artifact_ids) == 1
    assert session.storage.get_state(artifact_ids[0]) == "ABORTED"
    assert session.generations == ()


class _AccountingGraph:
    def __init__(self, report_session: AccountingReportSession, route_type: str):
        self.report_session = report_session
        self.route_type = route_type

    def invoke(self, state: dict[str, object]) -> dict[str, object]:
        decree_text = str(state["decree_text"])
        self.report_session.maybe_generate("户部", "会计司", decree_text)
        departments = ["户部"] if self.route_type == "single" else ["户部", "工部"]
        opinions = [
            {
                "department": department,
                "bureau_opinions": [
                    {
                        "bureau": "会计司" if department == "户部" else "技术司",
                        "opinion": "合成意见",
                    }
                ],
                "opinion": "合成部议",
            }
            for department in departments
        ]
        path = ["上书房", "丞相（首次分流）"]
        if self.route_type == "multi":
            path.append("军机处（召集）")
        for department in departments:
            bureau = "会计司" if department == "户部" else "技术司"
            path.extend(
                [department, f"{department}·{bureau}", f"{department}（部级补充）"]
            )
        if self.route_type == "multi":
            path.append("军机处（会审）")
        path.append("丞相（最终汇总）")
        return {
            "decree_text": decree_text,
            "chancellor_rationale": "合成分流",
            "route_type": self.route_type,
            "departments": departments,
            "processing_path": path,
            "ministry_opinions": opinions,
            "council_verdict": "合成会审" if self.route_type == "multi" else None,
            "final_verdict": "合成回奏",
            "recommendations": ["核验来源", "复核勾稽", "审阅报告"],
        }


def test_required_accounting_bureau_alone_generates_and_publishes_excel(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    session = AccountingReportSession(
        owner_user_id="owner-required-bureau",
        run_id="run-required-bureau",
        source_dir=tmp_path / "source",
        artifact_dir=tmp_path / "artifacts",
        db_path=tmp_path / "artifacts.sqlite3",
    )
    monkeypatch.setattr(
        "app.accounting_reports.session.load_ledger_rows",
        lambda _source, _period: SYNTHETIC_ROWS,
    )
    captured: list[list[dict[str, str]]] = []
    responses = iter(
        [
            '{"rationale":"生成财务报表","bureaus":["会计司"]}',
            '{"opinion":"会计司已生成管理报告"}',
            '{"opinion":"户部确认会计管理报告"}',
        ]
    )

    def chat_model(messages: list[dict[str, str]]) -> str:
        captured.append(messages)
        return next(responses)

    result = invoke_ministry_agent(
        "户部",
        "请生成2025年财务报表",
        "批准户部会计司办理",
        chat_model,
        required_bureaus=("会计司",),
        report_session=session,
    )
    published = session.publish("reply-required-bureau")

    assert len(captured) == 3
    assert "你是户部下属的会计司" in captured[1][0]["content"]
    assert "会计司确定性报表摘要" in captured[1][1]["content"]
    assert [item["bureau"] for item in result["bureau_opinions"]] == ["会计司"]
    assert len(published) == 1
    assert published[0].display_name.endswith("会计管理报告.xlsx")
    workbook = load_workbook(published[0].file_path, data_only=False)
    try:
        assert "管理摘要" in workbook.sheetnames
        assert "核心财务报表" in workbook.sheetnames
        assert workbook["核心财务报表"]["A2"].value == 2025
    finally:
        workbook.close()


@pytest.mark.parametrize(
    ("route_type", "expected_departments"),
    [("single", ["户部"])],
)
def test_synthetic_accounting_report_crosses_decree_archive_publish_and_download(
    route_type: str,
    expected_departments: list[str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    isolate_shiguan_default_db_path,
) -> None:
    del isolate_shiguan_default_db_path
    auth_db = tmp_path / "auth.sqlite3"
    artifact_db = tmp_path / "artifacts.sqlite3"
    configure_auth_db(auth_db)
    owner = create_user(f"owner-{route_type}", f"{route_type}@example.test", "six-or-more")
    other = create_user(f"other-{route_type}", f"other-{route_type}@example.test", "six-or-more")
    owner_headers = {"Authorization": f"Bearer {create_session(owner.id)}"}
    other_headers = {"Authorization": f"Bearer {create_session(other.id)}"}
    sessions: list[AccountingReportSession] = []
    accounting_context = AccountingAuthorityContext(
        AccountingRequestKind.ACCOUNTING_REPORT,
        ReportPeriod(2025, 2025),
        "a" * 64,
    )

    def build_session(
        *, owner_user_id: str, run_id: str, accounting_context=accounting_context
    ) -> AccountingReportSession:
        session = AccountingReportSession(
            owner_user_id=owner_user_id,
            run_id=run_id,
            source_dir=tmp_path / "synthetic-source",
            artifact_dir=tmp_path / "report_artifacts",
            db_path=artifact_db,
            request_kind=accounting_context.request_kind,
            period=accounting_context.period,
            dataset=SimpleNamespace(ledger_rows=SYNTHETIC_ROWS),
        )
        sessions.append(session)
        return session

    monkeypatch.setattr(decrees_api, "build_accounting_report_session", build_session)
    monkeypatch.setattr(
        "app.accounting_reports.session.load_ledger_rows",
        lambda _source, _period: SYNTHETIC_ROWS,
    )
    monkeypatch.setattr(
        decrees_api,
        "get_chancellor_graph",
        lambda *, report_session: _AccountingGraph(report_session, route_type),
    )
    monkeypatch.setattr(
        decrees_api.draft_authority_registry,
        "consume",
        lambda **_kwargs: ApprovedRouteSnapshot(
            departments=tuple(
                ApprovedDepartmentRoute(
                    department=department,
                    required_bureaus=(
                        ("会计司",) if department == "户部" else ("技术司",)
                    ),
                )
                for department in expected_departments
            )
        ),
    )
    monkeypatch.setattr(
        decrees_api.draft_authority_registry,
        "consume_with_context",
        lambda **kwargs: ConsumedDraftAuthority(
            decrees_api.draft_authority_registry.consume(**kwargs),
            accounting_context,
        ),
    )

    try:
        with TestClient(app) as client:
            configure_report_artifact_db(artifact_db)
            response = decrees_api.execute_decree_now(
                decrees_api.ChancellorDecreeRequest(
                    decree_text="请生成2025年财务报表"
                ),
                owner,
            )
            body = response.model_dump(mode="json")
            assert body["departments"] == expected_departments
            assert [
                item["department"] for item in body["ministry_opinions"]
            ] == expected_departments
            if route_type == "multi":
                assert body["processing_path"] == [
                    "上书房",
                    "丞相（首次分流）",
                    "军机处（召集）",
                    "户部",
                    "户部·会计司",
                    "户部（部级补充）",
                    "工部",
                    "工部·技术司",
                    "工部（部级补充）",
                    "军机处（会审）",
                    "丞相（最终汇总）",
                ]
            assert len(body["artifacts"]) == 1
            artifact_id = body["artifacts"][0]["artifact_id"]
            assert sessions[0].generations[0].artifact_id == artifact_id
            replies = list_archives(type="REPLY", owner_user_id=owner.id)
            assert len(replies) == 1
            assert sessions[0].storage.get_state(artifact_id) == "PUBLISHED"

            owner_download = client.get(
                f"/api/v1/report-artifacts/{artifact_id}/download",
                headers=owner_headers,
            )
            assert owner_download.status_code == 200
            assert (
                client.get(
                    f"/api/v1/report-artifacts/{artifact_id}/download",
                    headers=other_headers,
                ).status_code
                == 404
            )

        workbook = load_workbook(BytesIO(owner_download.content), data_only=False)
        try:
            assert workbook.sheetnames == [
                "管理摘要",
                "核心财务报表",
                "科目趋势",
                "异常分析",
                "科目明细",
                "校验结果",
                "数据来源",
            ]
            core = workbook["核心财务报表"]
            assert core["A2"].value == 2025
            assert core["B2"].value == (
                "=SUMIFS('科目明细'!$L$2:$L$7,'科目明细'!$A$2:$A$7,$A2,"
                "'科目明细'!$N$2:$N$7,\"assets\")-SUMIFS('科目明细'!$M$2:$M$7,"
                "'科目明细'!$A$2:$A$7,$A2,'科目明细'!$N$2:$N$7,\"assets\")"
            )
            assert core["H2"].value == "=E2-F2-G2"
            assert sum(
                int(str(workbook["科目明细"].cell(row, 12).value).lstrip("="))
                for row in range(2, 8)
                if workbook["科目明细"].cell(row, 14).value == "assets"
            ) == 1000
        finally:
            workbook.close()
    finally:
        configure_report_artifact_db(None)
        configure_auth_db(None)


def test_real_authority_graph_report_archive_download_trust_chain(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    isolate_shiguan_default_db_path,
) -> None:
    del isolate_shiguan_default_db_path
    auth_db = tmp_path / "auth.sqlite3"
    artifact_db = tmp_path / "artifacts.sqlite3"
    configure_auth_db(auth_db)
    owner = create_user("chain-owner", "chain-owner@example.test", "six-or-more")
    other = create_user("chain-other", "chain-other@example.test", "six-or-more")
    owner_headers = {"Authorization": f"Bearer {create_session(owner.id)}"}
    other_headers = {"Authorization": f"Bearer {create_session(other.id)}"}
    decree_text = (
        "请户部会计司根据现有财务数据，生成2024年至2025年管理层综合财务报表，"
        "并交付可下载的 Excel 文件。报告需包括管理摘要、核心财务报表、科目趋势、"
        "异常分析、科目明细、校验结果和数据来源；核对金额、同比变化及勾稽关系，"
        "列明数据缺口，不修改原始数据。"
    )
    draft_version = 9
    draft_fingerprint = "9" * 64
    approved_route = ApprovedRouteSnapshot(
        departments=(
            ApprovedDepartmentRoute(
                department="户部", required_bureaus=("会计司",)
            ),
        )
    )
    two_year_rows = (
        _row(2024, "asset", "1001", 0),
        _row(2024, "liability", "2001", 0),
        _row(2024, "equity", "3001", 0),
        _row(2024, "revenue", "4001", 0),
        _row(2024, "cost", "5001", 0),
        _row(2024, "expense", "6001", 0),
        *SYNTHETIC_ROWS,
    )
    registry = DraftAuthorityRegistry()
    registry.register(
        owner_user_id=owner.id,
        version=draft_version,
        fingerprint=draft_fingerprint,
        decree_text=decree_text,
        route_snapshot=approved_route,
        accounting_context=AccountingAuthorityContext(
            AccountingRequestKind.ACCOUNTING_REPORT,
            ReportPeriod(2024, 2025),
            "a" * 64,
        ),
    )
    monkeypatch.setattr(decrees_api, "draft_authority_registry", registry)
    monkeypatch.setattr(
        "app.agents.chancellor.graph.run_junjichu_council",
        lambda *_args, **_kwargs: pytest.fail(
            "single accounting route must not invoke 军机处"
        ),
    )
    monkeypatch.setattr(
        decrees_api,
        "open_case",
        lambda *_args, **_kwargs: pytest.fail(
            "single accounting route must not open a 军机处 case"
        ),
    )
    monkeypatch.setattr(
        "app.agents.evidence_protocol.invoke_bureau_with_evidence",
        lambda *_args, **_kwargs: pytest.fail(
            "deterministic accounting report must not request 锦衣卫 evidence"
        ),
    )
    monkeypatch.setattr(
        "app.accounting_reports.session.load_ledger_rows",
        lambda _source, _period: SYNTHETIC_ROWS,
    )

    def build_session(
        *, owner_user_id: str, run_id: str, accounting_context
    ) -> AccountingReportSession:
        return AccountingReportSession(
            owner_user_id=owner_user_id,
            run_id=run_id,
            source_dir=tmp_path / "synthetic-source",
                artifact_dir=tmp_path / "report_artifacts",
            db_path=artifact_db,
            request_kind=accounting_context.request_kind,
            period=accounting_context.period,
            dataset=SimpleNamespace(ledger_rows=two_year_rows),
        )

    responses = iter(
        [
            '{"rationale":"批准会计司办理","bureaus":["会计司"]}',
            '{"opinion":"会计司已生成并核验管理报告"}',
            '{"opinion":"户部确认会计管理报告"}',
            (
                '{"summary":"准予交付会计管理报告",'
                '"recommendations":["核验来源","复核勾稽","审阅报告"]}'
            ),
        ]
    )
    monkeypatch.setattr(
        decrees_api, "build_accounting_report_session", build_session
    )
    monkeypatch.setattr(
        decrees_api,
        "get_chancellor_graph",
        lambda *, report_session: build_chancellor_graph(
            owner_user_id=report_session.owner_user_id,
            chat_model=lambda _messages: next(responses),
            lifecycle_observer=decrees_api._lifecycle_observer_context.get(),
            report_session=report_session,
        ),
    )

    try:
        with TestClient(app) as client:
            configure_report_artifact_db(artifact_db)
            response = decrees_api.execute_decree_now(
                decrees_api.ChancellorDecreeRequest(
                    decree_text=decree_text,
                    draft_version=draft_version,
                    draft_fingerprint=draft_fingerprint,
                ),
                owner,
            )
            body = response.model_dump(mode="json")
            assert body["route_type"] == "single"
            assert body["departments"] == ["户部"]
            assert body["processing_path"] == [
                "上书房",
                "丞相（首次分流）",
                "户部",
                "户部·会计司",
                "户部（部级补充）",
                "丞相（最终汇总）",
            ]
            assert body["council_verdict"] is None
            assert [
                item["bureau"]
                for item in body["ministry_opinions"][0]["bureau_opinions"]
            ] == ["会计司"]
            artifact_id = body["artifacts"][0]["artifact_id"]
            download = client.get(
                f"/api/v1/report-artifacts/{artifact_id}/download",
                headers=owner_headers,
            )
            assert download.status_code == 200
            workbook = load_workbook(BytesIO(download.content), data_only=False)
            try:
                assert "管理摘要" in workbook.sheetnames
                assert workbook["核心财务报表"]["A2"].value == 2024
                assert workbook["核心财务报表"]["A3"].value == 2025
            finally:
                workbook.close()

            archives = client.get(
                "/api/v1/shiguan/archives?type=REPLY", headers=owner_headers
            ).json()
            assert len(archives) == 1
            authority = next(
                item
                for item in archives[0]["evidence"]
                if item["source"] == "approved_route_authority"
            )
            assert json.loads(authority["note"]) == {
                "draft_fingerprint": draft_fingerprint,
                "draft_version": draft_version,
                "departments": [
                    {
                        "department": "户部",
                        "required_bureaus": ["会计司"],
                    }
                ],
            }
            assert (
                client.get(
                    f"/api/v1/shiguan/archives/{archives[0]['id']}",
                    headers=other_headers,
                ).status_code
                == 404
            )
            assert (
                client.get(
                    f"/api/v1/report-artifacts/{artifact_id}/download",
                    headers=other_headers,
                ).status_code
                == 404
            )
            assert registry.consume(
                owner_user_id=owner.id,
                version=draft_version,
                fingerprint=draft_fingerprint,
                decree_text=decree_text,
            ) is None
    finally:
        configure_report_artifact_db(None)
        configure_auth_db(None)
